# -*- coding: utf-8 -*-
"""Évaluateur des formules que le moteur écrit dans ses classeurs — remplace le recalcul par Excel.

Périmètre (celui des classeurs de la DEF, prouvé contre Excel par une garde d'équivalence) :
  références (A1, $A$1, 'feuille'!A1, plages A1:B9), nombres, textes ; opérateurs + - * / ^ & = <> < > <= >= et %,
  moins unaire ; fonctions N, IF, ROUND, SUM, MAX, MIN, ABS, AND, OR, NOT, IFERROR.
Sémantique d'Excel : cellule vide = 0 en calcul et "" en concaténation ; comparaison insensible à la casse ;
ROUND arrondit « au plus loin de zéro » ; division par zéro -> "#DIV/0!", texte en calcul -> "#VALUE!".
Une formule hors périmètre lève FormuleNonSupportee : elle est SIGNALÉE, jamais devinée.

Usage : evaluer_classeur(xlsx_formules, xlsx_valeurs) -> {"formules": n, "non_supportees": [...]}
"""
import math
import re
from decimal import ROUND_HALF_UP, Decimal


class FormuleNonSupportee(Exception):
    pass


class Erreur(str):
    """Valeur d'erreur Excel (#DIV/0!, #VALUE!…) : se propage."""


TOKEN = re.compile(r"""
    (?P<ws>\s+)
  | (?P<str>"(?:[^"]|"")*")
  | (?P<ref>(?:(?:'(?:[^']|'')+'|[A-Za-z_][\w.]*)!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?)
  | (?P<func>[A-Z][A-Z0-9.]*)\(
  | (?P<num>\d+(?:\.\d+)?(?:[Ee][+-]?\d+)?)
  | (?P<op><>|<=|>=|[-+*/^&=<>(),%])
  | (?P<bool>TRUE|FALSE)
""", re.VERBOSE)


def tokens(f):
    pos, out = 0, []
    while pos < len(f):
        m = TOKEN.match(f, pos)
        if not m:
            raise FormuleNonSupportee(f"jeton illisible à « {f[pos:pos + 15]} »")
        pos = m.end()
        k = m.lastgroup
        if k != "ws":
            out.append((k, m.group(k)))
    return out


def col_num(s):
    n = 0
    for ch in s:
        n = n * 26 + ord(ch) - 64
    return n


def decouper_ref(r, feuille_courante):
    if "!" in r:
        f, c = r.rsplit("!", 1)
        f = f[1:-1].replace("''", "'") if f.startswith("'") else f
    else:
        f, c = feuille_courante, r
    c = c.replace("$", "")
    if ":" in c:
        a, b = c.split(":")
    else:
        a = b = c
    (ca, ra), (cb, rb) = [re.match(r"([A-Z]+)(\d+)", x).groups() for x in (a, b)]
    return f, col_num(ca), int(ra), col_num(cb), int(rb), ":" in c


def en_nombre(v):
    if isinstance(v, Erreur):
        return v
    if v is None or v == "":
        return 0.0
    if isinstance(v, bool):
        return 1.0 if v else 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return Erreur("#VALUE!")


def en_texte(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "VRAI" if v else "FAUX"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def vrai(v):
    if isinstance(v, str) and not isinstance(v, Erreur):
        return v.strip() != ""
    n = en_nombre(v)
    return not isinstance(n, Erreur) and n != 0


def arrondi(x, n):
    q = Decimal(1).scaleb(-int(n))
    return float(Decimal(repr(x)).quantize(q, rounding=ROUND_HALF_UP))


class Evaluateur:
    def __init__(self, wb):
        self.wb = wb
        self.cache = {}
        self.en_cours = set()
        self.non_supportees = []

    def cellule(self, feuille, col, lig):
        cle = (feuille, col, lig)
        if cle in self.cache:
            return self.cache[cle]
        ws = self.wb[feuille]
        c = ws.cell(lig, col)
        v = c.value
        if isinstance(v, str) and v.startswith("=") and c.data_type == "f":
            if cle in self.en_cours:
                raise FormuleNonSupportee(f"référence circulaire en {feuille}!{c.coordinate}")
            self.en_cours.add(cle)
            try:
                v = self.formule(v[1:], feuille)
            finally:
                self.en_cours.discard(cle)
        elif hasattr(v, "__str__") and type(v).__name__ == "CellRichText":
            v = str(v)
        self.cache[cle] = v
        return v

    def formule(self, f, feuille):
        self.t = tokens(f)
        self.i = 0
        self.feuille = feuille
        v = self.comparaison()
        if self.i != len(self.t):
            raise FormuleNonSupportee(f"reste non analysé : {self.t[self.i:]}")
        return v

    # ---- grammaire (priorités d'Excel)
    def voir(self):
        return self.t[self.i] if self.i < len(self.t) else (None, None)

    def prendre(self, attendu=None):
        k, v = self.voir()
        if attendu and v != attendu:
            raise FormuleNonSupportee(f"attendu « {attendu} », trouvé « {v} »")
        self.i += 1
        return k, v

    def comparaison(self):
        a = self.concat()
        while self.voir()[1] in ("=", "<>", "<", ">", "<=", ">="):
            op = self.prendre()[1]
            b = self.concat()
            a = self.comparer(a, b, op)
        return a

    def comparer(self, a, b, op):
        for x in (a, b):
            if isinstance(x, Erreur):
                return x
        na, nb = a if a is not None else "", b if b is not None else ""
        if isinstance(na, (int, float)) and not isinstance(na, bool) and nb == "":
            nb = 0.0
        if isinstance(nb, (int, float)) and not isinstance(nb, bool) and na == "":
            na = 0.0
        if isinstance(na, str) and isinstance(nb, str):
            na, nb = na.lower(), nb.lower()
        elif isinstance(na, str) != isinstance(nb, str):
            return op == "<>"                      # texte ≠ nombre
        return {"=": na == nb, "<>": na != nb, "<": na < nb, ">": na > nb, "<=": na <= nb, ">=": na >= nb}[op]

    def concat(self):
        a = self.somme()
        while self.voir()[1] == "&":
            self.prendre()
            b = self.somme()
            a = a if isinstance(a, Erreur) else (b if isinstance(b, Erreur) else en_texte(a) + en_texte(b))
        return a

    def somme(self):
        a = self.produit()
        while self.voir()[1] in ("+", "-"):
            op = self.prendre()[1]
            b = self.produit()
            x, y = en_nombre(a), en_nombre(b)
            a = x if isinstance(x, Erreur) else (y if isinstance(y, Erreur) else (x + y if op == "+" else x - y))
        return a

    def produit(self):
        a = self.puissance()
        while self.voir()[1] in ("*", "/"):
            op = self.prendre()[1]
            b = self.puissance()
            x, y = en_nombre(a), en_nombre(b)
            if isinstance(x, Erreur) or isinstance(y, Erreur):
                a = x if isinstance(x, Erreur) else y
            elif op == "/" and y == 0:
                a = Erreur("#DIV/0!")
            else:
                a = x * y if op == "*" else x / y
        return a

    def puissance(self):
        a = self.unaire()
        while self.voir()[1] == "^":
            self.prendre()
            b = self.unaire()
            x, y = en_nombre(a), en_nombre(b)
            a = x ** y if not (isinstance(x, Erreur) or isinstance(y, Erreur)) else Erreur("#VALUE!")
        return a

    def unaire(self):
        if self.voir()[1] in ("-", "+"):
            op = self.prendre()[1]
            v = en_nombre(self.unaire())
            return v if isinstance(v, Erreur) else (-v if op == "-" else v)
        v = self.primaire()
        while self.voir()[1] == "%":
            self.prendre()
            n = en_nombre(v)
            v = n if isinstance(n, Erreur) else n / 100
        return v

    def primaire(self):
        k, v = self.voir()
        if k == "num":
            self.prendre()
            return float(v)
        if k == "str":
            self.prendre()
            return v[1:-1].replace('""', '"')
        if k == "bool":
            self.prendre()
            return v == "TRUE"
        if k == "ref":
            self.prendre()
            f, c0, r0, c1, r1, plage = decouper_ref(v, self.feuille)
            if plage:
                return [self.cellule(f, c, r) for r in range(r0, r1 + 1) for c in range(c0, c1 + 1)]
            return self.cellule(f, c0, r0)
        if k == "func":
            self.prendre()
            return self.fonction(v)
        if v == "(":
            self.prendre()
            x = self.comparaison()
            self.prendre(")")
            return x
        raise FormuleNonSupportee(f"élément inattendu « {v} »")

    def sauter_argument(self):
        """Avance jusqu'à la fin de l'argument courant (virgule ou parenthèse fermante de même niveau)."""
        prof = 0
        while self.i < len(self.t):
            v = self.t[self.i][1]
            if v == "(" or self.t[self.i][0] == "func":
                prof += 1
            elif v == ")":
                if prof == 0:
                    return
                prof -= 1
            elif v == "," and prof == 0:
                return
            self.i += 1

    def fonction(self, nom):
        # repérer les bornes de chaque argument, puis les évaluer À LA DEMANDE (IF n'évalue que la branche retenue)
        bornes = []
        if self.voir()[1] == ")":
            self.prendre()
        else:
            while True:
                d = self.i
                self.sauter_argument()
                bornes.append((d, self.i))
                v = self.prendre()[1]
                if v == ")":
                    break
        fin = self.i

        def arg(k):
            d, f = bornes[k]
            sauve = (self.t, self.i)
            self.t, self.i = self.t[:f], d
            try:
                return self.comparaison()
            finally:
                self.t, self.i = sauve

        def valeurs(k):
            v = arg(k)
            return v if isinstance(v, list) else [v]

        n = len(bornes)
        if nom == "N":
            v = arg(0)
            v = v[0] if isinstance(v, list) else v
            r = float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else (1.0 if v is True else 0.0)
            if isinstance(v, Erreur):
                r = v
        elif nom == "IF":
            r = arg(1) if vrai(arg(0)) else (arg(2) if n > 2 else False)
        elif nom == "IFERROR":
            v = arg(0)
            r = arg(1) if isinstance(v, Erreur) else v
        elif nom in ("SUM", "MAX", "MIN"):
            nums = []
            for k in range(n):
                for v in valeurs(k):
                    if isinstance(v, Erreur):
                        self.i = fin
                        return v
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        nums.append(float(v))
            r = sum(nums) if nom == "SUM" else ((max(nums) if nom == "MAX" else min(nums)) if nums else 0.0)
        elif nom == "ROUND":
            x, d = en_nombre(arg(0)), en_nombre(arg(1)) if n > 1 else 0.0
            r = x if isinstance(x, Erreur) else arrondi(x, d)
        elif nom == "ABS":
            x = en_nombre(arg(0))
            r = x if isinstance(x, Erreur) else abs(x)
        elif nom in ("AND", "OR"):
            vs = [vrai(v) for k in range(n) for v in valeurs(k)]
            r = all(vs) if nom == "AND" else any(vs)
        elif nom == "NOT":
            r = not vrai(arg(0))
        else:
            raise FormuleNonSupportee(f"fonction {nom} hors périmètre")
        self.i = fin
        return r


def evaluer_classeur(source, cible):
    """Classeur à formules -> même classeur, formules remplacées par leurs VALEURS (mise en forme conservée)."""
    import openpyxl
    wb = openpyxl.load_workbook(source)
    ev = Evaluateur(wb)
    resultats, n = [], 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and c.data_type == "f":
                    n += 1
                    try:
                        v = ev.cellule(ws.title, c.column, c.row)
                    except FormuleNonSupportee as e:
                        ev.non_supportees.append(f"{ws.title}!{c.coordinate} : {e}")
                        v = None
                    resultats.append((c, v))
    for c, v in resultats:                           # écrire APRÈS avoir tout évalué (les formules lisent les formules)
        if isinstance(v, list):
            v = v[0] if v else None
        if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
            v = Erreur("#NUM!")
        c.value = str(v) if isinstance(v, Erreur) else v
        if isinstance(v, str):
            c.data_type = "s"
    wb.save(cible)
    return {"formules": n, "non_supportees": ev.non_supportees}
