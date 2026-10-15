"""Vérifie, pour chaque ligne de note que le moteur écrit (correspondances SN d'une couche), que le
libellé de la ligne est le même dans un autre formulaire (même ligne = même rubrique).

  lignes_notes.py <correspondances/sn.json> <reference.xlsx> <candidat.xlsx>
"""
from __future__ import annotations

import difflib
import json
import re
import sys
from pathlib import Path

import openpyxl


def norm(v) -> str:
    s = str(v or "").replace("’", "'").lower()
    return re.sub(r"[^a-zàâçéèêëîïôûùüÿ0-9]+", " ", s).strip()


def libelle(ws, ligne: int) -> str:
    for col in range(1, 4):  # le libellé est en A, B ou C selon les notes
        v = ws.cell(row=ligne, column=col).value
        if isinstance(v, str) and len(norm(v)) > 3:
            return norm(v)
    return ""


def main() -> int:
    sn = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    a = openpyxl.load_workbook(sys.argv[2], read_only=True)
    b = openpyxl.load_workbook(sys.argv[3], read_only=True)
    noms_b = {n.strip().upper(): n for n in b.sheetnames}
    noms_a = {n.strip().upper(): n for n in a.sheetnames}
    total = suspects = 0
    for cle, f in sn["feuilles"].items():
        nom = f["feuille"].strip().upper()
        if nom not in noms_a or nom not in noms_b:
            print(f"{nom} : absente d'un des formulaires"); continue
        wa, wb_ = a[noms_a[nom]], b[noms_b[nom]]
        for ligne in sorted(f["lignes"], key=int):
            total += 1
            la, lb = libelle(wa, int(ligne)), libelle(wb_, int(ligne))
            ratio = difflib.SequenceMatcher(None, la, lb).ratio()
            if ratio < 0.8:
                suspects += 1
                print(f"{nom:14} ligne {ligne:>3} ({ratio:.2f}) « {la[:50]} » ≠ « {lb[:50]} »")
    print(f"lignes écrites par le moteur : {total} · libellé différent (< 0,80) : {suspects}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
