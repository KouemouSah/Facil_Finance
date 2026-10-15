"""Compare deux formulaires Excel de liasse (feuilles, libellés, formules) — pour décider du gabarit d'un pack.

  comparer_formulaires.py <reference.xlsx> <candidat.xlsx> [--rapport fichier.md]

Rapporte : feuilles communes, ajoutées, retirées ; par feuille commune, les cellules dont le texte ou la
formule diffère. Les cellules numériques saisies sont ignorées (un formulaire vierge n'en a pas).
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

import openpyxl


def normaliser(v) -> str:
    if v is None:
        return ""
    s = str(v).replace("’", "'").replace("\xa0", " ")
    return re.sub(r"\s+", " ", s).strip()


def contenu(ws) -> dict[str, str]:
    cellules = {}
    for ligne in ws.iter_rows():
        for c in ligne:
            if c.value is None or isinstance(c.value, (int, float)):
                continue
            t = normaliser(c.value)
            if t:
                cellules[c.coordinate] = t
    return cellules


def comparer(ref: Path, cand: Path) -> dict:
    a = openpyxl.load_workbook(ref, data_only=False)
    b = openpyxl.load_workbook(cand, data_only=False)
    na = {normaliser(n).upper(): n for n in a.sheetnames}
    nb = {normaliser(n).upper(): n for n in b.sheetnames}
    communes = sorted(set(na) & set(nb))
    resultat = {"retirees": sorted(set(na) - set(nb)), "ajoutees": sorted(set(nb) - set(na)), "feuilles": {}}
    for cle in communes:
        ca, cb = contenu(a[na[cle]]), contenu(b[nb[cle]])
        diff = []
        for coord in sorted(set(ca) | set(cb), key=lambda c: (int(re.sub(r"\D", "", c)), c)):
            if ca.get(coord) != cb.get(coord):
                diff.append((coord, ca.get(coord, "∅"), cb.get(coord, "∅")))
        formules_a = sum(1 for v in ca.values() if v.startswith("="))
        formules_b = sum(1 for v in cb.values() if v.startswith("="))
        resultat["feuilles"][cle] = {"diff": diff, "cellules": (len(ca), len(cb)), "formules": (formules_a, formules_b),
                                     "dims": (a[na[cle]].dimensions, b[nb[cle]].dimensions)}
    return resultat


def rapport(r: dict, ref: Path, cand: Path) -> str:
    lignes = [f"# Comparaison de formulaires\n\n- Référence : `{ref.name}`\n- Candidat : `{cand.name}`\n",
              f"Feuilles communes : {len(r['feuilles'])} · retirées : {len(r['retirees'])} · ajoutées : {len(r['ajoutees'])}\n"]
    if r["retirees"]:
        lignes.append("**Retirées dans le candidat** : " + ", ".join(r["retirees"]) + "\n")
    if r["ajoutees"]:
        lignes.append("**Ajoutées dans le candidat** : " + ", ".join(r["ajoutees"]) + "\n")
    identiques = [f for f, d in r["feuilles"].items() if not d["diff"]]
    lignes.append(f"\n**Identiques (textes et formules)** : {len(identiques)} feuille(s)\n")
    lignes.append("\n| Feuille | Cellules diff. | Formules (réf. → cand.) | Exemples |\n|---|---|---|---|")
    for f, d in sorted(r["feuilles"].items(), key=lambda x: -len(x[1]["diff"])):
        if not d["diff"]:
            continue
        ex = " ; ".join(f"{c} « {x[:30]} » → « {y[:30]} »" for c, x, y in d["diff"][:3]).replace("|", "/")
        lignes.append(f"| {f} | {len(d['diff'])} | {d['formules'][0]} → {d['formules'][1]} | {ex} |")
    return "\n".join(lignes) + "\n"


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("reference"); p.add_argument("candidat"); p.add_argument("--rapport")
    a = p.parse_args()
    r = comparer(Path(a.reference), Path(a.candidat))
    texte = rapport(r, Path(a.reference), Path(a.candidat))
    if a.rapport:
        Path(a.rapport).write_text(texte, encoding="utf-8")
    print(texte[:6000])
    return 0


if __name__ == "__main__":
    sys.exit(main())
