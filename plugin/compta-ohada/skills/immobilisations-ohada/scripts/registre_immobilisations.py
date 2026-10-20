# -*- coding: utf-8 -*-
"""Registre d'immobilisations multi-biens (moteur déterministe) — SYSCOHADA.

INTELLIGENT SUR LES COLONNES : détecte automatiquement les colonnes du fichier réel
(alias + normalisation accents/casse/espaces, résolution unique) et repère la ligne
d'en-tête. Override manuel possible via --col-<champ>.

Traite une FLOTTE : dotation/cumul/VNC de chaque bien pour un exercice donné, puis
agrège par catégorie -> alimente la Note 3C de la DSF.

Usage :
  python registre_immobilisations.py --fichier parc.xlsx --exercice 2025
  python registre_immobilisations.py --fichier parc.csv --exercice 2025 --col-valeur_origine "Valeur brute"
"""
import argparse
import csv
import json
import os
import sys
import unicodedata
from datetime import date

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

ALIASES = {
    "designation": ["designation", "libelle", "intitule", "nom", "bien", "immobilisation", "description", "designationdubien"],
    "categorie": ["categorie", "famille", "classe", "typeimmobilisation", "rubrique", "nature", "groupe", "categorieimmo"],
    "valeur_origine": ["valeurorigine", "valeurdorigine", "vo", "valeurbrute", "valeuracquisition", "cout", "coutdacquisition", "valeurinitiale", "prixdacquisition", "montant"],
    "date_service": ["dateservice", "datemiseenservice", "datemise", "dateacquisition", "datedebut", "datedentree", "miseenservice", "dateacq"],
    "duree": ["duree", "dureeutilite", "dureeamortissement", "dureevie", "annees", "dureeenannees", "nbannees"],
    "methode": ["methode", "modeamortissement", "typeamortissement", "mode", "methodeamortissement"],
    "coef": ["coef", "coefficient", "coefdegressif"],
    "valeur_residuelle": ["valeurresiduelle", "vr", "residuelle", "valeurderecuperation"],
}
REQUIRED = ["designation", "valeur_origine", "date_service", "duree"]
OPTIONAL = ["categorie", "methode", "coef", "valeur_residuelle"]


def _norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return "".join(ch for ch in s.lower() if ch.isalnum())


def _score(header_norm, field):
    if not header_norm:
        return 0  # cellule vide : ne matche rien
    als = ALIASES.get(field, [])
    if header_norm in als:
        return 2
    if any(al in header_norm or (len(header_norm) >= 2 and header_norm in al) for al in als):
        return 1
    return 0


def resolve_columns(headers, fields, overrides):
    norm = [_norm(h) for h in headers]
    mapping, used, unresolved = {}, set(), []
    # overrides d'abord
    for f in fields:
        if overrides.get(f):
            ov = _norm(overrides[f])
            idx = next((i for i, h in enumerate(norm) if i not in used and (h == ov or ov in h or h in ov)), None)
            if idx is not None:
                mapping[f] = idx; used.add(idx)
    # puis auto, meilleur score, sans réutiliser une colonne
    for f in fields:
        if f in mapping:
            continue
        best, best_sc = None, 0
        for i, h in enumerate(norm):
            if i in used:
                continue
            sc = _score(h, f)
            if sc > best_sc:
                best, best_sc = i, sc
        if best is not None and best_sc > 0:
            mapping[f] = best; used.add(best)
        elif f in REQUIRED:
            unresolved.append(f)
    return mapping, unresolved


def load_rows(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        rows = [[("" if c is None else c) for c in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
        return rows
    with open(path, encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096); f.seek(0)
        delim = ";" if sample.count(";") >= sample.count(",") else ","
        return [row for row in csv.reader(f, delimiter=delim)]


def detect_header(rows):
    best_i, best_sc = 0, -1
    for i, row in enumerate(rows[:15]):
        headers = [str(c) for c in row]
        m, _ = resolve_columns(headers, REQUIRED, {})
        if len(m) > best_sc:
            best_i, best_sc = i, len(m)
    return best_i


def _num(v):
    if v in (None, ""):
        return 0.0
    s = str(v).replace(" ", "").replace("\xa0", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def _to_date(v):
    if isinstance(v, date):
        return v
    s = str(v).strip()[:10].replace("/", "-")
    parts = s.split("-")
    if len(parts) >= 3:
        if len(parts[0]) == 4:
            y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
        else:  # JJ-MM-AAAA
            d, m, y = int(parts[0]), int(parts[1]), int(parts[2])
        return date(y, m, d)
    raise ValueError(f"Date illisible : {v!r}")


def plan_bien(vo, ds, duree, methode, coef, vr):
    base = vo - vr
    if base <= 0 or duree <= 0:
        return {}
    y0 = ds.year
    prorata1 = (360 - ((ds.month - 1) * 30 + (ds.day - 1))) / 360.0
    annual = base / duree
    taux_deg = (coef / duree) if (methode == "degressif" and coef) else 0
    res, cumul, vnc = {}, 0.0, vo
    for i in range(duree + 1):
        reste = base - cumul
        if reste <= 0.009:
            break
        if methode == "degressif":
            ar = duree - i
            pleine = max((vnc - vr) * taux_deg, (vnc - vr) / ar if ar > 0 else (vnc - vr))
        else:
            pleine = annual
        dot = round(min(pleine * prorata1 if i == 0 else pleine, reste), 2)
        cumul = round(cumul + dot, 2)
        vnc = round(vo - cumul, 2)
        res[y0 + i] = (dot, cumul, vnc)
    return res


def read_biens(path, overrides):
    rows = load_rows(path)
    if not rows:
        raise SystemExit("Fichier vide.")
    hi = detect_header(rows)
    headers = [str(c) for c in rows[hi]]
    mapping, unresolved = resolve_columns(headers, REQUIRED + OPTIONAL, overrides)
    if unresolved:
        raise SystemExit("Colonnes obligatoires introuvables : " + ", ".join(unresolved)
                         + "\nEn-têtes détectés : " + " | ".join(headers)
                         + "\nUtilisez --col-<champ> \"<nom exact>\".")
    print(f"[colonnes détectées | en-tête ligne {hi+1}] "
          + ", ".join(f"{k}='{headers[v]}'" for k, v in mapping.items()), file=sys.stderr)
    biens = []
    for raw in rows[hi + 1:]:
        if len(raw) <= mapping["designation"] or not str(raw[mapping["designation"]]).strip():
            continue
        g = lambda f, d="": raw[mapping[f]] if f in mapping and mapping[f] < len(raw) else d
        biens.append({
            "designation": str(g("designation")).strip(),
            "categorie": (str(g("categorie", "Divers")).strip() or "Divers"),
            "valeur_origine": _num(g("valeur_origine")),
            "date_service": _to_date(g("date_service")),
            "duree": int(_num(g("duree"))),
            "methode": (str(g("methode", "lineaire")).strip().lower() or "lineaire"),
            "coef": _num(g("coef", 0)),
            "valeur_residuelle": _num(g("valeur_residuelle", 0)),
        })
    return biens


def registre(biens, exo):
    lignes, cats = [], {}
    for b in biens:
        plan = plan_bien(b["valeur_origine"], b["date_service"], b["duree"], b["methode"], b["coef"], b["valeur_residuelle"])
        cum_prev = plan.get(exo - 1, (0, 0, b["valeur_origine"]))[1]
        dot, cum, vnc = plan.get(exo, (0.0, cum_prev, round(b["valeur_origine"] - cum_prev, 2)))
        lignes.append({"designation": b["designation"], "categorie": b["categorie"],
                       "valeur_origine": round(b["valeur_origine"], 2), "cumul_anterieur": round(cum_prev, 2),
                       "dotation_exercice": round(dot, 2), "cumul_fin": round(cum, 2), "vnc_fin": round(vnc, 2)})
        c = cats.setdefault(b["categorie"], {"brut": 0.0, "amort": 0.0, "vnc": 0.0, "dotation": 0.0})
        c["brut"] = round(c["brut"] + b["valeur_origine"], 2)
        c["amort"] = round(c["amort"] + cum, 2)
        c["vnc"] = round(c["vnc"] + vnc, 2)
        c["dotation"] = round(c["dotation"] + dot, 2)
    return {"exercice": exo, "lignes": lignes, "note_3c": cats}


def main():
    p = argparse.ArgumentParser(description="Registre d'immobilisations (détection auto des colonnes).")
    p.add_argument("--fichier", required=True)
    p.add_argument("--exercice", type=int, required=True)
    p.add_argument("--json", action="store_true")
    for f in REQUIRED + OPTIONAL:
        p.add_argument(f"--col-{f}", default=None)
    a = p.parse_args()
    overrides = {f: getattr(a, f"col_{f}") for f in REQUIRED + OPTIONAL if getattr(a, f"col_{f}")}
    res = registre(read_biens(a.fichier, overrides), a.exercice)
    if a.json:
        print(json.dumps(res, ensure_ascii=False, indent=2)); return
    print(f"Registre — exercice {res['exercice']} — {len(res['lignes'])} biens")
    for l in res["lignes"]:
        print(f"  {l['designation'][:26]:26s} {l['categorie'][:16]:16s} VO {l['valeur_origine']:>14,.0f} "
              f"dot {l['dotation_exercice']:>13,.0f} cum {l['cumul_fin']:>14,.0f} VNC {l['vnc_fin']:>14,.0f}")
    print("\nNote 3C — par catégorie (brut / amort / VNC / dotation) :")
    for cat, d in res["note_3c"].items():
        print(f"  {cat[:22]:22s} {d['brut']:>15,.0f} {d['amort']:>15,.0f} {d['vnc']:>15,.0f} {d['dotation']:>14,.0f}")


if __name__ == "__main__":
    main()
