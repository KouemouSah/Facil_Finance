# -*- coding: utf-8 -*-
"""Note Markdown -> PDF A4 présentable (destiné au comptable, au dirigeant, à l'auditeur).

Même chaîne que les rapports de la DEF : un classeur mis en page (openpyxl) puis exporté en PDF par Excel,
LibreOffice à défaut. Mise en forme :
  - couverture (logo facultatif, titre, entité, date, mention de diffusion) et sommaire paginé ;
  - une section de niveau 2 (« ## ») = une page (au moins) ; sous-titres « ### » soulignés ;
  - tableaux : en-tête sur bandeau bleu, lignes alternées, montants alignés à droite, texte renvoyé à la ligne ;
  - repères d'état ÉCRITS, jamais la couleur seule : ✅ -> « OK », ⚠️ -> « À VOIR », ❌ -> « ÉCART » ;
  - encadrés : blocs « > » (compléments) et lignes « **Question** » ;
  - gras Markdown (**…**) restitué ; pied de page « titre — page n / N ».
Usage :
  python md_vers_pdf.py --md note.md --pdf note.pdf --titre "..." [--sous-titre ..] [--entite ..] [--logo ..]
         [--mention ..] [--pied ..] [--xlsx garder.xlsx]
"""
import argparse
import datetime
import math
import os
import re
import subprocess
import sys
import tempfile

BLEU, BLEU2, CLAIR, JAUNE, ZEBRE = "1F3864", "2F5597", "EEF3FA", "FFF8E5", "F7F9FC"
NCOL, LARG = 12, 7.6                       # grille de 12 colonnes (≈ 91 caractères de large, A4 portrait)
BADGES = {"✅": ("OK", "E2EFDA", "375623"), "⚠️": ("À VOIR", "FFF2CC", "7F6000"), "⚠": ("À VOIR", "FFF2CC", "7F6000"),
          "❌": ("ÉCART", "F8CBAD", "833C0B")}
NUM_RX = re.compile(r"^[\(\-+−]?\d[\d\s\u00a0\u202f.,/]*\)?\s*(%|XAF|M|Md)?$")


def sans_md(s):
    s = re.sub(r"`([^`]*)`", r"\1", s)
    s = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", s)
    return s.replace("\\|", "|")


def riche(texte, taille=9.5, couleur="1A1A1A", italique=False):
    """Texte avec **gras** Markdown -> texte riche Excel (repli : texte brut)."""
    from openpyxl.cell.rich_text import CellRichText, TextBlock
    from openpyxl.cell.text import InlineFont
    texte = re.sub(r"(?<!\*)\*(?!\*)", "", sans_md(texte))      # italique Markdown (*…*) : marques retirées
    for sym, (lab, _, _) in BADGES.items():
        texte = texte.replace(sym, f"[{lab}]")
    morceaux = re.split(r"(\*\*[^*]+\*\*)", texte)
    if len(morceaux) == 1:
        return texte.replace("*", "") if texte.count("*") >= 2 else texte
    rt = CellRichText()
    for m in morceaux:
        if not m:
            continue
        gras = m.startswith("**") and m.endswith("**")
        t = m[2:-2] if gras else re.sub(r"(?<!\w)\*([^*]+)\*(?!\w)", r"\1", m)
        rt.append(TextBlock(InlineFont(rFont="Calibri", sz=taille, b=gras, i=italique, color=couleur), t))
    return rt


def blocs(md):
    """Découpe Markdown -> [(type, contenu)] : h1, h2, h3, table, liste, citation, para."""
    lignes, out, i = md.splitlines(), [], 0
    while i < len(lignes):
        l = lignes[i].rstrip()
        if not l.strip() or l.strip() == "---":
            i += 1
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", l)
        if m:
            out.append((f"h{len(m.group(1))}", m.group(2).strip()))
            i += 1
            continue
        if l.lstrip().startswith("|"):
            t = []
            while i < len(lignes) and lignes[i].lstrip().startswith("|"):
                cells = [c.strip() for c in re.split(r"(?<!\\)\|", lignes[i].strip().strip("|"))]
                if not all(re.match(r"^:?-{2,}:?$", c) for c in cells if c):
                    t.append(cells)
                i += 1
            out.append(("table", t))
            continue
        if l.lstrip().startswith(">"):
            q = []
            while i < len(lignes) and lignes[i].lstrip().startswith(">"):
                q.append(lignes[i].lstrip()[1:].strip())
                i += 1
            out.append(("citation", q))
            continue
        if re.match(r"^\s*([-*]|\d+[.)])\s+", l):
            items = []
            while i < len(lignes) and (re.match(r"^\s*([-*]|\d+[.)])\s+", lignes[i]) or
                                       (lignes[i].startswith("  ") and lignes[i].strip() and items)):
                x = lignes[i]
                m2 = re.match(r"^(\s*)([-*]|\d+[.)])\s+(.*)", x)
                if m2:
                    items.append((len(m2.group(1)) // 2, m2.group(2), m2.group(3)))
                else:
                    niv, puce, txt = items[-1]
                    items[-1] = (niv, puce, txt + " " + x.strip())
                i += 1
            out.append(("liste", items))
            continue
        p = [l.strip()]
        i += 1
        while i < len(lignes) and lignes[i].strip() and not re.match(r"^(#|\||>|\s*([-*]|\d+[.)])\s)", lignes[i].lstrip()):
            p.append(lignes[i].strip())
            i += 1
        out.append(("para", " ".join(p)))
    return out


def construire(md_txt, a, xlsx):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.page import PageMargins
    fill = lambda c: PatternFill("solid", fgColor=c)
    fin = Side(style="thin", color="D0D7E2")
    wb = Workbook()
    feuilles, sections = [], []

    def feuille(titre):
        ws = wb.active if not feuilles else wb.create_sheet()
        ws.title = (re.sub(r"[\[\]\*\?/\\:]", "", titre)[:28] or "Section") + ("" if titre not in feuilles else f" {len(feuilles)}")
        feuilles.append(ws.title)
        for j in range(1, NCOL + 1):
            ws.column_dimensions[ws.cell(1, j).column_letter].width = LARG
        ws.page_setup.paperSize, ws.page_setup.orientation = ws.PAPERSIZE_A4, "portrait"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins = PageMargins(left=0.55, right=0.55, top=0.6, bottom=0.65, header=0.25, footer=0.3)
        ws.print_options.horizontalCentered = True
        ws.oddFooter.left.text = (a.pied or a.titre)[:80]
        ws.oddFooter.right.text = "Page &P / &N"
        for p in (ws.oddFooter.left, ws.oddFooter.right):
            p.size, p.color = 7.5, "7F7F7F"
        ws.sheet_view.showGridLines = False
        return ws

    def hauteur(texte, largeur_car, taille=9.5):
        n = sum(max(1, math.ceil(len(seg) * (taille / 9.5) / max(largeur_car, 1))) for seg in str(texte).split("\n"))
        return max(14, n * 12.8 * taille / 9.5 + 3)

    def zone(ws, r, c0, c1, valeur, font=None, fond=None, align=None, bord=None, taille=9.5):
        if c1 > c0:
            ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=c1)
        c = ws.cell(r, c0, valeur)
        c.font = font or Font(name="Calibri", size=taille)
        c.alignment = align or Alignment(wrap_text=True, vertical="top")
        if fond or bord:
            for j in range(c0, c1 + 1):
                if fond:
                    ws.cell(r, j).fill = fill(fond)
                if bord:
                    ws.cell(r, j).border = bord
        return c

    def texte_long(ws, texte, fond=None, gauche=1, italique=False, couleur="1A1A1A", gras=False, taille=9.5):
        r = ws.max_row + 1
        v = riche(texte, taille, couleur, italique) if not gras else sans_md(texte).replace("**", "")
        zone(ws, r, gauche, NCOL, v, font=Font(name="Calibri", size=taille, bold=gras, italic=italique, color=couleur), fond=fond)
        brut = re.sub(r"\*\*", "", sans_md(texte))
        ws.row_dimensions[r].height = hauteur(brut, (NCOL - gauche + 1) * LARG * 1.08, taille)
        return r

    def table(ws, rows):
        n = max(len(x) for x in rows)
        rows = [x + [""] * (n - len(x)) for x in rows]
        # largeur de chaque colonne proportionnelle à son contenu (au moins 1 colonne de grille)
        # colonnes de MONTANTS : assez larges pour le plus long montant sur une seule ligne ; le reste au prorata du texte
        def est_num(v):
            b = sans_md(v).replace("**", "").strip()
            return bool(NUM_RX.match(b)) and any(ch.isdigit() for ch in b)
        mini = []
        for j in range(n):
            vals = [sans_md(x[j]).replace("**", "").strip() for x in rows[1:]]
            nums = [v for v in vals if est_num(rows[1 + vals.index(v)][j]) or est_num(v)]
            if nums:
                mini.append(min(4, max(1, math.ceil((max(len(v) for v in nums) + 1.5) / LARG))))
            else:                                       # colonne de libellés : au moins 2 colonnes de grille si le texte est long
                mini.append(2 if max((len(v) for v in vals), default=0) > 10 else 1)
        poids = [max(4, min(45, max(len(sans_md(x[j]).replace("**", "")) for x in rows))) for j in range(n)]
        libre = NCOL - sum(mini)
        spans = list(mini)
        if libre > 0:
            tot = sum(poids)
            for j in range(n):
                spans[j] += round(libre * poids[j] / tot)
        while sum(spans) > NCOL:
            k = max(range(n), key=lambda j: spans[j] - mini[j])
            spans[k] -= 1
        while sum(spans) < NCOL:
            spans[max(range(n), key=lambda j: poids[j])] += 1
        bornes, c = [], 1
        for s in spans:
            bornes.append((c, c + s - 1))
            c += s
        for k, ligne in enumerate(rows):
            r = ws.max_row + 1
            h = 14
            for j, v in enumerate(ligne):
                c0, c1 = bornes[j]
                brut = sans_md(v).replace("**", "")
                badge = next((b for s, b in BADGES.items() if s in v), None)
                if k == 0:
                    cell = zone(ws, r, c0, c1, brut, font=Font(name="Calibri", size=9, bold=True, color="FFFFFF"), fond=BLEU,
                                align=Alignment(wrap_text=True, vertical="center", horizontal="center"))
                else:
                    num = bool(NUM_RX.match(brut.strip())) and any(ch.isdigit() for ch in brut)
                    cell = zone(ws, r, c0, c1, riche(v, 9), fond=ZEBRE if k % 2 == 0 else None,
                                align=Alignment(wrap_text=True, vertical="top", horizontal="right" if num else None),
                                bord=Border(bottom=fin), taille=9)
                    if badge:
                        lab, fnd, encre = badge
                        reste = brut
                        for s in BADGES:
                            reste = reste.replace(s, "").strip()
                        cell.value = lab + ((" — " + reste) if reste else "")
                        cell.font = Font(name="Calibri", size=9, bold=True, color=encre)
                        for jj in range(c0, c1 + 1):
                            ws.cell(r, jj).fill = fill(fnd)
                h = max(h, hauteur(brut, (c1 - c0 + 1) * LARG * 1.05, 9))
            ws.row_dimensions[r].height = h
        ws.append([])

    # ---------------- couverture
    ws = feuille("Couverture")
    for k in range(1, 15):                                  # espace réservé au logo : 14 lignes de 13 pt (≈ 240 px)
        ws.row_dimensions[k].height = 13
    ws.cell(14, 1, "")
    if a.logo and os.path.exists(a.logo):
        from openpyxl.drawing.image import Image
        from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
        from openpyxl.drawing.xdr import XDRPositiveSize2D
        from openpyxl.utils.units import pixels_to_EMU
        im = Image(a.logo)
        h = 120
        w = int(im.width * h / im.height)
        largeur_px = int(NCOL * LARG * 7 + 5)
        im.anchor = OneCellAnchor(_from=AnchorMarker(col=0, row=3, colOff=pixels_to_EMU(max((largeur_px - w) // 2, 0))),
                                  ext=XDRPositiveSize2D(pixels_to_EMU(w), pixels_to_EMU(h)))
        ws.add_image(im)
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for texte, taille, gras, couleur in ((a.titre, 24, True, BLEU), (a.sous_titre, 13, False, BLEU2), ("", 10, False, "000000"),
                                         (a.entite, 17, True, "000000"), ("", 10, False, "000000")):
        r = ws.max_row + 1
        zone(ws, r, 1, NCOL, texte, font=Font(name="Calibri", size=taille, bold=gras, color=couleur), align=centre)
        ws.row_dimensions[r].height = max(18, taille * 1.9) if texte else 12
    r = ws.max_row + 1
    zone(ws, r, 1, NCOL, f"Établi le {a.date or datetime.date.today().strftime('%d/%m/%Y')}",
         font=Font(name="Calibri", size=10, color="595959"), align=centre)
    if a.mention:
        r = ws.max_row + 2
        zone(ws, r, 3, NCOL - 2, a.mention, font=Font(name="Calibri", size=10, bold=True, color="C55A11"), align=centre,
             bord=Border(top=Side(style="thin", color="C55A11"), bottom=Side(style="thin", color="C55A11")))
        ws.row_dimensions[r].height = 20

    # ---------------- sommaire (paginé au second passage)
    som = feuille("Sommaire")

    # ---------------- corps
    md_txt = "\n".join(l for l in md_txt.splitlines() if not l.startswith("# "))
    intro, courant = [], None
    for typ, v in blocs(md_txt):
        if typ == "h2":
            courant = feuille(re.sub(r"^(Partie\s+\d+|[A-Z]\.|\d+\.)\s*[—–-]?\s*", "", v) or v)
            sections.append((v, courant.title, []))
            r = courant.max_row + 1
            zone(courant, r, 1, NCOL, v.replace("**", ""), font=Font(name="Calibri", size=14, bold=True, color="FFFFFF"), fond=BLEU,
                 align=Alignment(vertical="center", indent=1))
            courant.row_dimensions[r].height = 26
            courant.append([])
            continue
        if courant is None:
            intro.append((typ, v))
            continue
        ws = courant
        if typ in ("h3", "h4"):
            ws.append([])
            r = ws.max_row + 1
            zone(ws, r, 1, NCOL, sans_md(v).replace("**", ""), font=Font(name="Calibri", size=12 if typ == "h3" else 10.5,
                 bold=True, color=BLEU), bord=Border(bottom=Side(style="medium", color=BLEU2)) if typ == "h3" else None)
            ws.row_dimensions[r].height = 20
            sections[-1][2].append(sans_md(v).replace("**", ""))
        elif typ == "table":
            table(ws, v)
        elif typ == "citation":
            for k, q in enumerate(x for x in v if x):
                r = texte_long(ws, ("• " + q[2:]) if q.startswith("- ") else q, fond=CLAIR, gauche=1)
                ws.cell(r, 1).border = Border(left=Side(style="thick", color=BLEU2))
            ws.append([])
        elif typ == "liste":
            for niv, puce, txt in v:
                question = txt.lstrip().startswith("**Question")
                marque = puce if puce[0].isdigit() else ("▸" if niv else "•")
                texte_long(ws, f"{marque} {txt}", fond=JAUNE if question else None, gauche=1 + min(niv, 2))
            ws.append([])
        elif typ == "para":
            question = v.lstrip().startswith("**Question")
            italique = v.startswith("*") and not v.startswith("**")
            texte_long(ws, v.strip("*") if italique else v, fond=JAUNE if question else None, italique=italique,
                       couleur="595959" if italique else "1A1A1A")

    def ecrire_sommaire(pages):
        for rg in list(som.merged_cells.ranges):           # défusionner AVANT d'effacer (cellule fusionnée = lecture seule)
            som.unmerge_cells(str(rg))
        for row in som.iter_rows():
            for c in row:
                c.value = None
        r = 2
        zone(som, r, 1, NCOL, "Sommaire", font=Font(name="Calibri", size=16, bold=True, color=BLEU),
             bord=Border(bottom=Side(style="medium", color=BLEU)))
        som.row_dimensions[r].height = 26
        r += 2
        for titre, f, sous in sections:
            zone(som, r, 1, NCOL - 1, titre.replace("**", ""), font=Font(name="Calibri", size=10.5, bold=True))
            zone(som, r, NCOL, NCOL, pages.get(f, ""), font=Font(name="Calibri", size=10.5, bold=True),
                 align=Alignment(horizontal="right"))
            som.row_dimensions[r].height = 17
            r += 1
            for s in sous:
                zone(som, r, 2, NCOL, s, font=Font(name="Calibri", size=9, color="595959"))
                r += 1
            r += 1
        for typ, v in intro:                          # paragraphes d'introduction de la note, sous le sommaire
            if typ in ("para", "liste"):
                for t in ([v] if typ == "para" else [f"• {x[2]}" for x in v]):
                    zone(som, r, 1, NCOL, riche(t.strip("*"), 9, "595959", True))
                    som.row_dimensions[r].height = hauteur(sans_md(t), NCOL * LARG * 1.08, 9)
                    r += 1

    ecrire_sommaire({})
    wb.save(xlsx)
    return wb, feuilles, sections, ecrire_sommaire


def exporter_pdf(xlsx, pdf):
    ps = rf"""
$x = New-Object -ComObject Excel.Application; $x.Visible = $false; $x.DisplayAlerts = $false
try {{ $wb = $x.Workbooks.Open('{os.path.abspath(xlsx)}', 0, $true)
  $wb.ExportAsFixedFormat(0, '{os.path.abspath(pdf)}', 0, $true, $false); $wb.Close($false); 'OK' }} finally {{ $x.Quit() }}
"""
    if os.environ.get("FACIL_PDF", "").lower() != "natif":
        try:
            cp = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True, timeout=600)
            if "OK" in cp.stdout and os.path.exists(pdf):
                return "excel"
        except (OSError, subprocess.TimeoutExpired):
            pass
    try:                                                  # sans Excel : rendu natif (PyMuPDF)
        sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
        import rendu_pdf
        if rendu_pdf.classeur_vers_pdf(xlsx, pdf):
            return "natif"
    except Exception:
        pass
    for soffice in ("soffice", r"C:\Program Files\LibreOffice\program\soffice.exe"):
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", os.path.dirname(os.path.abspath(pdf)), xlsx],
                           capture_output=True, timeout=300)
            gen = os.path.join(os.path.dirname(os.path.abspath(pdf)), os.path.splitext(os.path.basename(xlsx))[0] + ".pdf")
            if os.path.exists(gen):
                os.replace(gen, pdf)
                return "libreoffice"
        except (OSError, subprocess.TimeoutExpired):
            continue
    return None


def main():
    p = argparse.ArgumentParser(description="Note Markdown -> PDF A4 présentable (via Excel).")
    p.add_argument("--md", required=True)
    p.add_argument("--pdf", required=True)
    p.add_argument("--titre", required=True)
    p.add_argument("--sous-titre", default="")
    p.add_argument("--entite", default="")
    p.add_argument("--logo")
    p.add_argument("--mention", default="Document de travail — diffusion restreinte")
    p.add_argument("--pied", default="")
    p.add_argument("--date", default="")
    p.add_argument("--xlsx", help="conserver le classeur intermédiaire à ce chemin")
    a = p.parse_args()
    x = a.xlsx or os.path.join(tempfile.mkdtemp(), "note.xlsx")
    wb, feuilles, sections, ecrire_sommaire = construire(open(a.md, encoding="utf-8").read(), a, x)
    if os.path.exists(a.pdf):
        os.remove(a.pdf)
    mode = exporter_pdf(x, a.pdf)
    if mode:
        try:                                            # second passage : numéros de page réels au sommaire
            try:
                import pymupdf as fitz            # nom actuel de PyMuPDF
            except ImportError:
                import fitz                       # ancien nom (versions < 1.24.3)
            doc = fitz.open(a.pdf)
            pages, i_sec = {}, 0
            # couverture et sommaire citent tous les titres : on ne cherche qu'APRÈS la page du sommaire
            p_som = next((i for i, pg in enumerate(doc) if pg.get_text().lstrip().startswith("Sommaire")), 1)
            import unicodedata                              # NFKC : la ligature « ﬁ » du rendu natif redevient « fi »
            for i, pg in enumerate(doc):
                txt = re.sub(r"\s+", " ", unicodedata.normalize("NFKC", pg.get_text()))
                if i <= p_som:
                    continue
                while i_sec < len(sections) and re.sub(r"\s+", " ", unicodedata.normalize("NFKC", sections[i_sec][0].replace("**", "")))[:30] in txt:
                    pages[sections[i_sec][1]] = i + 1
                    i_sec += 1
            doc.close()
            ecrire_sommaire(pages)
            wb.save(x)
            mode = exporter_pdf(x, a.pdf) or mode
        except Exception as e:                              # le PDF reste valable, sans numéros de page au sommaire
            print(f"Sommaire non paginé : {e}")
    print(f"PDF : {a.pdf} ({mode})" if mode else "PDF NON PRODUIT : ni Excel ni LibreOffice")
    sys.exit(0 if mode else 1)


if __name__ == "__main__":
    main()
