# -*- coding: utf-8 -*-
"""Détection d'une balance comptable quelconque par INVARIANTS COMPTABLES.

SOURCE UNIQUE : compta-ohada/lib/detection_balance.py. Les copies présentes dans
skills/*/scripts/ sont SYNCHRONISÉES par `python tools/sync_lib.py` (chaque skill est
zippé seul pour l'import Team) ; `tools/package.py` refuse d'empaqueter si une copie
diverge. Ne jamais éditer une copie.

Principe : les colonnes ne sont PAS reconnues par leur nom, mais par les propriétés
comptables de leur contenu. Les en-têtes (FR/ES/EN) ne donnent qu'un bonus plafonné :
ils manquent souvent, sont ambigus (« Débit » x4) ou dans une autre langue.

Invariants :
  - équilibre   : Σ débits = Σ crédits par paire (Σ = 0 pour une colonne signée) ;
  - exclusivité : une paire de SOLDES n'a jamais débit ET crédit sur une même ligne ;
  - continuité  : ouverture + mouvements = solde ;
  - cumul       : cumul = ouverture + mouvements ;
  - gestion     : classes 6-7 nulles à l'ouverture, actives dans un solde (N ou N-1) ;
  - hiérarchie  : une ligne n'est un cumul que si elle VAUT la somme de ses sous-comptes.

Statuts : « ok » ; « ambigu » (deux lectures donnent des SOLDES N différents : l'utilisateur
tranche, puis relance avec gabarit=...) ; « illisible ». « a_confirmer » liste les lectures
qui ne diffèrent que sur les colonnes secondaires (sans effet sur les soldes N).

Formats : .xls (xlrd), .xlsx/.xlsm (openpyxl), .csv/.txt. Toutes les feuilles d'un classeur
sont essayées ; la plus convaincante est retenue. PDF : non pris en charge ici (voir
documentation/analyse/decision-import-balances.md, D3).
"""
import csv
import os
import re
import unicodedata
from collections import Counter, defaultdict

__all__ = ["analyse", "analyse_grid", "LAYOUTS"]


# --------------------------------------------------------------------------- lecture
def read_sheets(path):
    """[(nom_feuille, lignes)]."""
    ext = os.path.splitext(path)[1].lower()
    if ext == ".xls":
        try:
            import xlrd
        except ImportError:
            raise SystemExit("Lecture .xls : installer xlrd (pip install xlrd).")
        wb = xlrd.open_workbook(path)
        return [(s.name, [s.row_values(r) for r in range(s.nrows)]) for s in wb.sheets()]
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        out = [(ws.title, [[("" if c is None else c) for c in r] for r in ws.iter_rows(values_only=True)])
               for ws in wb.worksheets]
        wb.close()
        return out
    if ext == ".pdf":
        raise SystemExit("PDF non pris en charge par ce moteur (extraction de tableau à ajouter). "
                         "Exporter la balance en Excel ou CSV.")
    with open(path, encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096); f.seek(0)
        delim = max(";,\t", key=sample.count)
        return [("csv", [row for row in csv.reader(f, delimiter=delim)])]


# --------------------------------------------------------------------------- typage des cellules
_NUM = re.compile(r"^\(?-?\d[\d\s .,']*\)?-?$")


def guess_decimal(rows):
    """Séparateur décimal des montants écrits en texte : vote sur le DERNIER séparateur."""
    votes = Counter()
    for r in rows[:600]:
        for v in r:
            if isinstance(v, str) and _NUM.match(v.strip()):
                s = v.strip().strip("()-")
                seps = [c for c in s if c in ".,"]
                if not seps:
                    continue
                last = seps[-1]; tail = s.rsplit(last, 1)[1]
                if len(set(seps)) == 2 or len(tail) != 3:
                    votes[last] += 1                          # « 1.234,56 » ou « 12,5 »
                else:
                    votes["." if last == "," else ","] += 1   # « 1,234 » : milliers
    return "," if votes[","] >= votes["."] else "."


def parse_number(v, dec):
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s or not _NUM.match(s):
        return None
    neg = (s.startswith("(") and s.endswith(")")) or s.endswith("-") or s.startswith("-")
    s = s.strip("()-").replace(" ", "").replace(" ", "").replace("'", "")
    s = s.replace("." if dec == "," else ",", "").replace(",", ".")
    try:
        x = float(s)
    except ValueError:
        return None
    return -x if neg else x


_ACC = re.compile(r"^([1-9]\d{1,11})(?:\.0)?(?:\s+(.*))?$")


def split_account(v):
    """(compte, libellé éventuellement collé) ou None. N'est appliquée qu'à la PREMIÈRE
    cellule non vide d'une ligne : un montant entier plus à droite n'est jamais un compte."""
    if isinstance(v, float):
        return (str(int(v)), "") if v.is_integer() and 10 <= v < 1e12 else None
    m = _ACC.match(str(v).strip())
    return (m.group(1), (m.group(2) or "").strip()) if m else None


def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", s)


HINTS = {  # indices FR / ES / EN — bonus seulement
    "O": ["ouverture", "anouveau", "report", "initial", "debut", "apertura", "opening"],
    "M": ["mouvement", "mvt", "periode", "movimiento", "movement"],
    "C": ["cumul", "acumulado"],
    "S": ["solde", "saldo", "cloture", "final", "closing"],
    "P": ["n1", "precedent", "anterior", "previous"],
    "D": ["debit", "debe", "deudor"],
    "K": ["credit", "haber", "acreedor"],
}

# Rôles : O ouverture, M mouvements, C cumuls, S soldes N, P soldes N-1 ;
# suffixe D/K = débit/crédit, s = colonne signée (D - C).
LAYOUTS = {
    "8 col. (ouverture, mouvements, cumuls, soldes)": ["OD", "OK", "MD", "MK", "CD", "CK", "SD", "SK"],
    "6 col. (ouverture, mouvements, soldes)": ["OD", "OK", "MD", "MK", "SD", "SK"],
    "6 col. (mouvements, cumuls, soldes)": ["MD", "MK", "CD", "CK", "SD", "SK"],
    "4 col. (mouvements, soldes)": ["MD", "MK", "SD", "SK"],
    "4 col. (ouverture, soldes)": ["OD", "OK", "SD", "SK"],
    "4 col. (soldes N-1, soldes N)": ["PD", "PK", "SD", "SK"],
    "4 col. (soldes N, soldes N-1)": ["SD", "SK", "PD", "PK"],
    "2 col. (soldes)": ["SD", "SK"],
    "3 col. signées (ouverture, mouvements, solde)": ["Os", "Ms", "Ss"],
    "2 col. signées (solde N-1, solde N)": ["Ps", "Ss"],
    "2 col. signées (solde N, solde N-1)": ["Ss", "Ps"],
    "1 col. (solde signé)": ["Ss"],
}


def _net(v, idx, f):
    if f + "D" in idx:
        return v[idx[f + "D"]] - v[idx[f + "K"]]
    if f + "s" in idx:
        return v[idx[f + "s"]]
    return None


def _solde(v, idx):
    for f in ("S", "C", "M"):
        x = _net(v, idx, f)
        if x is not None:
            return x
    raise ValueError("gabarit sans solde")


# --------------------------------------------------------------------------- analyse
def analyse(path, gabarit=None):
    """Analyse toutes les feuilles et retient la plus convaincante (confiance, puis nb de comptes)."""
    best = None
    for name, raw in read_sheets(path):
        r = analyse_grid(raw, gabarit)
        r["feuille"] = name
        key = (r["statut"] != "illisible", r.get("confiance", 0), r.get("feuilles", 0))
        if best is None or key > best[0]:
            best = (key, r)
    return best[1] if best else {"statut": "illisible", "raison": "fichier vide", "soldes": {}}


def analyse_grid(raw, gabarit=None):
    if gabarit is not None and gabarit not in LAYOUTS:
        raise ValueError(f"gabarit inconnu : {gabarit!r}. Choix : {list(LAYOUTS)}")
    dec = guess_decimal(raw)
    acc_rows, header_text = [], defaultdict(list)
    for i, r in enumerate(raw):
        acc = None; lab = ""; nums = []
        first = next((j for j, v in enumerate(r) if str(v).strip()), None)
        for j, v in enumerate(r):
            if acc is None and j == first and j <= 3:
                sa = split_account(v)
                if sa:
                    acc, lab = sa
                    continue
            n = parse_number(v, dec)
            if acc is not None and n is not None:
                nums.append((j, n))
            elif isinstance(v, str) and v.strip():
                if acc is not None and not lab:
                    lab = v.strip()
                elif acc is None:
                    header_text[j].append(norm(v))
        if acc is not None and nums:
            acc_rows.append({"i": i, "acc": acc, "lab": lab, "nums": nums})
    if not acc_rows:
        return {"statut": "illisible", "raison": "aucune ligne « compte + montants » reconnue", "soldes": {}}

    # 1) feuilles candidates (comptes qui n'en prolongent aucun autre) : elles fixent les colonnes
    accs = sorted({r["acc"] for r in acc_rows})
    has_child = {a for a in accs if any(b != a and b.startswith(a) for b in accs)}
    cand = [r for r in acc_rows if r["acc"] not in has_child]
    presence = Counter(c for r in cand for c, _ in r["nums"])
    cols = sorted(c for c, n in presence.items() if n >= 0.05 * len(cand))
    k = len(cols)
    if not k:
        return {"statut": "illisible", "raison": "aucune colonne de montants stable", "soldes": {}}
    vec = lambda r, shift=0: [dict(r["nums"]).get(c + shift, 0.0) for c in cols]
    # une ligne dont AUCUN montant ne tombe dans les colonnes de la balance n'est pas un compte
    # (date d'édition « 46210 », n° de page…) : sans cela, elle devient un compte fantôme à 0
    cand = [r for r in cand if any(c in cols for c, _ in r["nums"])]

    # 2) parents : cumul SI ET SEULEMENT SI ils valent la somme de leurs sous-comptes
    leaves = [(r, vec(r)) for r in cand]

    def same(v, tot):
        """Égalité colonne à colonne, OU du net de chaque paire (« solde des soldes »)."""
        if all(abs(a - b) < 1 for a, b in zip(v, tot)):
            return True
        return k % 2 == 0 and all(abs((v[p] - v[p + 1]) - (tot[p] - tot[p + 1])) < 1 for p in range(0, k, 2))

    cumuls, parents_reels = 0, []
    for r in sorted((x for x in acc_rows if x["acc"] in has_child), key=lambda x: -len(x["acc"])):
        shifted = [vec(r, s) for s in (0, -1, 1, -2, 2)]   # lignes de cumul parfois décalées
        if not any(any(x) for x in shifted):
            continue
        kids = [v for q, v in leaves if q["acc"].startswith(r["acc"])]
        tot = [sum(col) for col in zip(*kids)] if kids else None
        if tot and any(same(x, tot) for x in shifted):
            cumuls += 1
        elif any(vec(r)):                                   # parent portant son propre solde
            parents_reels.append(r["acc"])
            leaves.append((r, vec(r)))
        # sinon : montants hors colonnes (date d'édition lue comme compte…) -> ignorée

    # 3) gabarit : score d'invariants
    vs = [(r["acc"], v) for r, v in leaves]
    cands = []
    for name, lay in LAYOUTS.items():
        if len(lay) > k or (gabarit and name != gabarit):
            continue
        for off in range(k - len(lay) + 1):
            idx = {role: off + p for p, role in enumerate(lay)}
            tests = invariants(lay, idx, vs)
            if not tests:
                continue
            sc = sum(tests.values()) / len(tests) - 0.15 * (k - len(lay)) / k
            sc += hint_bonus(idx, cols, header_text)
            cands.append((sc, name, idx, tests))
    if not cands:
        return {"statut": "illisible", "raison": f"aucun gabarit applicable à {k} colonne(s)", "soldes": {}}
    cands.sort(key=lambda c: -c[0])
    sc, name, idx, tests = cands[0]

    def solde_cols(ix):
        return tuple(sorted((r, c) for r, c in ix.items() if r[0] == "S"))

    rivaux, a_confirmer = [], []
    for sc2, n2, idx2, _ in cands[1:]:
        if sc2 < sc - 0.02:
            break
        (rivaux if solde_cols(idx2) != solde_cols(idx) else a_confirmer).append(n2)
    if gabarit:                     # choix imposé par l'utilisateur : plus d'ambiguïté à lever
        rivaux = []

    soldes, n1, ouv, libelles, mvts = {}, {}, {}, {}, {}
    for (r, v) in leaves:
        a = r["acc"]
        soldes[a] = soldes.get(a, 0.0) + _solde(v, idx)
        if r["lab"]:
            libelles[a] = r["lab"]
        if "MD" in idx:                       # mouvements débit / crédit de l'exercice (note 3, TFT)
            md, mk = mvts.get(a, (0.0, 0.0))
            mvts[a] = (md + v[idx["MD"]], mk + v[idx["MK"]])
        if _net(v, idx, "P") is not None:
            n1[a] = n1.get(a, 0.0) + _net(v, idx, "P")
        if _net(v, idx, "O") is not None:
            ouv[a] = ouv.get(a, 0.0) + _net(v, idx, "O")
    return {
        "statut": "ambigu" if rivaux else "ok",
        "gabarit": name, "confiance": round(min(max(sc, 0), 1.0), 3), "invariants": tests,
        "colonnes_numeriques": k, "separateur_decimal": dec,
        "feuilles": len(vs), "cumuls_prouves": cumuls, "parents_avec_solde_propre": parents_reels,
        "rivaux": rivaux, "a_confirmer": sorted(set(a_confirmer) - set(rivaux)),
        "soldes": soldes, "soldes_n1": n1 or None, "soldes_ouverture": ouv or None, "libelles": libelles,
        "mouvements": mvts or None,
        "equilibre": round(sum(soldes.values()), 2),
    }


def invariants(lay, idx, vs):
    t = {}
    fam = {r[0] for r in lay}
    n = len(vs)
    net = lambda v, f: _net(v, idx, f)
    for f in fam:
        tot = sum(net(v, f) for _, v in vs)
        scale = sum(abs(net(v, f)) for _, v in vs) or 1.0
        t["equilibre_" + f] = 1.0 if abs(tot) <= max(1.0, 1e-9 * scale) else 0.0
    for f in ("S", "P"):
        if f + "D" in idx:
            t["exclusivite_" + f] = sum(1 for _, v in vs if not (v[idx[f + "D"]] and v[idx[f + "K"]])) / n
    if "MD" in idx:
        both = sum(1 for _, v in vs if v[idx["MD"]] and v[idx["MK"]]) / n
        t["mouvements_non_exclusifs"] = 1.0 if both > 0.02 else 0.3
    gestion = [v for a, v in vs if a[0] in "67"]
    if gestion and "S" in fam:
        t["solde_gestion_active"] = 1.0 if sum(1 for v in gestion if abs(net(v, "S")) >= 1) > 0.3 * len(gestion) else 0.3
    if gestion and "O" in fam:
        t["ouverture_gestion_nulle"] = sum(1 for v in gestion if abs(net(v, "O")) < 1) / len(gestion)
    if gestion and "P" in fam:
        t["precedent_gestion_active"] = 1.0 if sum(1 for v in gestion if abs(net(v, "P")) >= 1) > 0.3 * len(gestion) else 0.0
    if {"O", "M", "S"} <= fam:
        t["continuite"] = sum(1 for _, v in vs if abs(net(v, "O") + net(v, "M") - net(v, "S")) < 1) / n
    if {"O", "M", "C"} <= fam:
        t["cumul"] = sum(1 for _, v in vs if abs(v[idx["OD"]] + v[idx["MD"]] - v[idx["CD"]]) < 1
                         and abs(v[idx["OK"]] + v[idx["MK"]] - v[idx["CK"]]) < 1) / n
    if {"C", "S"} <= fam:
        t["cumul_solde"] = sum(1 for _, v in vs if abs(net(v, "C") - net(v, "S")) < 1) / n
    return t


def hint_bonus(idx, cols, header_text):
    """Au plus +0.05 : un en-tête départage deux lectures, il ne suffit jamais seul."""
    hits = 0
    for role, p in idx.items():
        words = "".join(header_text.get(cols[p], []))
        if words and any(h in words for h in HINTS.get(role[0], [])):
            hits += 1
    return 0.05 * hits / max(1, len(idx))


def rapport(res):
    """Résumé lisible (stderr des scripts) : ce qui a été compris, et ce qui reste à trancher."""
    if res["statut"] == "illisible":
        return f"[balance] ILLISIBLE : {res.get('raison')}"
    lignes = [f"[balance] feuille « {res.get('feuille', '?')} » — {res['gabarit']} — confiance {res['confiance']:.2f}"
              f" — {res['feuilles']} comptes, {res['cumuls_prouves']} lignes de cumul prouvées et écartées"
              f" — équilibre {res['equilibre']:,.2f}"]
    if res["parents_avec_solde_propre"]:
        lignes.append(f"[balance] comptes parents portant leur propre solde : {res['parents_avec_solde_propre'][:10]}")
    if res["a_confirmer"]:
        lignes.append(f"[balance] lecture des colonnes secondaires à confirmer (soldes N inchangés) : {res['a_confirmer']}")
    if res["statut"] == "ambigu":
        lignes.append(f"[balance] AMBIGU : ces lectures donnent des soldes N différents : {[res['gabarit']] + res['rivaux']}"
                      " -> demander à l'utilisateur, puis relancer avec --gabarit \"<nom>\".")
    if abs(res["equilibre"]) >= 1:
        lignes.append("[balance] ATTENTION : la balance n'est pas équilibrée (Σ soldes ≠ 0).")
    return "\n".join(lignes)
