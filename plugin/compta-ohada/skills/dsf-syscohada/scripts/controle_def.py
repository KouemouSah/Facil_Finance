# -*- coding: utf-8 -*-
"""RAPPORT DE CONTRÔLE de la DEF (interne, jamais déposé) — A4, par sections et catégories, lecture rapide.

Structure :
  Couverture (verdict, compteurs) · 1. Synthèse et points prioritaires · 2. Cohérence comptable ·
  3. Fiscalité (IS, CMF, réintégrations, interprétations, questions) · 4. Tiers (TF5) ·
  5. Présentation de la DEF (comptes hors nomenclature, lignes sans REF, cellules à saisir) ·
  6. Avertissements classés · 7. Sources utilisées et méthode.
Chaque point porte un STATUT écrit (OK / À VOIR / ÉCART) — jamais la couleur seule — et, quand il existe, son IMPACT
chiffré. Le rapport dit aussi ce qui N'EST PAS intégré à la DEF (propositions en attente, points d'expert).
"""
import datetime
import os
import tempfile

BLEU, BLEU_CLAIR, GRIS = "1F3864", "D9E1F2", "F2F2F2"
STATUTS = {"OK": ("E2EFDA", "375623"), "À VOIR": ("FFF2CC", "7F6000"), "ÉCART": ("F8CBAD", "833C0B"), "INFO": ("DDEBF7", "1F3864")}
NUM = '#,##0;(#,##0);"–"'


# contrôles dont l'échec rend la DEF fausse (bloquants) ; les autres sont des anomalies à examiner
BLOQUANTS = ("balance équilibrée", "actif = total passif", "résultat du bilan", "impôt du compte de résultat",
             "tft réconcilié", "bruts de clôture", "ouverture + mouvements", "dotations =")


def fr(n):
    """Montant au format français : espaces pour les milliers."""
    return format(n, ",.0f").replace(",", " ")


def _cat_avert(w):
    w2 = w.lower()
    for mots, cat in ((("cmf", "cuota", "acompte", "is ", "impôt", "réintégr", "déficit", "interprétation", "don"), "Fiscalité — IS et CMF"),
                      (("iva", "tva"), "Fiscalité — IVA (TF3)"), (("irpf", "salaire"), "Fiscalité — IRPF (TF4)"),
                      (("tiers", "tf5", "nif", "bénéficiaire"), "Tiers (TF5)"),
                      (("balance", "compte", "clôtur", "13", "mapping", "non mappé", "éclatement"), "Comptabilité et balance")):
        if any(m in w2 for m in mots):
            return cat
    return "Divers"


def construire(gen, sortie_xlsx, sortie_pdf=None):
    """gen : Generateur (def_modele) déjà rendu. Écrit le classeur de contrôle et, si demandé, son PDF A4."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.page import PageMargins

    D, a = gen.D, gen.a
    f = D["fisc"]
    t1, t2 = f["tf1"], f["tf2"]
    ex = a.exercice
    I = D.get("identite") or {}
    nom = I.get("raison_sociale") or os.path.basename(a.balance)
    fill = lambda c: PatternFill("solid", fgColor=c)
    wb = Workbook()
    feuilles = []

    def page(titre, largeurs):
        ws = wb.active if not feuilles else wb.create_sheet()
        ws.title = titre[:31]
        feuilles.append(ws.title)
        ws.append([f"{nom} — DEF {ex}"]); ws.cell(1, 1).font = Font(name="Calibri", size=10, bold=True, color=BLEU)
        ws.append(["Rapport de contrôle interne — ne pas joindre au dépôt"]); ws.cell(2, 1).font = Font(name="Calibri", size=9, italic=True, color="7F7F7F")
        ws.append([titre]); ws.cell(3, 1).font = Font(name="Calibri", size=15, bold=True, color=BLEU)
        ws.append([])
        for j, w in enumerate(largeurs, 1):
            ws.column_dimensions[ws.cell(1, j).column_letter].width = w
        ws.page_setup.paperSize, ws.page_setup.orientation = ws.PAPERSIZE_A4, "portrait"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.6, bottom=0.6, header=0.25, footer=0.25)
        ws.print_options.horizontalCentered = True
        ws.oddFooter.left.text = f"Contrôle DEF {ex} — interne"
        ws.oddFooter.center.text = "Page &P / &N"
        ws.oddFooter.right.text = titre[:50]
        for p in (ws.oddFooter.left, ws.oddFooter.center, ws.oddFooter.right):
            p.size, p.color = 8, "7F7F7F"
        ws.sheet_view.showGridLines = False
        return ws

    def section(ws, texte):
        ws.append([])
        ws.append([texte])
        c = ws.cell(ws.max_row, 1)
        c.font = Font(name="Calibri", size=11, bold=True, color=BLEU)
        for j in range(1, 5):
            ws.cell(ws.max_row, j).border = Border(bottom=Side(style="medium", color=BLEU))

    def entetes(ws, cols):
        ws.append(cols)
        for j in range(1, len(cols) + 1):
            c = ws.cell(ws.max_row, j)
            c.font, c.fill = Font(name="Calibri", size=9, bold=True, color="FFFFFF"), fill(BLEU)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def ligne(ws, vals, statut=None, gras=False):
        ws.append(vals)
        r = ws.max_row
        for j in range(1, len(vals) + 1):
            c = ws.cell(r, j)
            c.font = Font(name="Calibri", size=9, bold=gras)
            c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="right" if isinstance(c.value, (int, float)) else None)
            c.border = Border(bottom=Side(style="hair", color="BFBFBF"))
            if isinstance(c.value, (int, float)):
                c.number_format = NUM
        if statut:
            j = vals.index(statut) + 1
            fond, encre = STATUTS[statut]
            c = ws.cell(r, j)
            c.fill, c.font = fill(fond), Font(name="Calibri", size=9, bold=True, color=encre)
            c.alignment = Alignment(horizontal="center", vertical="top")

    # ---------------- données
    coherence = [(x["controle"], "OK" if x["ok"] else ("ÉCART" if any(b in x["controle"].lower() for b in BLOQUANTS) else "À VOIR"), x["detail"]) for x in D["controles"]]
    tiers = [(n, "OK" if ok else "À VOIR", d) for n, ok, d in gen.controles if n.startswith("TF5")]
    presentation = [(n, "OK" if ok else "À VOIR", d) for n, ok, d in gen.controles if not n.startswith("TF5")]
    props = f.get("propositions", [])
    taux = t2[4]
    is_sup_cmf = t2[5] >= t2[14]
    en_attente = [p for p in props if p.get("statut") != "acceptée" and p["type"] != "information"]
    acceptees = [p for p in props if p.get("statut") == "acceptée"]
    impact_attente = round(sum(p["montant"] for p in en_attente if p["type"] == "reintegration") * taux) if is_sup_cmf else 0
    interp = [w for w in D["avertissements"] if "INTERPRÉTATION D'ENTREPRISE" in w]
    questions = f.get("questions", [])
    n_ecart = sum(1 for _, s, _ in coherence if s == "ÉCART")
    n_voir = sum(1 for _, s, _ in coherence + tiers + presentation if s == "À VOIR") + len(en_attente) + len(questions) + len(interp)
    prioritaires = []
    for p in en_attente:
        prioritaires.append(("Fiscalité", f"Réintégration proposée NON intégrée : {p['compte']} — {fr(p['montant'])} ({p['article']})",
                             f"IS +{fr(round(p['montant'] * taux))}" if p["type"] == "reintegration" and is_sup_cmf else "à examiner"))
    for n, s, d in coherence:
        if s != "OK":
            prioritaires.append(("Cohérence comptable" + (" — BLOQUANT" if s == "ÉCART" else ""), n, d[:200]))
    if f.get("versements_cmf") == [] and t2[14]:
        prioritaires.append(("Fiscalité", "Aucune CMF ni acompte d'IS versé selon les fichiers", f"IS à payer {fr(t2[7])}"))
    for q in questions:
        prioritaires.append(("Fiscalité", f"Question d'entrée ouverte ({q['article']})", q["question"][:260]))
    nif = next((d for n, ok, d in tiers if "bénéficiaires identifiés" in n), "")
    if nif and "NIF manquants : aucun" not in nif:
        prioritaires.append(("Tiers", "NIF manquants au TF5 (art. 27.2)", nif.split("NIF manquants : ")[-1][:120]))

    # ---------------- couverture
    ws = page("Rapport de contrôle", [4, 90])
    for _ in range(4):
        ws.append([])
    for texte, fnt in (("RAPPORT DE CONTRÔLE", Font(name="Calibri", size=24, bold=True, color=BLEU)),
                       (f"Déclaration statistique et fiscale — exercice {ex}", Font(name="Calibri", size=13, color=BLEU)),
                       ("", None), (nom, Font(name="Calibri", size=16, bold=True)), ("", None)):
        ws.append(["", texte])
        if fnt:
            ws.cell(ws.max_row, 2).font = fnt
        ws.cell(ws.max_row, 2).alignment = Alignment(horizontal="center")
    verdict = ("ÉCART", "Des contrôles de cohérence sont en écart : la DEF ne doit pas être déposée en l'état.") if n_ecart else \
              ("À VOIR", f"Les états sont cohérents ; {n_voir} point(s) demandent une décision ou une vérification avant le dépôt.") if n_voir else \
              ("OK", "Aucun point ouvert.")
    ws.append(["", f"Verdict : {verdict[0]}"])
    ws.cell(ws.max_row, 2).font = Font(name="Calibri", size=14, bold=True, color=STATUTS[verdict[0]][1])
    ws.cell(ws.max_row, 2).fill = fill(STATUTS[verdict[0]][0])
    ws.cell(ws.max_row, 2).alignment = Alignment(horizontal="center")
    ws.append(["", verdict[1]]); ws.cell(ws.max_row, 2).alignment = Alignment(horizontal="center", wrap_text=True)
    ws.append([])
    for lab, v in (("Contrôles de cohérence", f"{len(coherence) - n_ecart} OK / {n_ecart} en écart"),
                   ("Réintégrations acceptées", f"{len(acceptees)} ({fr(sum(p['montant'] for p in acceptees))})"),
                   ("Réintégrations proposées NON intégrées", f"{len(en_attente)} — impact IS {fr(impact_attente)}"),
                   ("Résultat fiscal / IS dû / IS à payer", f"{fr(t1[22])} / {fr(f['impot_du'])} / {fr(t2[7])}"),
                   ("Cellules à saisir par l'entité", str(len(gen.a_completer)))):
        ws.append(["", f"{lab} :  {v}"]); ws.cell(ws.max_row, 2).alignment = Alignment(horizontal="center")
        ws.cell(ws.max_row, 2).font = Font(name="Calibri", size=10)
    tr = getattr(a, "tracabilite", None) or {}
    if tr.get("valide_par"):
        ws.append(["", f"Version définitive — validée par {tr['valide_par']} le "
                       f"{str(tr.get('valide_le') or '').replace('T', ' ')} (réponses appliquées : section 7)"])
        ws.cell(ws.max_row, 2).alignment = Alignment(horizontal="center")
        ws.cell(ws.max_row, 2).font = Font(name="Calibri", size=10, bold=True)
    ws.append([]); ws.append(["", f"Généré le {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')} — document interne, à relire avec le comptable."])
    ws.cell(ws.max_row, 2).font = Font(name="Calibri", size=9, italic=True, color="7F7F7F")
    ws.cell(ws.max_row, 2).alignment = Alignment(horizontal="center")

    # ---------------- 1. synthèse
    ws = page("1. Synthèse", [24, 70, 30])
    section(ws, "Points à traiter en priorité")
    entetes(ws, ["Catégorie", "Point", "Impact / détail"])
    for cat, pt, imp in prioritaires or [("—", "Aucun point prioritaire", "")]:
        ligne(ws, [cat, pt, imp])
    section(ws, "Tableau de bord par catégorie")
    entetes(ws, ["Catégorie", "Contenu", "OK / À VOIR / ÉCART"])
    for cat, items in (("Cohérence comptable", coherence), ("Tiers (TF5)", tiers), ("Présentation de la DEF", presentation)):
        ok_ = sum(1 for _, s, _ in items if s == "OK")
        ligne(ws, [cat, f"{len(items)} contrôle(s)", f"{ok_} / {sum(1 for _, s, _ in items if s == 'À VOIR')} / {sum(1 for _, s, _ in items if s == 'ÉCART')}"])
    ligne(ws, ["Fiscalité", "propositions, questions, interprétations", f"{len(acceptees)} acceptée(s) / {len(en_attente) + len(questions) + len(interp)} à décider / 0"])

    # ---------------- 2. cohérence comptable
    ws = page("2. Cohérence comptable", [52, 12, 60])
    section(ws, "Contrôles arithmétiques et d'articulation des états")
    entetes(ws, ["Contrôle", "Statut", "Détail"])
    for n, s, d in coherence:
        ligne(ws, [n, s, d], statut=s)
    section(ws, "Écriture d'impôt")
    entetes(ws, ["Point", "Statut", "Détail"])
    ligne(ws, ["Impôt sur le résultat dans la balance", "INFO", D["info_is"]["message"][:300]], statut="INFO")

    # ---------------- 3. fiscalité
    ws = page("3. Fiscalité", [46, 14, 20, 44])
    section(ws, "Chiffres clés (TF1 / TF2)")
    entetes(ws, ["Rubrique", "Ligne", "Montant", "Commentaire"])
    for lab, lg, v, com in (("Résultat comptable avant impôt", "TF1-1", t1[1], ""), ("Total des réintégrations", "TF1-13", t1[13], f"{len(acceptees)} acceptée(s)"),
                            ("Résultat fiscal", "TF1-22", t1[22], ""), ("Impôt sur les sociétés", "TF2-5", t2[5], f"taux {taux:.0%}"),
                            ("Cuota Mínima Fiscal", "TF2-14", t2[14], "1,5 % de la base"), ("CMF / acomptes versés (fichiers)", "TF2-6", t2[6], f.get("source_versements", "")),
                            ("IS à payer", "TF2-7", t2[7], "")):
        ligne(ws, [lab, lg, v, com], gras=lg in ("TF1-22", "TF2-7"))
    section(ws, "Réintégrations")
    entetes(ws, ["Compte / origine", "Statut", "Montant", "Motif (article)"])
    for p in props:
        if p["type"] == "information":
            continue
        s = "OK" if p.get("statut") == "acceptée" else "À VOIR"
        ligne(ws, [f"{p['compte']} — {p.get('source', 'balance')}\n{(p.get('libelle') or '')[:120]}",
                   s, p["montant"], f"{p['motif'][:220]} ({p['article']})" + ("" if s == "OK" else " — NON intégrée à la DEF")], statut=s)
    if en_attente:
        ligne(ws, ["Impact si toutes les propositions en attente sont acceptées", "À VOIR", impact_attente,
                   f"IS à payer passerait de {fr(t2[7])} à {fr(t2[7] + impact_attente)}"], statut="À VOIR", gras=True)
    section(ws, "Interprétations de l'entreprise et questions d'entrée")
    entetes(ws, ["Point", "Statut", "", "Détail"])
    for w in interp:
        ligne(ws, ["Interprétation d'entreprise", "À VOIR", "", w[:400]], statut="À VOIR")
    for q in questions:
        ligne(ws, [f"Question ({q['article']})", "À VOIR", "", q["question"][:300]], statut="À VOIR")
    for x in f.get("versements_exclus", []):
        ligne(ws, ["Débit du compte d'IS non retenu comme CMF", "INFO", x["montant"], f"{x['piece']} — {x['motif']}"], statut="INFO")

    # ---------------- 4. tiers
    ws = page("4. Tiers (TF5)", [46, 12, 70])
    section(ws, "Rapprochement des bénéficiaires déclarés avec le grand livre")
    entetes(ws, ["Bénéficiaire", "Statut", "Détail"])
    for n, s, d in tiers:
        ligne(ws, [n.replace("TF5 ", "").replace(" : rapprochement grand livre", ""), s, d], statut=s)

    # ---------------- 5. présentation
    ws = page("5. Présentation de la DEF", [46, 12, 70])
    section(ws, "Placements et structure (traçabilité des choix automatiques)")
    entetes(ws, ["Point", "Statut", "Détail"])
    for n, s, d in presentation:
        ligne(ws, [n, s, d], statut=s)
    section(ws, "Couverture des cellules chiffrées du modèle")
    entetes(ws, ["Nature", "", "Nombre"])
    for k, n in gen.couverture.items():
        ligne(ws, [k, "", n])
    section(ws, "Cellules à saisir par l'entité (fond ambre dans le classeur)")
    entetes(ws, ["Feuille", "", "Cellules"])
    from collections import Counter
    for fe, n in Counter(x for x, _ in gen.a_completer).most_common():
        ligne(ws, [fe, "", n])

    # ---------------- 6. avertissements classés
    ws = page("6. Avertissements", [28, 100])
    par = {}
    for w in D["avertissements"]:
        par.setdefault(_cat_avert(w), []).append(w)
    for cat in sorted(par):
        section(ws, cat)
        for w in par[cat]:
            ligne(ws, ["", w])

    # ---------------- 7. sources et méthode
    ws = page("7. Sources et méthode", [34, 94])
    section(ws, "Fichiers utilisés")
    entetes(ws, ["Rôle", "Fichier et usage"])
    for role, chemin, usage in (("Balance", a.balance, "soldes, mouvements, ouvertures : TOUS les montants de l'exercice"),
                                ("Grand livre", a.grand_livre, "TF3 à TF5, CMF versée, IS passé en charge (structure des pièces)"),
                                (("Formulaire officiel (gabarit embarqué)", a.modele,
                                  "structure seule (feuilles, lignes, libellés, formules des totaux) ; aucune DEF déposée : "
                                  "bilan N-1 depuis l'ouverture de la balance, N-1 des notes vide, TF5 depuis le grand livre")
                                 if getattr(a, "gabarit", False) else
                                 ("DEF de référence (modèle)", a.modele, "structure (feuilles, lignes, libellés), colonne N-1, tiers nommés du TF5 "
                                  + ("et leurs montants déclarés (même exercice)" if a.colonne_precedente == "N-1" else ""))),
                                ("Profil de l'entreprise", a.profil, "interprétations fiscales choisies, faits, identité du rapport"),
                                ("Décisions fiscales", a.decisions, "réintégrations acceptées, déficits, CMF saisie")):
        ligne(ws, [role, f"{os.path.basename(chemin) if chemin else 'non fourni'} — {usage}"])
    if D.get("src_n1"):                                   # d'où viennent les comparatifs N-1 (balance, DEF, scan lu)
        ligne(ws, ["Comparatifs N-1", D["src_n1"]])
    tr = getattr(a, "tracabilite", None)
    if tr:
        # chaque réponse APPLIQUÉE, son origine (formulaire, conversation, reprise d'une source) et sa date ; puis qui
        # a validé et quand — sans quoi une version définitive ne dit pas sur quelles décisions elle repose
        section(ws, "Validation")
        entetes(ws, ["Point", "Valeur"])
        ligne(ws, ["Statut de la version", tr.get("statut") or ""])
        if tr.get("valide_par"):
            ligne(ws, ["Validée par", tr["valide_par"]])
            ligne(ws, ["Validée le", str(tr.get("valide_le") or "").replace("T", " ")])
        section(ws, f"Réponses appliquées ({len(tr.get('reponses') or [])})")
        entetes(ws, ["Question", "Réponse — origine, date"])
        for x in tr.get("reponses") or [{"question": "—", "reponse": "aucune", "origine": "", "le": ""}]:
            quand = f", le {x['le'].replace('T', ' ')}" if x.get("le") else ""
            ligne(ws, [x["question"], f"{x['reponse']} — {x['origine']}{quand}"])
    section(ws, "Méthode")
    for t in ("Aucun montant de l'exercice n'est recopié du modèle : chaque cellule est recalculée depuis la balance "
              "(ou le grand livre pour TF3 à TF5), sauf les montants DÉCLARÉS du TF5 quand le modèle est du même exercice.",
              "Les totaux sont des formules dont la structure est prouvée par les montants du modèle ; les écarts avec la DEF de "
              "référence sont listés et expliqués par la garde de référence.",
              "Une proposition de réintégration n'est JAMAIS appliquée sans acceptation (fichier de décisions) : celles en attente "
              "figurent en section 3 avec leur impact."):
        ligne(ws, ["", t])

    wb.save(sortie_xlsx)
    if sortie_pdf:
        import def_ge
        return def_ge.exporter_pdf(sortie_xlsx, sortie_pdf, feuilles)
    return True
