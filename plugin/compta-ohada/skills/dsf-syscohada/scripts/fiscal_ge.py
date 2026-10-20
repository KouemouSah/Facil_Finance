# -*- coding: utf-8 -*-
"""Moteur fiscal Guinée équatoriale : TF1 (résultat fiscal) et TF2 (IS / Cuota Mínima Fiscal).

SOURCE UNIQUE : compta-ohada/lib/fiscal_ge.py (+ parametres_fiscaux_ge.json). Copies
synchronisées dans les skills par tools/sync_lib.py ; ne jamais éditer une copie.

Principes :
- AUCUN taux dans le code : tout vient de parametres_fiscaux_ge.json, chaque valeur avec son article.
- Les réintégrations tirées de la balance sont des PROPOSITIONS (jugement) : elles ne s'appliquent
  que si l'utilisateur les accepte (fichier de décisions). Rien ne bloque : une donnée absente
  prend une valeur par défaut VISIBLE (ex. CMF versée = 0) et produit un avertissement chiffré.
- Le calcul part du résultat AVANT impôt (l'IS n'est pas déductible, art. 62.3.a) : c'est la
  ligne 1 de TF1 dans la DEF déposée.

Numérotation des lignes : celle des tableaux TF1/TF2 de la DEF GE (modèle GE 2025 du dossier pilote).
"""
import datetime
import json
import math
import os
import re

HERE = os.path.dirname(os.path.abspath(__file__))
PARAMS = os.path.join(HERE, "parametres_fiscaux_ge.json")

# TF1 : lignes de réintégration (II) et de déduction (III) — code -> (n° de ligne, libellé ES, libellé FR)
TF1_REINT = {
    "amort_non_deductibles": (4, "Amortizaciones no deducibles", "Amortissements non déductibles"),
    "amort_differes": (5, "Amortizaciones consideradas diferidas fiscalmente", "Amortissements réputés différés"),
    "provisions": (6, "Provisiones no deducibles", "Provisions non déductibles"),
    "remuneration_exploitant": (7, "Remuneraciones del explotante individual y su cónyuge", "Rémunérations de l'exploitant"),
    "interets_excedentaires": (8, "Intereses excedentarios de las cuentas corrientes", "Intérêts excédentaires des comptes courants"),
    "impuestos": (9, "Impuestos no deducibles", "Impôts non déductibles"),
    "multas": (10, "Multas y penalidades", "Amendes et pénalités"),
    "donaciones": (11, "Donaciones y liberalidades no deducibles", "Dons et libéralités non déductibles"),
    "diversos": (12, "Diversos (a detallar)", "Divers (à détailler)"),
}
TF1_DEDUC = {
    "amort_differes_anterieurs": (15, "Amortizaciones anteriormente diferidas", "Amortissements antérieurement différés"),
    "provisions_anterieures": (16, "Provisiones anteriormente integradas", "Provisions antérieurement réintégrées"),
    "plus_values": (17, "Plus-valías no imponibles", "Plus-values non imposables"),
    "diversos_1": (18, "Diversos I", "Divers I"),
    "diversos_2": (19, "Diversos II", "Divers II"),
}


def load_params(path=PARAMS):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def P(params, cle):
    return params["parametres"][cle]["valeur"]


def cite(params, cle):
    return params["parametres"][cle]["article"]


def entrees_depuis_soldes(soldes, exercice, date_cloture=None):
    """Résultat avant IS, CA et autres produits depuis les soldes D-C d'une balance."""
    res_avant_is = -sum(s for a, s in soldes.items() if a[:1] in "678" and not a.startswith("89"))
    ca = -sum(s for a, s in soldes.items() if a.startswith("70"))
    autres = -sum(s for a, s in soldes.items() if a[:1] in "78" and not a.startswith("70") and s < 0)
    return {"exercice": exercice, "date_cloture": date_cloture or f"{exercice}-12-31",
            "resultat_avant_is": round(res_avant_is, 2), "ca": round(ca, 2), "autres_produits": round(autres, 2)}


def _norm(s):
    import unicodedata
    return unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()


def propositions(soldes, params, libelles=None):
    """Réintégrations proposées depuis la balance : jamais appliquées sans acceptation.
    Le numéro de compte ne suffit pas : le LIBELLÉ doit confirmer la nature (mots_cles),
    sinon la proposition devient une question — libellé réel à l'appui."""
    libelles = libelles or {}
    out = []
    for regle in params["propositions"]:
        for a, s in sorted(soldes.items()):
            if not (any(a.startswith(p) for p in regle["prefixes"]) and s > 0.5):
                continue
            lib = libelles.get(a, "")
            typ, motif = regle["type"], regle["motif"]
            if typ == "reintegration" and regle.get("mots_cles"):
                if not lib:
                    typ, motif = "question", "Libellé du compte inconnu : nature à confirmer. " + motif
                # mot-clé en DÉBUT de mot : « don » ne doit pas matcher « abandon »
                elif not any(re.search(r"\b" + re.escape(k), _norm(lib)) for k in regle["mots_cles"]):
                    typ = "question"
                    motif = (f"Le libellé « {lib} » ne correspond pas à la nomenclature ({'/'.join(regle['prefixes'])} = "
                             f"{regle['ligne']}) : la nature réelle des dépenses est à examiner. " + motif)
            out.append({"compte": a, "libelle": lib, "montant": round(s, 2), "ligne": regle["ligne"], "type": typ,
                        "article": regle["article"], "motif": motif})
    return out


def resoudre_profil(params, profil, decisions):
    """Interprétations (défaut = lecture stricte) et faits (None = inconnu) ; les décisions de
    l'exercice priment sur le profil d'entreprise, qui prime sur les défauts."""
    profil = profil or {}
    pi, pf = profil.get("interpretations", {}), profil.get("faits", {})
    interp = {}
    for k, v in params.get("interpretations", {}).items():
        interp[k] = decisions.get(k, pi.get(k, v["defaut"]))
    faits = {k: decisions.get(k, pf.get(k)) for k in params.get("faits", {})}
    justif = profil.get("justifications", {})
    return interp, faits, justif


def questions_entree(params, faits, soldes, libelles, portee="tf2"):
    """Faits inconnus -> questions, avec les indices trouvés dans la balance (libellés des comptes).
    Seules les questions de la PORTÉE demandée sont posées (tf2 par défaut : IS / CMF)."""
    out = []
    for k, spec in params.get("faits", {}).items():
        if faits.get(k) is not None or spec.get("portee", "tf2") != portee:
            continue
        indices = []
        for a, s in sorted((soldes or {}).items()):
            lib = _norm((libelles or {}).get(a, ""))
            if a[:2] in ("40", "41", "46", "47") and abs(s) >= 1 and any(re.search(r"\b" + re.escape(m), lib)
                                                                          for m in spec.get("indices_libelles", [])):
                indices.append({"compte": a, "libelle": (libelles or {}).get(a, ""), "solde": round(s, 2)})
        out.append({"fait": k, "article": spec["article"], "question": spec["question"], "indices": indices})
    return out


def propositions_gl(gl_lignes, params, pieces_exclues=()):
    """Réintégrations proposées depuis le GRAND LIVRE, quand seul le libellé de l'écriture révèle la nature de la charge
    (ex. principal d'IS d'un contrôle fiscal passé au 6480 « autres impôts et taxes »). Regroupées par compte ; jamais
    appliquées sans acceptation (même règle que les propositions de la balance)."""
    out = []
    for regle in params.get("propositions_gl", []):
        rx = re.compile(regle["libelle_regex"])
        excl = [_norm(e) for e in regle.get("exclure", [])]
        par_compte = {}
        for x in gl_lignes or []:
            if x.get("piece") in pieces_exclues or not x["compte"].startswith(tuple(regle["prefixes"])):
                continue
            lib = _norm(x.get("libelle") or "")
            if not rx.search(lib) or any(e in lib for e in excl):
                continue
            m = round((x.get("debit") or 0) - (x.get("credit") or 0), 2)
            if abs(m) < 1:
                continue
            p = par_compte.setdefault(x["compte"], {"montant": 0.0, "pieces": []})
            p["montant"] += m
            p["pieces"].append(f"{x.get('piece')} « {x.get('libelle')} » {m:,.0f}")
        for compte, p in sorted(par_compte.items()):
            out.append({"compte": compte, "libelle": "; ".join(p["pieces"][:4]), "montant": round(p["montant"], 2),
                        "ligne": regle["ligne"], "type": regle["type"], "article": regle["article"],
                        "motif": regle["motif"] + " Source : grand livre.", "source": "grand livre"})
    out += is_en_charge_par_piece(gl_lignes, params, pieces_exclues, deja={p["compte"] for p in out})
    return out


def is_en_charge_par_piece(gl_lignes, params, pieces_exclues=(), deja=()):
    """IS passé EN CHARGE, reconnu par la STRUCTURE de la pièce, pas par son libellé (raisonnement de réviseur) :
    une pièce qui crédite le compte d'IS (441) sans débiter l'impôt sur le résultat (89) a pour contrepartie des
    charges. Principal d'IS en charge = crédit 441 de la pièce − majorations d'IS de la même pièce (647 dont le
    libellé cite l'IS), plafonné aux charges non pénales débitées par la pièce. Exemple du dossier pilote :
    469 884 995 − 197 545 272 = 272 339 723, passés au 6480 « autres impôts et taxes »."""
    c = params["comptes"]
    racines_is = tuple(c.get("impot_benefices", ["441"]))
    rx_is = re.compile(r"\biss?\b|impuesto sobre (las )?sociedades|impot sur (les |le )?(benefices|resultat|societes)")
    pieces = {}
    for x in gl_lignes or []:
        if x.get("piece") and x.get("piece") not in pieces_exclues:
            pieces.setdefault(x["piece"], []).append(x)
    out = {}
    for piece, ls in pieces.items():
        credit_is = sum(x.get("credit") or 0 for x in ls if x["compte"].startswith(racines_is))
        if credit_is < 1 or any(x["compte"].startswith("89") and (x.get("debit") or 0) for x in ls):
            continue
        maj = sum(x.get("debit") or 0 for x in ls if x["compte"].startswith("647") and rx_is.search(_norm(x.get("libelle") or "")))
        charges = [x for x in ls if x["compte"][:2] in ("64", "65", "67", "81", "83", "84") and not x["compte"].startswith("647")
                   and (x.get("debit") or 0)]
        principal = round(min(credit_is - maj, sum(x["debit"] for x in charges)), 2)
        if principal < 1 or not charges:
            continue
        compte = max({x["compte"] for x in charges}, key=lambda a: sum(x["debit"] for x in charges if x["compte"] == a))
        if compte in deja:
            continue
        lib_is = next((x.get("libelle") for x in ls if x["compte"].startswith(racines_is) and x.get("credit")), "")
        o = out.setdefault(compte, {"montant": 0.0, "pieces": []})
        o["montant"] += principal
        o["pieces"].append(f"{piece} (« {lib_is} ») : crédit IS {credit_is:,.0f} − majorations d'IS {maj:,.0f} = {principal:,.0f}")
    return [{"compte": a, "libelle": "; ".join(o["pieces"][:4]), "montant": round(o["montant"], 2), "ligne": "impuestos",
             "type": "reintegration", "article": "art. 62.3.a", "source": "grand livre (structure des pièces)",
             "motif": "IS (principal) comptabilisé en charge, reconnu par la structure de la pièce : l'IS n'est pas déductible, "
                      "y compris celui d'un exercice antérieur (rappel, contrôle fiscal). À confirmer avec le comptable."}
            for a, o in sorted(out.items())]


def versements_depuis_fichiers(exercice, params, mouvements=None, libelles=None, gl_lignes=None, pieces_exclues=()):
    """CMF / acomptes d'IS VERSÉS, lus dans les fichiers fournis (règle du porteur : ce que la comptabilité enregistre
    comme versé est versé ; la quittance n'est pas exigée pour calculer, elle se joint au dépôt, art. 80.5).
    Un versement = un DÉBIT d'un compte d'impôt sur les bénéfices (441). Qualification :
      - libellé d'un exercice antérieur (contrôle fiscal, redressement, régularisation, année < N) -> exclu ;
      - versement de janvier N sans mention de l'exercice N -> 2e acompte de N-1 (art. 80.2 : 15 janvier N) -> exclu ;
      - sinon -> acompte / CMF de l'exercice N, déductible en TF2 ligne 6.
    Sans grand livre, les débits de la balance (441) sont retenus, sans date ni libellé : signalé.
    Retourne (versements retenus, exclus, source)."""
    c = params["comptes"]
    racines = tuple(c.get("impot_benefices", ["441"]))
    mots_cmf = [_norm(m) for m in c.get("cmf_mots", [])]
    mots_ant = [_norm(m) for m in c.get("is_anterieur_mots", [])]
    n = int(exercice)
    retenus, exclus = [], []
    if gl_lignes:
        for x in gl_lignes:
            if not x["compte"].startswith(racines) or not x.get("debit") or x.get("piece") in pieces_exclues:
                continue
            lib = _norm(x.get("libelle") or "")
            annees = [int(y) for y in re.findall(r"(?<!\d)(20\d{2})(?!\d)", lib)] + \
                     [2000 + int(y) for y in re.findall(r"\b(?:is|iss|cmf)\s?(\d{2})\b", lib)]
            date = str(x.get("date") or "")
            v = {"date": date, "piece": x.get("piece", ""), "libelle": x.get("libelle", ""), "montant": round(x["debit"], 2),
                 "compte": x["compte"]}
            if any(m in lib for m in mots_ant) or any(a < n for a in annees):
                exclus.append({**v, "motif": "exercice antérieur (libellé)"})
            elif date[5:7] == "01" and n not in annees and not any(m in lib for m in mots_cmf):
                exclus.append({**v, "motif": f"janvier {n} : 2e acompte de {n - 1} (art. 80.2)"})
            else:
                retenus.append({**v, "qualifie": any(m in lib for m in mots_cmf) or n in annees})
        return retenus, exclus, "grand livre"
    for a, (md, _mk) in sorted((mouvements or {}).items()):
        if a.startswith(racines) and md:
            retenus.append({"date": "", "piece": "", "libelle": (libelles or {}).get(a, ""), "montant": round(md, 2),
                            "compte": a, "qualifie": False})
    return retenus, exclus, "balance (débits de l'exercice, sans date)"


def calculer(entrees, decisions=None, soldes=None, params=None, libelles=None, profil=None,
             mouvements=None, gl_lignes=None, pieces_exclues=()):
    params = params or load_params()
    d = decisions or {}
    avert = []
    interp, faits, justif = resoudre_profil(params, profil, d)
    questions = questions_entree(params, faits, soldes, libelles)
    props = propositions(soldes, params, libelles) if soldes else []
    props += propositions_gl(gl_lignes, params, pieces_exclues)

    # --- II. réintégrations : propositions acceptées + saisies de l'utilisateur
    acc = d.get("propositions_acceptees", [])
    reint = {k: 0.0 for k in TF1_REINT}
    detail = []
    for p in props:
        if p["type"] == "information":
            continue
        pris = acc == "toutes" and p["type"] == "reintegration" or p["compte"] in (acc if isinstance(acc, list) else [])
        p["statut"] = "acceptée" if pris else ("en attente de réponse" if p["type"] == "question" else "proposée, non retenue")
        if pris:
            reint[p["ligne"]] += p["montant"]
            detail.append({**p, "origine": "proposition"})
    for x in d.get("reintegrations", []):
        reint[x["ligne"]] += x["montant"]
        detail.append({**x, "origine": "saisie"})
    deduc = {k: 0.0 for k in TF1_DEDUC}
    for x in d.get("deductions", []):
        deduc[x["ligne"]] += x["montant"]
        detail.append({**x, "origine": "saisie", "sens": "déduction"})
    attente = [p for p in props if p.get("statut") == "en attente de réponse"]
    if attente:
        avert.append("Question(s) en attente : " + "; ".join(f"{p['compte']} {p['montant']:,.0f} ({p['article']})" for p in attente)
                     + " — non réintégré(s) tant que l'utilisateur n'a pas répondu.")

    r1 = entrees["resultat_avant_is"]
    tot_reint = round(sum(reint.values()), 2)
    tot_deduc = round(sum(deduc.values()), 2)
    rf = round(r1 + tot_reint - tot_deduc, 2)                     # IV. résultat net fiscal (ligne 22)

    # --- V. déficits antérieurs (art. 73) : imputation du plus ancien au plus récent, dans le délai
    n = int(entrees["exercice"])
    dur = P(params, "deficit_report_exercices")
    deficits = sorted(d.get("deficits", []), key=lambda x: x["exercice"])
    imputable, perimes = [], []
    for x in deficits:
        (imputable if n - int(x["exercice"]) <= dur else perimes).append(x)
    if perimes:
        avert.append(f"Déficits hors délai de {dur} exercices ({cite(params, 'deficit_report_exercices')}), non imputés : "
                     + ", ".join(f"{x['exercice']} ({x['montant']:,.0f})" for x in perimes))
    if deficits:
        avert.append("Déficits imputés : ils doivent avoir été reconnus par l'Administration fiscale (art. 73.1).")
    impute, reste = 0.0, []
    for x in imputable:
        u = min(x["montant"], max(rf - impute, 0))
        impute += u
        if x["montant"] - u > 0.5:
            reste.append({"exercice": x["exercice"], "montant": round(x["montant"] - u, 2),
                          "imputable_jusqu_en": int(x["exercice"]) + dur})
    if rf < 0:
        reste.append({"exercice": n, "montant": round(-rf, 2), "imputable_jusqu_en": n + dur})
    rn = round(rf - impute, 2)                                    # VI. résultat net (ligne 26)

    # --- TF2 I. impôt sur les sociétés
    exonere = round(d.get("benefice_exonere", 0.0), 2)
    imposable = max(rn - exonere, 0.0)
    pas = P(params, "is_arrondi_base")

    def impot(arr_base, arr_franc):
        b = math.floor(imposable / pas) * pas if arr_base else imposable   # art. 75.1 : fraction < 1 000 négligeable
        x = b * P(params, "is_taux")
        return b, (round(x) if arr_franc else round(x, 2))

    base, is_brut = impot(interp["arrondi_base_is"], interp["arrondi_is_franc"])
    _, is_strict = impot(True, True)
    if abs(is_brut - is_strict) >= 0.01:
        ecarts = [k for k in ("arrondi_base_is", "arrondi_is_franc") if interp[k] != params["interpretations"][k]["defaut"]]
        avert.append(f"INTERPRÉTATION D'ENTREPRISE ({', '.join(ecarts)}) : IS {is_brut:,.2f} au lieu de {is_strict:,.2f} "
                     f"selon la lecture stricte ({cite(params, 'is_arrondi_base')}) — écart {is_brut - is_strict:+,.2f}."
                     + (f" Justification : {justif.get(ecarts[0])}" if ecarts and justif.get(ecarts[0]) else
                        " Aucune justification consignée."))
    nouvelle = bool(faits.get("nouvelle_societe") or False)
    if nouvelle and d.get("bonification_nouvelle", True):
        is_brut = round(is_brut * (1 - P(params, "bonification_nouvelles")))
        avert.append(f"Bonification de {P(params, 'bonification_nouvelles'):.0%} ({cite(params, 'bonification_nouvelles')}) "
                     "appliquée : vérifier que l'exercice est parmi les cinq premiers.")

    # --- TF2 II. Cuota Mínima Fiscal
    base_cmf_choix = interp["base_cmf"]
    ca = entrees["ca"]
    diversos = entrees.get("autres_produits", 0.0) if base_cmf_choix == "produits_totaux" else 0.0
    base_cmf = round(ca + diversos, 2)
    arr = (lambda x: round(x)) if interp["arrondi_is_franc"] else (lambda x: round(x, 2))
    cmf = 0 if nouvelle else arr(base_cmf * P(params, "cmf_taux"))
    if not nouvelle:
        alt = round((ca + entrees.get("autres_produits", 0.0)) * P(params, "cmf_taux"))
        if base_cmf_choix == "ca" and abs(alt - cmf) >= 1:
            avert.append(f"Base de la CMF = chiffre d'affaires (comptes 70), comme la DEF déposée. La loi parle de « volumen de "
                         f"ingresos » (art. 75.3) / « ingresos globales » (art. 80.2) : avec tous les produits, la CMF serait "
                         f"{alt:,.0f} au lieu de {cmf:,.0f}. À trancher par un fiscaliste (sans effet si l'IS dépasse la CMF).")
    else:
        avert.append(f"Société nouvelle : Cuota Mínima non applicable ({cite(params, 'cmf_exoneration_nouvelles')}).")
    du = max(is_brut, cmf)

    versee = d.get("cmf_versee")                      # saisie explicite de l'utilisateur : elle prime
    versements_exclus, source_versements = [], "saisie de l'utilisateur"
    if versee is None and (mouvements is not None or gl_lignes):
        versee, versements_exclus, source_versements = versements_depuis_fichiers(
            entrees["exercice"], params, mouvements, libelles, gl_lignes, pieces_exclues)
        a_confirmer = [x for x in versee if not x.get("qualifie")]
        if versee:
            avert.append(f"CMF / acomptes d'IS VERSÉS selon les fichiers ({source_versements}) : "
                         f"{sum(x['montant'] for x in versee):,.0f}, déduits en TF2 ligne 6 ({cite(params, 'cmf_acomptes')})."
                         + (f" {len(a_confirmer)} versement(s) retenu(s) sans libellé explicite de CMF : à confirmer." if a_confirmer else "")
                         + " Quittances à joindre au dépôt (art. 80.5), non exigées pour le calcul.")
        else:
            avert.append(f"Aucun versement de CMF ni d'acompte d'IS dans les fichiers ({source_versements}) : aucun débit du "
                         f"compte d'impôt sur les bénéfices sur l'exercice. TF2 ligne 6 = 0, IS à payer = impôt dû.")
        if versements_exclus:
            avert.append("Débits du compte d'IS NON retenus comme CMF de l'exercice : " + "; ".join(
                f"{x['date']} {x['piece']} {x['montant']:,.0f} ({x['motif']})" for x in versements_exclus))
    if versee is None:
        versee, source_versements = [], "aucune source (ni balance à mouvements, ni grand livre)"
        avert.append(f"CMF versée inconnue (aucun mouvement fourni) : 0 retenu. Si des acomptes ({cite(params, 'cmf_acomptes')}) ou "
                     f"des retenues sur paiements de l'État ({cite(params, 'cmf_retenue_etat')}) existent, l'IS à payer est "
                     f"surévalué (jusqu'à {cmf:,.0f}).")
    versee_tot = round(sum(x["montant"] for x in versee), 2)
    a_payer = round(du - versee_tot, 2)
    if faits.get("marches_publics") and not versee:
        avert.append(f"Société payée par l'État : des retenues de 1,5 % ({cite(params, 'cmf_retenue_etat')}) ont en principe "
                     "été subies sur chaque encaissement public. Elles valent paiement de CMF : les recenser.")

    # --- obligations annexes (non incluses dans TF2) : retenue sur non-résidents, art. 76
    annexes = []
    if faits.get("paiements_non_residents"):
        fact = d.get("facturation_non_residents")
        if fact is None:
            avert.append(f"Paiements à des non-résidents déclarés : montant de la facturation brute à fournir pour calculer "
                         f"la retenue de {P(params, 'non_residents_taux'):.0%} ({cite(params, 'non_residents_taux')}).")
        else:
            annexes.append({"objet": "Impôt des non-résidents (retenue sur facturation brute)", "base": fact,
                            "taux": P(params, "non_residents_taux"), "montant": round(fact * P(params, "non_residents_taux")),
                            "article": cite(params, "non_residents_taux")})
    for q in questions:
        ind = f" Indice(s) : {', '.join(i['compte'] + ' « ' + i['libelle'] + ' »' for i in q['indices'])}." if q["indices"] else ""
        avert.append(f"QUESTION ({q['article']}) : {q['question']}{ind} Sans réponse, la règle n'est pas appliquée.")
    if a_payer < 0:
        avert.append(f"Versements de CMF supérieurs à l'impôt dû : excédent de {-a_payer:,.0f}. Son sort (imputation, "
                     "restitution) n'est pas réglé par ce moteur : à examiner.")

    # --- dons : économie potentielle si les bénéficiaires sont éligibles (art. 64.2)
    dons = reint["donaciones"]
    if dons > 0:
        plafond = round(ca * P(params, "dons_plafond_ca"))
        eco = round(min(dons, plafond) * P(params, "is_taux")) if is_brut >= cmf else 0
        avert.append(f"Dons réintégrés en totalité ({dons:,.0f}). Si les bénéficiaires sont des organismes d'intérêt public "
                     f"situés en GE ({cite(params, 'dons_plafond_ca')}), {min(dons, plafond):,.0f} seraient déductibles "
                     f"(plafond 0,25 % du CA = {plafond:,.0f}) : économie d'IS possible ≈ {eco:,.0f}.")

    tf1 = {1: r1, **{v[0]: round(reint[k], 2) for k, v in TF1_REINT.items()}, 13: tot_reint,
           **{v[0]: round(deduc[k], 2) for k, v in TF1_DEDUC.items()}, 20: tot_deduc, 22: rf,
           24: round(impute, 2), 25: round(sum(x["montant"] for x in reste), 2), 26: rn}
    tf2 = {1: max(rn, 0.0), 2: exonere, 3: imposable, 4: P(params, "is_taux"), 5: is_brut, 6: versee_tot, 7: a_payer,
           10: ca, 11: diversos, 12: base_cmf, 13: P(params, "cmf_taux"), 14: cmf}
    return {"entrees": entrees, "tf1": tf1, "tf2": tf2, "base_is_arrondie": base, "impot_du": du,
            "interpretations": interp, "faits": faits, "questions": questions, "obligations_annexes": annexes,
            "reintegrations_detail": detail, "propositions": props, "deficits_reportables": reste,
            "versements_cmf": versee, "versements_exclus": versements_exclus, "source_versements": source_versements,
            "echeances": echeances(entrees, params), "avertissements": avert,
            "citations": {k: v["article"] for k, v in params["parametres"].items()}}


def echeances(entrees, params):
    clo = datetime.date.fromisoformat(entrees["date_cloture"])
    m = clo.month - 1 + P(params, "def_delai_mois")
    y, m = clo.year + m // 12, m % 12 + 1
    import calendar
    dlim = datetime.date(y, m, min(clo.day, calendar.monthrange(y, m)[1]))
    n = int(entrees["exercice"])
    return [
        {"objet": "Dépôt de la DEF", "date": dlim.isoformat(), "article": cite(params, "def_delai_mois")},
        {"objet": "Paiement de l'IS", "date": f"{P(params, 'is_paiement_jours')} jours après le dépôt effectif",
         "article": cite(params, "is_paiement_jours")},
        {"objet": f"1er acompte CMF {n}", "date": f"{n}-07-15", "article": cite(params, "cmf_acomptes")},
        {"objet": f"2e acompte CMF {n}", "date": f"{n + 1}-01-15", "article": cite(params, "cmf_acomptes")},
    ]


# --------------------------------------------------------------------------- Excel (formules vivantes)
def ecrire_xlsx(res, path, entite="", identifiant=""):
    """Classeur Paramètres / TF1 / TF2 : saisies en valeurs, TOUT le reste en formules
    (totaux, arrondi, taux lus dans l'onglet Paramètres — jamais écrits dans une formule)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    params = load_params()
    wb = Workbook()
    wp = wb.active; wp.title = "Parametros"
    wp.append(["Paramètre", "Valeur", "Article (Ley 1/2024)", "Texte"])
    row_of = {}
    for i, (k, v) in enumerate(params["parametres"].items(), start=2):
        val = v["valeur"] if not isinstance(v["valeur"], list) else ", ".join(v["valeur"])
        wp.append([k, val, v["article"], v["texte"]])
        row_of[k] = i
    # interprétations RETENUES pour cette entité (1 = oui, 0 = non) : visibles, et lues par les formules
    for k in ("arrondi_base_is", "arrondi_is_franc"):
        spec = params["interpretations"][k]
        v = res.get("interpretations", {}).get(k, spec["defaut"])
        wp.append([f"interpretation.{k}", 1 if v else 0, spec["article"],
                   spec["question"] + ("" if v == spec["defaut"] else "  [≠ lecture stricte : choix de l'entreprise]")])
        row_of["interpretation." + k] = wp.max_row
    ref = lambda k: f"Parametros!$B${row_of[k]}"
    saisie = PatternFill("solid", fgColor="FFF2CC")
    gras = Font(bold=True)
    tf1 = res["tf1"]

    ws = wb.create_sheet("TF1")
    ws.append(["Designación entidad:", entite, "", "Año terminado el", res["entrees"]["date_cloture"]])
    ws.append(["Número de identificación", identifiant])
    ws.append(["DETERMINACIÓN DEL RESULTADO FISCAL / DÉTERMINATION DU RÉSULTAT FISCAL"])
    ws.append(["Rubrique", "Ligne", "Montant"])
    lig = {}

    def put(label, n, val=None, formule=None, bold=False, input_=False):
        ws.append([label, n, formule if formule else val])
        r = ws.max_row
        lig[n] = r
        if bold:
            ws.cell(r, 1).font = gras; ws.cell(r, 3).font = gras
        if input_:
            ws.cell(r, 3).fill = saisie
        ws.cell(r, 3).number_format = "#,##0"

    put("I. Resultado contable antes de impuesto / Résultat comptable avant impôt", 1, tf1[1], input_=True, bold=True)
    ws.append(["II. REINTEGRACIONES / RÉINTÉGRATIONS"])
    for k, (n, es, fr) in TF1_REINT.items():
        put(f"{es} / {fr}", n, tf1[n], input_=True)
    put("TOTAL REINTEGRACIONES", 13, formule="=SUM(" + ",".join(f"C{lig[v[0]]}" for v in TF1_REINT.values()) + ")", bold=True)
    ws.append(["III. DEDUCCIONES / DÉDUCTIONS"])
    for k, (n, es, fr) in TF1_DEDUC.items():
        put(f"{es} / {fr}", n, tf1[n], input_=True)
    put("TOTAL DEDUCCIONES", 20, formule="=SUM(" + ",".join(f"C{lig[v[0]]}" for v in TF1_DEDUC.values()) + ")", bold=True)
    put("IV. RESULTADO NETO FISCAL / RÉSULTAT FISCAL", 22, formule=f"=C{lig[1]}+C{lig[13]}-C{lig[20]}", bold=True)
    put("V. Déficits antérieurs imputés (art. 73)", 24, tf1[24], input_=True)
    put("VI. RESULTADO NETO / RÉSULTAT NET FISCAL", 26, formule=f"=C{lig[22]}-C{lig[24]}", bold=True)
    ws.column_dimensions["A"].width = 78; ws.column_dimensions["C"].width = 20

    t2 = res["tf2"]
    w2 = wb.create_sheet("TF2")
    w2.append(["DETERMINACIÓN DEL IMPUESTO SOBRE SOCIEDADES / IS ET CUOTA MÍNIMA FISCAL"])
    w2.append(["Rubrique", "Ligne", "Montant"])
    l2 = {}

    def put2(label, n, val=None, formule=None, input_=False, fmt="#,##0"):
        w2.append([label, n, formule if formule else val])
        r = w2.max_row
        l2[n] = r
        w2.cell(r, 3).number_format = fmt
        if input_:
            w2.cell(r, 3).fill = saisie

    put2("Beneficio fiscal / Bénéfice fiscal", 1, formule=f"=MAX(TF1!C{lig[26]},0)")
    put2("Beneficio exonerado / Bénéfice exonéré", 2, t2[2], input_=True)
    put2("Beneficio imponible / Bénéfice imposable", 3, formule=f"=MAX(C3-C4,0)")
    put2("Tasa / Taux (art. 75.2)", 4, formule=f"={ref('is_taux')}", fmt="0.0%")
    put2("Impuesto sobre sociedades (arrondi de la base : voir Parametros, art. 75.1)", 5,
         formule=f"=ROUND(FLOOR(C5,{ref('is_arrondi_base')})*C6,0)")
    put2("Cuota mínima fiscal, a deducir / CMF versée (art. 80.2-80.3)", 6, t2[6], input_=True)
    put2("Impuesto sobre sociedades, a pagar / IS à payer", 7, formule=f"=MAX(C7,C{0})-C8")
    w2.append(["II. CUOTA MÍNIMA FISCAL"])
    put2("Cifra de negocio (cuentas 70) / Chiffre d'affaires", 10, t2[10], input_=True)
    put2("Ingresos y beneficios diversos / Autres produits (si retenus)", 11, t2[11], input_=True)
    put2("Cifra de negocio imponible / Base CMF", 12, formule=f"=C{0}+C{0}")
    put2("Tasa / Taux (art. 75.3)", 13, formule=f"={ref('cmf_taux')}", fmt="0.0%")
    put2("Cuota mínima a pagar / CMF", 14, formule=f"=ROUND(C{0}*C{0},0)")
    # formules qui dépendent de lignes écrites plus bas : références résolues maintenant
    w2.cell(l2[7], 3).value = f"=MAX(C{l2[5]},C{l2[14]})-C{l2[6]}"
    w2.cell(l2[12], 3).value = f"=C{l2[10]}+C{l2[11]}"
    arr_franc = ref("interpretation.arrondi_is_franc")
    arr_base = ref("interpretation.arrondi_base_is")
    w2.cell(l2[14], 3).value = f"=ROUND(C{l2[12]}*C{l2[13]},IF({arr_franc}=1,0,2))"
    w2.cell(l2[3], 3).value = f"=MAX(C{l2[1]}-C{l2[2]},0)"
    w2.cell(l2[5], 3).value = (f"=ROUND(IF({arr_base}=1,FLOOR(C{l2[3]},{ref('is_arrondi_base')}),C{l2[3]})"
                               f"*C{l2[4]},IF({arr_franc}=1,0,2))")
    w2.column_dimensions["A"].width = 70; w2.column_dimensions["C"].width = 20

    wa = wb.create_sheet("Avertissements")
    wa.append(["Avertissement / hypothèse par défaut"])
    for a in res["avertissements"]:
        wa.append([a])
    wa.append([]); wa.append(["Échéance", "Date", "Article"])
    for e in res["echeances"]:
        wa.append([e["objet"], e["date"], e["article"]])
    wa.column_dimensions["A"].width = 120
    wb.save(path)
    return path
