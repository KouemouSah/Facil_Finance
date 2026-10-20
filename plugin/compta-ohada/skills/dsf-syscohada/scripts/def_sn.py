# -*- coding: utf-8 -*-
"""DEF sur le modèle SYSCOHADA révisé — SYSTÈME NORMAL (formulaire officiel OHADA, datas/DSF_Normal_MODELE-PAR-DEFAUT.xlsx),
adaptée à la Guinée équatoriale et AUX DONNÉES de l'entité.

UN SEUL MOTEUR : les montants viennent du Generateur de def_modele (le même que le format GE, prouvé) ; ce module ne
fait que la PRÉSENTATION sur le formulaire officiel, dont il GARDE les formules (totaux, variations : ce sont les
formules officielles). Il écrit :
  - l'identité (page de garde, en-têtes « Désignation entité / Numéro d'identification ») ; mentions de la GE
    (République, ministère, DGIC) à la place du Cameroun ;
  - les ÉTATS par code de poste (bilan : brut, amortissements, net N-1 ; passif et compte de résultat : N ; TFT : N),
    en VALEURS, jamais par renvoi aux notes (le bilan reste juste même si une note est omise) ;
  - les NOTES ligne par ligne : correspondance ligne -> comptes (reprise du format GE, complétée à la main :
    references/correspondance-sn*.json), N sur la balance, N-1 des notes de bilan sur l'ouverture ;
  - les tableaux de MOUVEMENTS 3A (brut) et 3C (amortissements) depuis les mouvements du moteur ;
  - l'historique (31) et la fiche de synthèse (34) repris du calcul du format GE, libellé par libellé ;
  - les tableaux FISCAUX GE (TF1, Reintegraciones, TF2 à TF5) et la balance générale (art. 79.2), à la place des
    annexes fiscales camerounaises (CF1 à CF2 ter, C1/C2 des notes), sans objet en GE.
LIASSE ADAPTÉE AUX DONNÉES : une note sans montant est omise (sauf choix « conserver ») ; une note extra-comptable
n'est gardée (à saisir, fond ambre) que si la balance montre qu'elle s'applique (sûretés si dettes financières,
retraites si provision 196, effectifs si charges de personnel…) ; chaque omission est dite dans la grille d'analyse.
CONTRÔLE : pour chaque poste qu'une note détaille, somme des lignes = montant du poste au bilan ; un écart est dit.
"""
import json
import os
import re
import sys
import tempfile
import difflib
from copy import copy

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import def_modele as dm  # noqa: E402

# formulaire officiel OHADA système normal (liasse DGI Cameroun, vierge) : embarqué avec les gabarits, donc aussi
# dans l'extension Desktop
MODELE_SN = os.path.join(HERE, "..", "templates", "syscohada_sn_modele.xlsx")
REF = os.path.join(HERE, "..", "references")
AMBRE = "FFF2CC"
NUM = dm.NUM

# feuilles camerounaises sans objet en GE (annexes fiscales CF, tableaux fiscaux C1/C2 des notes)
CAMEROUN = re.compile(r"^(cf\d|autres annexes fisca|c[o]?\d ?-? ?note)", re.I)
# notes extra-comptables : gardées (à saisir) seulement si la balance montre qu'elles s'appliquent
DECLENCHEURS = {"note 1": ("16", "17", "18"), "note 3b": ("17",), "note 3d": ("81", "82"), "note 3e": ("106",),
                "note 16b": ("196",), "note 16bbis": ("196",), "note 27b": ("66",), "note 13": ("10",)}
TOUJOURS = ("note 2",)                     # informations obligatoires : toujours produites
SANS_DONNEE = ("note 3f", "note 16c", "note 35", "notes stat sociale et envi", "notes stat a caractere commerci",
               "note 32", "note 33")       # seulement si l'entité les renseigne (liste de la grille)
GE_MENTIONS = {"republique du cameroun": "RÉPUBLIQUE DE GUINÉE ÉQUATORIALE",
               "ministere des finances": "MINISTÈRE DE L'ÉCONOMIE, DES FINANCES ET DE LA PLANIFICATION",
               "direction generale des impots": "DIRECTION GÉNÉRALE DES IMPÔTS ET CONTRIBUTIONS (DGIC)"}


def plat(s):
    return re.sub(r"[^a-z0-9]+", " ", dm.norm(s or "")).strip()


def code(titre):
    t = plat(titre).replace("co1", "c1")
    m = re.match(r"(c\d )?note (\d+ ?[a-z]?( bis)?)", t)
    return ((m.group(1) or "") + "note " + m.group(2).replace(" ", "")).strip() if m else t


def correspondance():
    auto = json.load(open(os.path.join(REF, "correspondance-sn.json"), encoding="utf-8"))
    man = json.load(open(os.path.join(REF, "correspondance-sn-manuelle.json"), encoding="utf-8"))
    out = {}
    for k, f in auto["feuilles"].items():
        out[k] = {"roles": {int(c): ro for c, ro in f["roles"].items()}, "colonne_n": f["colonne_n"],
                  "lignes": {int(r): x["expr"] for r, x in f["lignes"].items()}}
    for k, lignes in man["feuilles"].items():            # la main prime toujours sur l'automatique
        out.setdefault(k, {"roles": {}, "colonne_n": None, "lignes": {}})["lignes"].update({int(r): e for r, e in lignes.items()})
    return out


class ProducteurSN:
    def __init__(self, g, ge_evalue=None):
        """g : Generateur (def_modele) déjà construit ; ge_evalue : classeur format GE produit ET évalué (valeurs),
        source des tableaux fiscaux, de l'historique (31) et de la synthèse (34)."""
        self.g, self.D, self.ge = g, g.D, ge_evalue
        rk = (lambda a: "129~" + a if a.startswith("13") else a) if not self.D["agg"]["controles"].get("balance_cloturee") else (lambda a: a)
        self.detail = g.detail
        agg1 = self.D.get("agg1")
        self.detail1 = [(rk(a), rf, col, v) for a, rf, col, v in agg1["detail"]] if agg1 else None
        self.corr = correspondance()
        self.controles, self.a_completer, self.omises, self.ajoutees = [], [], [], []
        soldes = self.g.soldes
        self.presents = {a[:k] for a, v in soldes.items() if abs(v) >= 1 for k in range(1, len(a) + 1)}

    # ------------------------------------------------------------------ utilitaires
    def a_saisir(self, ws, cell):
        from openpyxl.styles import PatternFill
        if type(cell).__name__ == "MergedCell":
            return
        cell.fill = PatternFill("solid", fgColor=AMBRE)
        self.a_completer.append((ws.title, cell.coordinate))

    @staticmethod
    def ecrire(cell, v):
        """Écrit un montant, sauf sur une formule INTERNE à la feuille (total officiel) : celle-ci est gardée."""
        if type(cell).__name__ == "MergedCell":
            return False
        f = cell.value
        if isinstance(f, str) and f.startswith("=") and "!" not in f:
            return False
        cell.value = v
        cell.number_format = NUM
        return True

    # ------------------------------------------------------------------ identité et page de garde
    def identite(self, wb):
        ident = self.D.get("identite") or {}
        nom, nif = ident.get("raison_sociale"), ident.get("nif")
        N = self.D["exercice"]
        for ws in wb.worksheets:
            for row in ws.iter_rows(max_row=8):
                for c in row:
                    if not isinstance(c.value, str):
                        continue
                    p = plat(c.value)
                    # l'en-tête porte aussi l'exercice et la durée : toujours écrits ; nom et NIF dès qu'ils sont saisis
                    if p.startswith("designation entite"):
                        c.value = f"Désignation entité : {nom or '……………………………'}      Exercice clos le 31-12-{N}"
                    elif p.startswith("numero d") and "identification" in p:   # d'/d’ : l'apostrophe varie
                        c.value = f"Numéro d'identification : {nif or '……………………………'}      Durée (en mois) : 12"
        pg = next((w for w in wb.worksheets if code(w.title) == "page de garde"), None)
        if pg is None:
            return
        for row in pg.iter_rows():
            for c in row:
                if not isinstance(c.value, str):
                    continue
                p = plat(c.value)
                for k, v in GE_MENTIONS.items():
                    if p.startswith(k):
                        c.value = v
                if p.startswith("denomination sociale"):
                    c.value = f"DÉNOMINATION SOCIALE : {nom}" if nom else c.value
                    if not nom:
                        self.a_saisir(pg, c)
                elif p.startswith("adresse complete"):
                    c.value = f"ADRESSE COMPLÈTE : {ident['adresse']}" if ident.get("adresse") else c.value
                elif p.startswith(("n d identification", "n didentification")):
                    c.value = f"N° D'IDENTIFICATION FISCALE : {nif}" if nif else c.value
                elif p.startswith("sigle usuel") and ident.get("sigle"):
                    c.value = f"SIGLE USUEL : {ident['sigle']}"
                elif p.startswith("exercice clos le"):
                    c.value = f"EXERCICE CLOS LE : 31/12/{self.D['exercice']}"

    # ------------------------------------------------------------------ états
    def etats(self, wb):
        pN, p1, agg = self.D["pN"], self.D["p1"], self.D["agg"]["par_ref"]
        ecrits = {}
        for ws in wb.worksheets:
            t = code(ws.title)
            if t not in ("bilan paysage", "compte de resultat", "tableau des flux de tresorerie"):
                continue
            h = next((r for r in range(1, 20) if any(plat(ws.cell(r, c).value) == "ref" for c in range(1, ws.max_column + 1))), 7)
            roles = dm.roles_colonnes(ws, h, h + 3, 2025)
            refs = [c for c in range(1, ws.max_column + 1) if plat(ws.cell(h, c).value) == "ref"]
            for r in range(h + 1, ws.max_row + 1):
                for rc in refs:
                    ref = str(ws.cell(r, rc).value or "").strip()
                    if not dm.est_ref(ref):
                        continue
                    cols = [c for c in roles if c > rc and (not [x for x in refs if rc < x < c])]
                    for c in cols:
                        ro = roles[c]
                        if t == "tableau des flux de tresorerie":
                            v = (self.D["tft"] or {}).get(ref) if ro == "n" else None
                        elif ro == "brut":
                            v = agg.get(ref, {}).get("brut", pN.get(ref))
                        elif ro == "amort":
                            v = agg.get(ref, {}).get("amort", 0.0)
                        elif ro in ("net_n", "n"):
                            v = pN.get(ref)
                        elif ro in ("net_n1", "n1"):
                            v = p1.get(ref)
                        else:
                            continue
                        # (une formule du formulaire — net = brut - amortissements, sous-total — est gardée par ecrire)
                        if v is not None and ref in pN | p1 | (self.D["tft"] or {}):
                            if self.ecrire(ws.cell(r, c), round(v, 2)) or ro in ("net_n", "n"):
                                ecrits.setdefault(t, set()).add(ref)
        # tout poste CHIFFRÉ du bilan / compte de résultat doit trouver sa ligne dans le formulaire : une ligne sans
        # code (ex. « Ecart de conversion - Passif ») est retrouvée par son LIBELLÉ ; sinon, c'est dit
        for t, etat in (("bilan paysage", ("actif", "passif")), ("compte de resultat", ("resultat",))):
            ws = next((w for w in wb.worksheets if code(w.title) == t), None)
            manque = [(ref, v) for ref, v in pN.items() if abs(v) >= 1 and self.D["postes"].get(ref, {}).get("etat") in etat
                      and not self.D["postes"].get(ref, {}).get("total") and ref not in ecrits.get(t, set())]
            for ref, v in list(manque):
                cible = plat(self.D["postes"][ref].get("libelle", ""))
                cols = (11, 12) if ref[0] in "CD" and t == "bilan paysage" else (6, 7) if t == "bilan paysage" else (5, 6)
                lab_col = 9 if cols[0] == 11 else 2
                ligne = next((r for r in range(8, ws.max_row + 1) if cible and difflib.SequenceMatcher(
                    None, plat(ws.cell(r, lab_col).value)[:30], cible[:30]).ratio() >= 0.75), None) if ws else None
                if ligne and self.ecrire(ws.cell(ligne, cols[0]), round(v, 2)):
                    if p1.get(ref) is not None:
                        self.ecrire(ws.cell(ligne, cols[1]), round(p1[ref], 2))
                    manque.remove((ref, v))
                    self.controles.append((f"{t} : poste {ref} sans code dans le formulaire, rattaché par son libellé",
                                           True, f"ligne {ligne} « {str(ws.cell(ligne, lab_col).value).strip()[:40]} »"))
            self.controles.append((f"{t} : chaque poste chiffré a sa ligne dans le formulaire", not manque,
                                   ", ".join(f"{r} {v:,.0f}" for r, v in manque) or "tous présents"))
            if ws is not None and t == "bilan paysage":
                self.corriger_total_general(ws)

    def corriger_total_general(self, ws):
        """Défaut du formulaire officiel : le TOTAL GÉNÉRAL (=K38+K34+K27) oublie la ligne « Écart de conversion ».
        Chaque ligne d'écart de conversion non comptée est ajoutée à la formule du total de SON côté — et dit."""
        tot = next((r for r in range(8, ws.max_row + 1) for c in (2, 9) if plat(ws.cell(r, c).value).startswith("total general")), None)
        if not tot:
            return
        for lab_col, cols in ((2, (4, 5, 6, 7)), (9, (11, 12))):
            ecarts = [r for r in range(8, ws.max_row + 1) if plat(ws.cell(r, lab_col).value).startswith("ecart de conversion")]
            for c in cols:
                f = ws.cell(tot, c).value
                L = dm.Generateur.L(c)
                if not (isinstance(f, str) and f.startswith("=")):
                    continue
                oublies = [r for r in ecarts if not re.search(rf"\b{L}{r}\b", f)]
                if oublies:
                    ws.cell(tot, c).value = f + "".join(f"+N({L}{r})" for r in oublies)
                    self.controles.append((f"Formule du formulaire CORRIGÉE : {L}{tot} (total général) oubliait l'écart "
                                           f"de conversion", True, f"{f} -> {ws.cell(tot, c).value}"))

    # ------------------------------------------------------------------ notes par correspondance
    def notes(self, wb):
        for ws in wb.worksheets:
            k = code(ws.title)
            f = self.corr.get(k)
            if not f or not f["lignes"]:
                continue
            roles = f["roles"] or dm.roles_colonnes(ws, 7, 10, 2025)
            cn = f["colonne_n"] or next((c for c, ro in sorted(roles.items()) if ro in ("n", "net_n")), None)
            c1 = next((c for c, ro in sorted(roles.items()) if ro in ("n1", "net_n1") and c > (cn or 0)), None)
            par_poste, autres = {}, {}
            for r, expr in sorted(f["lignes"].items()):
                ex = dm.expr_comptes(expr)
                if not ex or not cn:
                    continue
                ref, termes = ex
                vn = dm.valeur_detail(ref, termes, self.detail)
                self.ecrire(ws.cell(r, cn), vn if abs(vn) >= 0.5 else 0)
                if c1 and self.detail1 is not None and ref[0] in "ABCD":
                    v1 = dm.valeur_detail(ref, termes, self.detail1)
                    self.ecrire(ws.cell(r, c1), v1 if abs(v1) >= 0.5 else 0)
                self.echeances(ws, r, roles, cn)
                sgn = -1 if termes[0][1][:2] in ("28", "29", "39", "49", "59") else 1
                par_poste.setdefault(ref, 0.0)
                par_poste[ref] += sgn * vn
                if sgn > 0 and plat(ws.cell(r, 1).value).startswith("autre"):
                    autres[ref] = r                       # la DERNIÈRE ligne « Autres… » du poste
            for ref, tot in par_poste.items():            # la note retombe-t-elle sur son poste ?
                attendu = self.D["pN"].get(ref)
                if attendu is None:
                    continue
                ecart = round(abs(attendu) - abs(tot), 2)
                if abs(ecart) >= 1 and ref in autres:
                    # comptes que le poste contient mais qu'aucune ligne de la note ne cite (hors nomenclature du
                    # formulaire, ex. 611 transports sur achats) : versés à la ligne « Autres… » du poste, comme le
                    # format GE — et DIT
                    cell = ws.cell(autres[ref], cn)
                    if self.ecrire(cell, round((cell.value if dm.num(cell.value) else 0) + ecart, 2)):
                        self.controles.append((f"{ws.title.strip()} : comptes du poste {ref} hors nomenclature du formulaire",
                                               True, f"{ecart:,.0f} présentés ligne {autres[ref]} « "
                                               f"{str(ws.cell(autres[ref], 1).value).strip()[:40]} »"))
                        tot = tot + ecart if tot >= 0 else tot - ecart
                ok = abs(abs(tot) - abs(attendu)) < 1
                self.controles.append((f"{ws.title.strip()} : lignes du poste {ref} = poste au bilan / résultat", ok,
                                       f"note {abs(tot):,.0f} / poste {abs(attendu):,.0f}"))

    def echeances(self, ws, r, roles, cn):
        """Échéancier : tout « à un an au plus » PAR DÉFAUT, par formule (= N − plus d'un an − plus de deux ans) ; les
        deux autres colonnes à 0, à ventiler par l'entreprise — comme le format GE."""
        col = {ro: c for c, ro in sorted(roles.items(), reverse=True)}
        if "ech1" not in col:
            return
        L = dm.Generateur.L
        autres = "".join(f"-N({L(col[k])}{r})" for k in ("ech2", "ech3") if k in col)
        if self.ecrire(ws.cell(r, col["ech1"]), f"=N({L(cn)}{r}){autres}"):
            for k in ("ech2", "ech3"):
                if k in col:
                    self.ecrire(ws.cell(r, col[k]), 0)

    # ------------------------------------------------------------------ mouvements d'immobilisations
    def mouvements(self, wb):
        n3, postes = self.D.get("n3") or {}, self.D["postes"]
        lib = {ref: plat(p.get("libelle", "")) for ref, p in postes.items() if ref in n3}
        for ws in wb.worksheets:
            k = code(ws.title)
            if k not in ("note 3a", "note 3c"):
                continue
            cols = ({"b_ouv": 4, "b_aug": 5, "b_vir+": 6, "b_dim": 8, "b_vir-": 9} if k == "note 3a"
                    else {"a_ouv": 2, "a_dot": 3, "a_dim": 4})
            faits = set()

            def score(lb, ref):                         # « Terrains hors immeubles de placement » -> poste « Terrains »
                l2 = lib[ref]
                if l2 and (lb.startswith(l2) or l2.startswith(lb)):
                    return 1.0
                return difflib.SequenceMatcher(None, lb[:30], l2[:30]).ratio()
            for r in range(9, ws.max_row + 1):
                lb = plat(ws.cell(r, 1).value)
                if not lb or lb.startswith(("sous total", "total", "immobilisations")):
                    continue
                ref = max((x for x in lib if x not in faits), key=lambda x: score(lb, x), default=None)
                if not ref or score(lb, ref) < 0.6:
                    continue
                faits.add(ref)
                L = n3[ref]
                vals = ({"b_ouv": L["b_ouv"], "b_aug": L["b_aug"], "b_vir+": max(L["b_vir"], 0), "b_dim": L["b_dim"],
                         "b_vir-": max(-L["b_vir"], 0)} if k == "note 3a" else
                        {"a_ouv": L["a_ouv"], "a_dot": L["a_dot"], "a_dim": L["a_rep"] + max(-L["a_vir"], 0) - max(L["a_vir"], 0)})
                for cle, c in cols.items():
                    self.ecrire(ws.cell(r, c), round(vals[cle], 2))
            # chaque poste qui a des mouvements doit avoir trouvé sa ligne : sinon la note ne retombe pas sur le bilan
            # (3A : mouvements BRUTS « b_* » ; 3C : AMORTISSEMENTS « a_* » — un titre de participation ne s'amortit pas)
            pre = "b_" if k == "note 3a" else "a_"
            oublies = [ref for ref, L in n3.items() if ref not in faits
                       and any(abs(v) >= 1 for c_, v in L.items() if c_.startswith(pre) and dm.num(v))]
            self.controles.append((f"{ws.title.strip()} : chaque poste immobilisé a sa ligne de mouvements", not oublies,
                                   ", ".join(f"{r} {postes.get(r, {}).get('libelle', '')}" for r in oublies) or "tous placés"))

    # ------------------------------------------------------------------ 31 et 34 : repris du calcul du format GE
    def depuis_ge(self, wb):
        if self.ge is None:
            return
        for ws in wb.worksheets:
            k = code(ws.title)
            if k not in ("note 31", "note 34"):
                continue
            src = next((w for w in self.ge.worksheets if code(w.title) == k), None)
            if src is None:
                continue
            # colonne N du format GE repérée par son EN-TÊTE (jamais « le premier nombre » : ce peut être un n° de ligne)
            G = dm.Generateur.__new__(dm.Generateur)
            h0, h1 = G.entetes(src, 1, src.max_row)
            roles = dm.roles_colonnes(src, h0, h1, self.D["exercice"])
            cg = next((c for c, ro in sorted(roles.items()) if ro in ("n", "net_n")), None)
            if cg is None:
                continue
            sl = {}
            for r in range(h1 + 1, src.max_row + 1):
                lb = next((plat(src.cell(r, c).value) for c in (1, 2, 3) if isinstance(src.cell(r, c).value, str)
                           and not dm.expr_comptes(src.cell(r, c).value)), "")
                if lb and dm.num(src.cell(r, cg).value):
                    sl[lb] = src.cell(r, cg).value
            for r in range(1, ws.max_row + 1):
                lb = plat(ws.cell(r, 1).value)
                if not lb:
                    continue
                m = difflib.get_close_matches(lb, list(sl), n=1, cutoff=0.8)
                if m:
                    self.ecrire(ws.cell(r, 2), round(sl[m[0]], 2))

    # ------------------------------------------------------------------ fiscalité GE, balance, feuilles camerounaises
    def fiscal_ge(self, wb):
        for ws in list(wb.worksheets):
            if CAMEROUN.match(plat(ws.title)):
                self.omises.append((ws.title.strip(), "annexe fiscale camerounaise, sans objet en GE"))
                del wb[ws.title]
        if self.ge is None:
            return
        for src in self.ge.worksheets:
            t = code(src.title)
            # (« Réintégrations » en français, « Reintegraciones » en espagnol : la même feuille)
            if t in ("tf1", "reintegraciones", "reintegrations", "tf2", "tf3", "tf4", "tf5") or t.startswith("balance genera") \
                    or t.startswith("bg "):
                ws = wb.create_sheet(src.title[:31])
                copier_feuille(src, ws)
                self.ajoutees.append(ws.title)

    # ------------------------------------------------------------------ extra-comptable : à saisir si applicable
    def extra(self, wb):
        for ws in wb.worksheets:
            k = code(ws.title)
            if k in DECLENCHEURS and any(p in self.presents for p in DECLENCHEURS[k]):
                self.marquer_saisie(ws)

    def marquer_saisie(self, ws):
        for r in range(8, ws.max_row + 1):
            for c in range(2, min(ws.max_column, 16) + 1):
                cell = ws.cell(r, c)
                if cell.value is None and type(cell).__name__ != "MergedCell" and ws.cell(r, 1).value:
                    self.a_saisir(ws, cell)
                    break

    # ------------------------------------------------------------------ liasse adaptée aux données
    def adapter(self, wb, chemin_evalue, conserver=False):
        import openpyxl
        ev = openpyxl.load_workbook(chemin_evalue, data_only=True)
        a_saisir = {s for s, _ in self.a_completer}
        retirer = []
        for ws in wb.worksheets:
            k = code(ws.title)
            if not (k.startswith("note") or k.startswith("notes stat")) or k in TOUJOURS or ws.title in a_saisir:
                continue
            if k in DECLENCHEURS:                      # extra-comptable non déclenchée
                retirer.append((ws.title, "non applicable d'après la balance"))
                continue
            if k in SANS_DONNEE:
                retirer.append((ws.title, "aucune information reçue — à produire si l'entité la renseigne"))
                continue
            e = ev[ws.title]
            montant = any(dm.num(c.value) and abs(c.value) >= 0.5 for row in e.iter_rows(min_row=8) for c in row[1:])
            if not montant:
                retirer.append((ws.title, "néant — aucun montant N ni N-1"))
        if conserver:
            return []
        noms = {t for t, _ in retirer}
        for ws in wb.worksheets:                        # renvois vers une feuille retirée : figés à leur valeur
            if ws.title in noms:
                continue
            e = ev[ws.title] if ws.title in ev.sheetnames else None
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("=") and any(f"'{n}'!" in c.value or f"{n}!" in c.value for n in noms):
                        c.value = e[c.coordinate].value if e is not None else None
        for t, motif in retirer:
            del wb[t]
            self.omises.append((t.strip(), motif))
        self.grille(wb)
        return retirer

    def grille(self, wb):
        g = next((w for w in wb.worksheets if code(w.title).startswith("grille")), None)
        if g is None:
            return
        omis = {code(t): m for t, m in self.omises}
        presents = {code(w.title) for w in wb.worksheets}
        col = g.max_column + 1
        h = next((r for r in range(1, 15) if plat(g.cell(r, 1).value).startswith("notes")), 6)
        g.cell(h, col, "Présentation").font = copy(g.cell(h, 2).font)
        from openpyxl.utils import get_column_letter
        g.column_dimensions[get_column_letter(col)].width = 44
        for r in range(h + 1, g.max_row + 1):
            k = code(g.cell(r, 1).value or "")
            if k in presents:
                g.cell(r, col, "présentée")
            elif k in omis:
                g.cell(r, col, omis[k])

    # ------------------------------------------------------------------ production
    def produire(self, sortie, conserver=False):
        import openpyxl
        import formules_xl
        wb = openpyxl.load_workbook(MODELE_SN)
        self.identite(wb)
        self.etats(wb)
        self.notes(wb)
        self.mouvements(wb)
        self.depuis_ge(wb)
        self.extra(wb)
        self.fiscal_ge(wb)
        wb.save(sortie)
        ev = os.path.join(tempfile.mkdtemp(), "sn_eval.xlsx")
        rap = formules_xl.evaluer_classeur(sortie, ev)
        self.adapter(wb, ev, conserver)
        wb.active = next((i for i, w in enumerate(wb.worksheets) if code(w.title) == "page de garde"), 0)
        wb.save(sortie)
        formules_xl.evaluer_classeur(sortie, ev)
        self.controles.append(("Formules du formulaire évaluées par le moteur", not rap["non_supportees"],
                               f"{rap['formules']} formules ; non supportées : {len(rap['non_supportees'])}"))
        return {"xlsx": sortie, "evalue": ev, "feuilles": len(wb.worksheets), "omises": self.omises,
                "ajoutees": self.ajoutees, "controles": self.controles, "a_completer": len(self.a_completer)}


def copier_feuille(src, ws):
    """Copie VALEURS + mise en forme d'une feuille évaluée (tableaux fiscaux GE) vers le classeur SN."""
    for row in src.iter_rows():
        for c in row:
            if type(c).__name__ == "MergedCell":
                continue
            d = ws.cell(c.row, c.column, c.value)
            if c.has_style:
                d.font, d.fill, d.border = copy(c.font), copy(c.fill), copy(c.border)
                d.alignment, d.number_format = copy(c.alignment), c.number_format
    for rg in src.merged_cells.ranges:
        ws.merge_cells(str(rg))
    for k, dim in src.column_dimensions.items():
        ws.column_dimensions[k].width = dim.width
    for k, dim in src.row_dimensions.items():
        if dim.height:
            ws.row_dimensions[k].height = dim.height
    ws.page_setup.orientation = src.page_setup.orientation
    if src.print_area:
        try:
            ws.print_area = src.print_area if isinstance(src.print_area, str) else ",".join(src.print_area)
        except Exception:
            pass
