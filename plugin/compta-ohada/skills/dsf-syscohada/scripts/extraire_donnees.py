# -*- coding: utf-8 -*-
"""Extracteur de données extra-comptables (mode adaptatif) — LIT les documents de
l'utilisateur (quel que soit leur format clé-valeur) et FORMALISE vers le besoin
normalisé (`besoin-extra-comptable.json`), au lieu d'imposer un gabarit à remplir.

Couvre le TABULAIRE / CLÉ-VALEUR (identité, effectifs, informations générales) :
scanne les cellules, repère un libellé connu et prend la valeur voisine (droite,
sinon dessous). Alias + normalisation accents/casse.

Ne couvre PAS le non structuré (contrats/engagements) ni le jugement (retraitements) :
ceux-là = lecture par Claude + confirmation, ou saisie humaine (voir le skill).

Usage :
  python extraire_donnees.py --fichier "Informations generales.xlsx"
  python extraire_donnees.py --fichier infos.csv --json
"""
import argparse
import csv
import json
import os
import sys
import unicodedata

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

KEY_ALIASES = {
    "raison_sociale": ["raisonsociale", "denomination", "nomentreprise", "designationentite", "designationdelentite", "entite", "nomdelasociete"],
    "nif": ["nif", "numeroidentification", "numeroidentifiant", "identifiantfiscal", "niu", "numerodidentification"],
    "forme_juridique": ["formejuridique", "forme"],
    "regime_fiscal": ["regimefiscal", "regimedimposition", "regime"],
    "exercice_clos": ["exercice", "exerciceclos", "datecloture", "clotureexercice", "periode"],
    "adresse": ["adresse", "siege", "siegesocial"],
    "effectif_total": ["effectiftotal", "effectif", "nombresalaries", "nombredemployes", "totaleffectif", "nombredetravailleurs"],
    "effectif_cadres": ["effectifcadres", "cadres", "nombredecadres"],
    "masse_salariale": ["massesalariale", "totalsalaires", "salairesbruts", "totalmassesalariale"],
}


def _norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return "".join(ch for ch in s.lower() if ch.isalnum())


def _match_key(cell_norm):
    if not cell_norm or len(cell_norm) < 3:
        return None
    for key, als in KEY_ALIASES.items():
        if cell_norm in als:
            return key
    for key, als in KEY_ALIASES.items():
        if any(al in cell_norm or (len(cell_norm) >= 3 and cell_norm in al) for al in als):
            return key
    return None


def load_grid(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        grid = []
        for s in wb.sheetnames[:6]:
            for r in wb[s].iter_rows(values_only=True):
                grid.append([("" if c is None else str(c).strip()) for c in r])
        wb.close()
        return grid
    with open(path, encoding="utf-8-sig", newline="") as f:
        sample = f.read(4096); f.seek(0)
        delim = ";" if sample.count(";") >= sample.count(",") else ","
        return [[str(c).strip() for c in row] for row in csv.reader(f, delimiter=delim)]


def extract(grid):
    found, evidence = {}, {}
    for r, row in enumerate(grid):
        for c, cell in enumerate(row):
            key = _match_key(_norm(cell))
            if not key or key in found:
                continue
            # valeur = voisin droit non vide, sinon dessous
            val = ""
            if c + 1 < len(row) and str(row[c + 1]).strip():
                val = str(row[c + 1]).strip()
            elif r + 1 < len(grid) and c < len(grid[r + 1]) and str(grid[r + 1][c]).strip():
                val = str(grid[r + 1][c]).strip()
            if val:
                found[key] = val
                evidence[key] = cell
    return found, evidence


def main():
    p = argparse.ArgumentParser(description="Extraction extra-comptable (clé-valeur, adaptatif).")
    p.add_argument("--fichier", required=True)
    p.add_argument("--json", action="store_true")
    a = p.parse_args()
    grid = load_grid(a.fichier)
    found, evidence = extract(grid)
    besoin_tabulaire = list(KEY_ALIASES.keys())
    manquants = [k for k in besoin_tabulaire if k not in found]
    if a.json:
        print(json.dumps({"extrait": found, "manquants": manquants}, ensure_ascii=False, indent=2))
        return
    print("== Données extraites (À CONFIRMER) ==")
    for k, v in found.items():
        print(f"  {k:20s} = {v!r:30s}  (libellé source : {evidence[k]!r})")
    print("\n== Clés clé-valeur non trouvées (à lire ailleurs / demander) ==")
    print("  " + (", ".join(manquants) if manquants else "aucune"))
    print("\nRappel : engagements/sûretés (non structuré) = lecture par Claude + confirmation ; "
          "retraitements = jugement fourni. La détection est une proposition à valider.")


if __name__ == "__main__":
    main()
