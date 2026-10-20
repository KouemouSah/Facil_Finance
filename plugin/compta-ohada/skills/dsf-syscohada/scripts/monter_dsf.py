# -*- coding: utf-8 -*-
"""Orchestrateur de montage DSF — chaîne unifiée « fichiers -> DSF ».

Enchaîne automatiquement :
  1. classer_documents.py  -> devine le rôle de chaque fichier (balance, registre,
     tiers, modèle DSF, inconnu) — à CONFIRMER.
  2. extraire_donnees.py   -> lit les données extra-comptables (clé-valeur) des
     fichiers non-comptables et les formalise vers le besoin.
  3. agreger_balance_dsf   -> squelette Bilan/CR depuis la balance.
  4. registre_immobilisations -> Note 3C (si registre détecté/fourni).
  5. PORTE DE COMPLÉTUDE    -> liste ce qui manque, ne produit jamais une DSF
     présentée comme complète si elle ne l'est pas.
  6. Peuplement du modèle   -> si un cell-map est fourni (copie du modèle).

La détection est une PROPOSITION à confirmer. Le jugement (retraitements, fiscal)
et les engagements/sûretés non structurés restent hors périmètre auto.

Usage :
  python monter_dsf.py --dossier "C:/mes_docs" --exercice 2025
  python monter_dsf.py --fichiers bal.xlsx parc.csv infos.csv --exercice 2025
  python monter_dsf.py --balance bal.xlsx --registre parc.csv --exercice 2025 --cell-map cm.json --sortie out.xlsx
"""
import argparse
import glob
import json
import os
import shutil
import subprocess
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

HERE = os.path.dirname(os.path.abspath(__file__))
AGREGER = os.path.join(HERE, "agreger_balance_dsf.py")
CLASSER = os.path.join(HERE, "classer_documents.py")
EXTRAIRE = os.path.join(HERE, "extraire_donnees.py")
REGISTRE = os.path.normpath(os.path.join(HERE, "..", "..", "immobilisations-ohada", "scripts", "registre_immobilisations.py"))
DEFAULT_MODELE = os.path.normpath(os.path.join(HERE, "..", "..", "..", "datas", "DSF_Normal_MODELE-PAR-DEFAUT.xlsx"))


def run_json(script, args):
    # errors="replace" : les sous-scripts écrivent leur diagnostic sur stderr en
    # encodage console (cp1252 sous Windows) ; on ne veut pas planter le décodage.
    cp = subprocess.run([sys.executable, script] + args + ["--json"],
                        capture_output=True, text=True, encoding="utf-8", errors="replace")
    if cp.returncode != 0:
        raise SystemExit(f"Échec {os.path.basename(script)} : {(cp.stderr or '').strip()}")
    return json.loads(cp.stdout)


def gather_files(a):
    files = list(a.fichiers or [])
    if a.dossier:
        for ext in ("*.xls", "*.xlsx", "*.xlsm", "*.csv"):
            files += glob.glob(os.path.join(a.dossier, ext))
    # inclure les overrides explicites
    for f in (a.balance, a.registre, a.tiers, a.modele):
        if f and f not in files:
            files.append(f)
    # dédoublonner en gardant l'ordre
    seen, out = set(), []
    for f in files:
        if f not in seen:
            seen.add(f); out.append(f)
    return out


def classify_all(files):
    """Retourne {basename: role} via classer_documents (proposition à confirmer)."""
    if not files:
        return {}
    res = run_json(CLASSER, ["--fichiers"] + files)
    return {d["fichier"]: d["role"] for d in res.get("documents", [])}


def main():
    p = argparse.ArgumentParser(description="Montage DSF orchestré (fichiers -> DSF).")
    p.add_argument("--fichiers", nargs="*", default=[])
    p.add_argument("--dossier", default=None)
    p.add_argument("--exercice", type=int, default=None)
    # overrides explicites (prioritaires sur la détection)
    p.add_argument("--balance", default=None)
    p.add_argument("--registre", default=None)
    p.add_argument("--tiers", default=None)
    p.add_argument("--intake", default=None)
    p.add_argument("--modele", default=None)
    p.add_argument("--gabarit", default=None,
                   help="lecture imposée des colonnes de la balance, si la détection l'a signalée ambiguë")
    p.add_argument("--cell-map", default=None)
    p.add_argument("--sortie", default=None)
    p.add_argument("--decisions-fiscales", default=None,
                   help="JSON des jugements fiscaux (propositions acceptées, réintégrations, déficits, CMF versée)")
    p.add_argument("--xlsx-tf", default=None, help="classeur TF1/TF2 à produire (formules vivantes)")
    p.add_argument("--profil", default=None, help="profil d'entreprise : interprétations, faits, justifications")
    p.add_argument("--grand-livre", default=None, help="grand livre (sinon : celui reconnu parmi les fichiers)")
    p.add_argument("--xlsx-tf345", default=None, help="classeur TF3/TF4/TF5 à produire (formules vivantes + traçabilité)")
    a = p.parse_args()

    files = gather_files(a)
    by_name = {os.path.basename(f): f for f in files}

    # 1) CLASSIFICATION (proposition à confirmer)
    print("== 1. Classification des documents (À CONFIRMER par l'utilisateur) ==")
    roles = classify_all(files)
    for name, role in roles.items():
        print(f"  {name[:42]:42s} -> {role}")

    def first_of(role):
        return next((by_name[n] for n, r in roles.items() if r == role), None)

    balance = a.balance or first_of("balance")
    registre = a.registre or first_of("registre_immobilisations")
    tiers = a.tiers or first_of("detail_tiers")
    modele = a.modele or first_of("modele_dsf") or DEFAULT_MODELE
    autres = [by_name[n] for n, r in roles.items() if r in ("inconnu", "detail_tiers")]

    if not balance:
        raise SystemExit("Aucune balance détectée. Fournir --balance ou un fichier balance.")

    # 2) EXTRACTION extra-comptable (lire -> formaliser)
    print("\n== 2. Lecture des données extra-comptables ==")
    intake = {}
    if a.intake:
        intake = json.load(open(a.intake, encoding="utf-8"))
        print(f"  Intake fourni : {len(intake)} clés")
    else:
        for f in autres:
            try:
                ext = run_json(EXTRAIRE, ["--fichier", f]).get("extrait", {})
            except SystemExit:
                ext = {}
            for k, v in ext.items():
                intake.setdefault(k, v)
        print(f"  Extrait des documents : {len(intake)} clés " + (str(list(intake)) if intake else "(aucune)"))

    # 3) ASSEMBLAGE
    print("\n== 3. Assemblage du pré-remplissage ==")
    prefill = run_json(AGREGER, ["--balance", balance] + (["--gabarit", a.gabarit] if a.gabarit else [])
                       + (["--profil", a.profil] if a.profil else []))
    print(f"  Balance '{os.path.basename(balance)}' : {len(prefill['par_ref'])} postes, {prefill['nb_non_mappes']} non mappés")
    if registre and a.exercice:
        reg = run_json(REGISTRE, ["--fichier", registre, "--exercice", str(a.exercice)])
        prefill["note_3c"] = reg.get("note_3c", {})
        print(f"  Registre '{os.path.basename(registre)}' : Note 3C sur {len(prefill['note_3c'])} catégories")
    elif registre and not a.exercice:
        print("  Registre détecté mais --exercice manquant : Note 3C non calculée.")
    prefill["intake"] = intake

    # 3bis) TF1/TF2 (Guinée équatoriale) : résultat fiscal, IS, Cuota Mínima — moteur fiscal_ge
    if a.exercice:
        sys.path.insert(0, HERE)
        import fiscal_ge
        from detection_balance import analyse as _analyse
        det = _analyse(balance, a.gabarit)
        dec = json.load(open(a.decisions_fiscales, encoding="utf-8")) if a.decisions_fiscales else None
        profil = json.load(open(a.profil, encoding="utf-8")) if a.profil else None
        tf = fiscal_ge.calculer(fiscal_ge.entrees_depuis_soldes(det["soldes"], a.exercice), dec,
                                det["soldes"], libelles=det.get("libelles"), profil=profil,
                                mouvements=det.get("mouvements"))   # CMF versée lue dans la balance
        prefill["tf"] = tf
        print(f"  TF1/TF2 : résultat fiscal {tf['tf1'][22]:,.0f} | IS {tf['tf2'][5]:,.0f} | CMF {tf['tf2'][14]:,.0f}"
              f" | IS à payer {tf['tf2'][7]:,.0f}" + ("" if dec else "  (aucune décision fiscale : propositions NON appliquées)"))
        for w in tf["avertissements"]:
            print("  ⚠ " + w)
        if a.xlsx_tf:
            fiscal_ge.ecrire_xlsx(tf, a.xlsx_tf)
            print(f"  Classeur TF1/TF2 (formules vivantes) : {a.xlsx_tf}")
    else:
        print("  TF1/TF2 non calculés : --exercice manquant.")

    # 3ter) TF3/TF4/TF5 depuis le grand livre (fourni, ou reconnu par classer_documents)
    gl_path = a.grand_livre or first_of("grand_livre")
    if gl_path and a.exercice:
        import grand_livre
        import tf345_ge
        from detection_balance import analyse as _analyse
        gl = grand_livre.lire(gl_path)
        if gl["detection"].get("statut") == "ok":
            b = _analyse(balance, a.gabarit)
            profil = json.load(open(a.profil, encoding="utf-8")) if a.profil else None
            t = tf345_ge.calculer(gl, profil=profil, libelles=b.get("libelles"), soldes_balance=b["soldes"])
            prefill["tf345"] = {k: {kk: vv for kk, vv in v.items() if kk != "trace"} for k, v in t.items() if k.startswith("tf")}
            n, x = t["tf4"]["categories"]["national"], t["tf4"]["categories"]["non_national"]
            print(f"  TF3 : IVA {t['tf3']['iva_theorique']:,.0f} | déductible {t['tf3']['iva_deductible']:,.0f} | à payer "
                  f"{t['tf3']['ligne_29']:,.0f} — TF4 : IRPF {n['irpf_retenu'] + x['irpf_retenu']:,.0f} — TF5 : "
                  f"{len(t['tf5']['lignes'])} tiers, {t['tf5']['total_ht']:,.0f} HT")
            for w in t["avertissements"]:
                print("  ⚠ " + w)
            if a.xlsx_tf345:
                tf345_ge.ecrire_xlsx(t, a.xlsx_tf345)
                print(f"  Classeur TF3/TF4/TF5 : {a.xlsx_tf345}")
        else:
            print(f"  TF3/TF4/TF5 non calculés : grand livre illisible ({gl['detection'].get('raison')}).")
    else:
        print("  TF3/TF4/TF5 non calculés : grand livre absent (--grand-livre) ou --exercice manquant.")

    # 4) PORTE DE COMPLÉTUDE
    print("\n== 4. Porte de complétude ==")
    checks = [
        ("Balance après inventaire", bool(balance)),
        ("Registre d'immobilisations (Note 3)", bool(registre) and bool(a.exercice)),
        ("Détail tiers (Notes 7/17)", bool(tiers)),
        ("Données extra-comptables (identité/effectifs)", bool(intake)),
        ("Cell-map du modèle", bool(a.cell_map)),
        ("Modèle DSF", bool(modele)),
    ]
    manque = [lbl for lbl, ok in checks if not ok]
    for lbl, ok in checks:
        print(f"  [{'x' if ok else ' '}] {lbl}")
    print(f"\nÉtat : {'COMPLET' if not manque else 'INCOMPLET — ' + ' ; '.join(manque)}")
    print("Rappel : engagements/sûretés (non structuré) + retraitements (jugement) = à fournir/confirmer.")

    # 5) PEUPLEMENT ou pré-remplissage
    if a.cell_map and os.path.exists(modele):
        from openpyxl import load_workbook
        cm = json.load(open(a.cell_map, encoding="utf-8"))
        sortie = a.sortie or os.path.join(os.getcwd(), "DSF_pre-remplie.xlsx")
        shutil.copyfile(modele, sortie)
        wb = load_workbook(sortie)
        ecrits, ignores = 0, []
        for c in cm.get("cells", []):
            val = None
            if c["source"] == "ref":
                val = prefill["par_ref"].get(c["ref"], {}).get("montant")
            elif c["source"] == "note3c":
                val = prefill.get("note_3c", {}).get(c["categorie"], {}).get(c["champ"])
            elif c["source"] == "intake":
                val = prefill["intake"].get(c["cle"])
            if val is None or c["sheet"] not in wb.sheetnames:
                ignores.append(c.get("libelle", "?")); continue
            wb[c["sheet"]][c["cell_n"]] = val
            ecrits += 1
        wb.save(sortie)
        print(f"\n== 5. Peuplement : {ecrits} cellules écrites -> {sortie} ==")
        if ignores:
            print(f"  {len(ignores)} valeur(s) non placées : " + ", ".join(str(x) for x in ignores[:8]))
        print("  Ouvrir dans Excel : les formules de la liasse recalculent Bilan/CR/TFT.")
    else:
        print("\n== 5. Peuplement NON effectué (cell-map absent, propre au modèle) ==")
        print("  Pré-remplissage disponible (postes agrégés + Note 3C + intake).")

    print("\n⚠ Validation chef comptable obligatoire. Annexes fiscales : format GE à confirmer.")


if __name__ == "__main__":
    main()
