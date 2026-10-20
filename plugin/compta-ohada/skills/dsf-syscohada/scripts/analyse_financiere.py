# -*- coding: utf-8 -*-
"""RAPPORT D'ANALYSE FINANCIÈRE PLURIANNUEL (interne, jamais déposé) — N, N-1, N-2… selon les sources présentes.

Années : N (balance), N-1 (ouvertures de la balance pour le bilan ; DEF précédente — Excel ou scannée — pour le
compte de résultat), puis les années antérieures lues sur des DEF scannées (chacune en donne deux : son exercice et
le précédent). Seuls des montants PROUVÉS ou CONFIRMÉS entrent ; un montant manquant s'affiche « — », jamais estimé.

Contenu : synthèse (chiffres clés, tendance), bilan et compte de résultat comparés, soldes intermédiaires de gestion,
ratios, variations significatives EXPLIQUÉES pour guider le personnel :
  • faits chiffrés : les comptes qui portent l'écart (bilan : clôture − ouverture de la balance, compte par compte ;
    compte de résultat : composition de l'exercice N par compte — la balance N-1 n'est pas connue) ;
  • points à vérifier propres au poste (coupure, provisions, inventaire…) ;
  • cause métier : JAMAIS inventée — demandée en question facultative (« analyse.<code> »), reprise ici avec son
    origine et sa date quand elle est donnée.
"""
import datetime
import os

BLEU, GRIS = "1F3864", "F2F2F2"
NUM, PCT = '#,##0;(#,##0);"–"', '0.0%;-0.0%;"–"'
SEUIL_PCT = 0.10                  # variation significative : ≥ 10 % …
SEUIL_PART = 0.01                 # … et ≥ 1 % du total du bilan (postes de bilan) ou du chiffre d'affaires (résultat)

BILAN_ACTIF = ["AD", "AI", "AP", "AQ", "AZ", "BA", "BB", "BG", "BH", "BI", "BJ", "BK", "BQ", "BR", "BS", "BT", "BU", "BZ"]
BILAN_PASSIF = ["CA", "CB", "CD", "CE", "CF", "CG", "CH", "CJ", "CL", "CM", "CP", "DA", "DB", "DC", "DD", "DF",
                "DH", "DI", "DJ", "DK", "DM", "DN", "DP", "DQ", "DR", "DT", "DV", "DZ"]
RESULTAT = ["TA", "RA", "RB", "XA", "TB", "TC", "TD", "XB", "TE", "TF", "TG", "TH", "TI", "RC", "RD", "RE", "RF", "RG",
            "RH", "RI", "RJ", "XC", "RK", "XD", "TJ", "RL", "XE", "TK", "TL", "TM", "RM", "RN", "XF", "XG", "TN", "TO",
            "RO", "RP", "XH", "RQ", "RS", "XI"]
SIG = [("XB", "Chiffre d'affaires"), ("XA", "Marge commerciale"), ("XC", "Valeur ajoutée"),
       ("XD", "Excédent brut d'exploitation"), ("XE", "Résultat d'exploitation"), ("XF", "Résultat financier"),
       ("XG", "Résultat des activités ordinaires"), ("XH", "Résultat hors activités ordinaires"), ("XI", "Résultat net")]

VERIFIER = [  # (codes, points à vérifier) — guides de révision usuels, par nature de poste
    (("BI", "BH", "BJ"), "créances anciennes et provisions pour créances douteuses ; encaissements après la clôture ; "
                         "coupure des ventes"),
    (("BB",), "inventaire physique et sa valorisation ; dépréciation des stocks lents ; coupure des achats"),
    (("AD", "AI", "AJ", "AK", "AL", "AM", "AN", "AP", "AQ", "AR", "AS", "AZ"),
     "acquisitions et cessions de l'exercice (note 3) ; dotations ; mises au rebut ; immobilisations en cours"),
    (("DJ", "DI"), "factures non parvenues ; délais de paiement ; litiges fournisseurs ; avances clients"),
    (("DK",), "déclarations et paiements fiscaux et sociaux après la clôture ; redressements en cours"),
    (("DA", "DB", "DC", "DD"), "nouveaux emprunts et remboursements (tableaux d'amortissement) ; provisions pour risques"),
    (("BS", "BT", "BQ", "BR", "DQ", "DR", "DT"), "rapprochements bancaires à la clôture ; découverts ; placements"),
    (("CA", "CB", "CD", "CE", "CF", "CG", "CH", "CJ", "CL", "CM", "CP", "DF"),
     "affectation du résultat N-1 (procès-verbal) ; dividendes ; augmentation de capital ; subventions"),
    (("TA", "TB", "TC", "TD", "XB"), "effet volume et effet prix ; contrats nouveaux ou perdus ; coupure des ventes"),
    (("RA", "RB", "RC", "RD", "RE", "RF", "XA"), "niveau d'activité ; évolution des prix d'achat ; variation des stocks"),
    (("RG", "RH", "RI", "RJ"), "contrats de services ; charges exceptionnelles enregistrées en exploitation"),
    (("RK",), "effectifs ; augmentations, primes, indemnités ; charges sociales"),
    (("RL", "TJ"), "investissements de l'exercice ; durées d'amortissement ; provisions nouvelles ou reprises"),
    (("RM", "RN", "TK", "TL", "TM", "XF"), "niveau d'endettement et taux ; pertes et gains de change"),
    (("TN", "TO", "RO", "RP", "XH"), "opérations exceptionnelles : cessions d'immobilisations, régularisations"),
    (("RS",), "base imposable, réintégrations, minimum de perception (CMF) ; acomptes"),
]


def a_verifier(ref):
    return next((t for codes, t in VERIFIER if ref in codes), "pièces justificatives des principales écritures du poste")


def ratios(v):
    """Ratios d'une année (dict code -> montant, charges en négatif). None quand une donnée manque ou un dénominateur
    est nul : un ratio ne s'invente pas."""
    g = lambda *cs: None if any(v.get(c) is None for c in cs) else [v[c] for c in cs]   # noqa: E731
    # base négative ou nulle (CA, valeur ajoutée, capitaux propres…) : le ratio n'a pas de sens, il n'est pas calculé
    # (constaté sur un dossier réel : « charges de personnel / valeur ajoutée » = −1,5 avec une valeur ajoutée négative)
    div = lambda a, b: None if a is None or not b or b < 0 else a / b                  # noqa: E731
    out = {}
    x = g("XB", "XD")
    out["Marge d'EBE (EBE / CA)"] = div(x[1], x[0]) if x else None
    x = g("XB", "XI")
    out["Marge nette (résultat / CA)"] = div(x[1], x[0]) if x else None
    x = g("XB", "XC")
    out["Valeur ajoutée / CA"] = div(x[1], x[0]) if x else None
    x = g("XC", "RK")
    out["Charges de personnel / valeur ajoutée"] = div(abs(x[1]), x[0]) if x else None
    x = g("CP", "BZ")
    out["Autonomie financière (capitaux propres / total bilan)"] = div(x[0], x[1]) if x else None
    x = g("DD", "DT", "CP")
    out["Endettement (dettes financières + trésorerie passif) / capitaux propres"] = div(x[0] + x[1], x[2]) if x else None
    x = g("DD", "XD")
    out["Capacité de remboursement (dettes financières / EBE), années"] = div(x[0], x[1]) if x and x[1] > 0 else None
    x = g("BK", "BT", "DP", "DT")
    out["Liquidité générale (actif circulant + trésorerie) / passif à court terme"] = div(x[0] + x[1], x[2] + x[3]) if x else None
    x = g("BT", "DP", "DT")
    out["Liquidité immédiate (trésorerie actif / passif à court terme)"] = div(x[0], x[1] + x[2]) if x else None
    x = g("BI", "XB")
    out["Délai clients (jours de CA)"] = div(x[0] * 360, x[1]) if x else None
    x = g("DJ", "RA", "RC", "RE")
    out["Délai fournisseurs (jours d'achats)"] = div(x[0] * 360, abs(x[1]) + abs(x[2]) + abs(x[3])) if x else None
    x = g("BB", "XB")
    out["Stocks (jours de CA)"] = div(x[0] * 360, x[1]) if x else None
    x = g("XI", "CP")
    out["Rentabilité des capitaux propres (résultat / capitaux propres)"] = div(x[0], x[1]) if x else None
    x = g("XE", "BZ")
    out["Rentabilité économique (résultat d'exploitation / total bilan)"] = div(x[0], x[1]) if x else None
    x = g("DF", "AZ")
    out["Fonds de roulement (ressources stables − actif immobilisé)"] = x[0] - x[1] if x else None
    x = g("BK", "DP")
    out["Besoin en fonds de roulement (actif circulant − passif circulant)"] = x[0] - x[1] if x else None
    x = g("BT", "DT")
    out["Trésorerie nette (trésorerie actif − trésorerie passif)"] = x[0] - x[1] if x else None
    return out


EN_MONTANT = {"Fonds de roulement (ressources stables − actif immobilisé)",
              "Besoin en fonds de roulement (actif circulant − passif circulant)",
              "Trésorerie nette (trésorerie actif − trésorerie passif)"}
EN_JOURS = {"Délai clients (jours de CA)", "Délai fournisseurs (jours d'achats)", "Stocks (jours de CA)",
            "Capacité de remboursement (dettes financières / EBE), années"}


def comptes_par_poste(detail):
    """agg["detail"] -> {poste: {compte: montant net du poste}} (actif : brut − amortissements)."""
    out = {}
    for compte, ref, colonne, v in detail or []:
        d = out.setdefault(ref, {})
        d[compte] = d.get(compte, 0.0) + (-v if colonne == "amort" else v)
    return out


def expliquer(ref, vN, v1, cN, c1, libelles):
    """Faits chiffrés d'une variation : comptes qui la portent (bilan) ou composition de N (résultat)."""
    def nom(c):
        return f"{c} {libelles.get(c, '')}".strip()
    if ref in cN and c1 is not None and ref[0] in "ABCD":
        ecarts = sorted(((c, cN[ref].get(c, 0.0) - (c1.get(ref, {}).get(c, 0.0))) for c in set(cN[ref]) | set(c1.get(ref, {}))),
                        key=lambda x: -abs(x[1]))
        tot = sum(e for _, e in ecarts) or 1
        top = [(c, e) for c, e in ecarts if abs(e) >= 1][:3]
        if top:
            return "écart porté par : " + " ; ".join(f"{nom(c)} {e:+,.0f} ({e / tot:.0%} de l'écart)".replace(",", " ")
                                                     for c, e in top)
    if ref in cN:
        comp = sorted(cN[ref].items(), key=lambda x: -abs(x[1]))
        tot = sum(abs(x) for _, x in comp) or 1
        top = [(c, x) for c, x in comp if abs(x) >= 1][:3]
        if top:
            return ("composition en N : "                 # pourquoi pas l'écart par compte : dit une fois, aux sources
                    + " ; ".join(f"{nom(c)} {abs(x):,.0f} ({abs(x) / tot:.0%})".replace(",", " ") for c, x in top))
    return "détail par compte non disponible pour ce poste"


def variations(annees, postes, cN, c1, libelles, causes):
    """Variations significatives N / N-1, expliquées. annees : {an: {"v": {code: montant}, ...}} trié."""
    ans = sorted(annees)
    if len(ans) < 2:
        return []
    N, P = ans[-1], ans[-2]
    vN, vP = annees[N]["v"], annees[P]["v"]
    base_b = abs(vN.get("BZ") or 0) or 1
    base_r = abs(vN.get("XB") or 0) or 1
    out = []
    for ref in BILAN_ACTIF + BILAN_PASSIF + RESULTAT:
        # les TOTAUX (et CJ, doublon du résultat) se lisent en synthèse : on explique les lignes qui les composent
        if "total" in postes.get(ref, {}) or ref == "CJ":
            continue
        a, b = vN.get(ref), vP.get(ref)
        if a is None or b is None:
            continue
        d = a - b
        base = base_b if ref[0] in "ABCD" else base_r
        pct = d / abs(b) if b else (1.0 if d else 0.0)
        if abs(d) >= SEUIL_PART * base and abs(pct) >= SEUIL_PCT:
            out.append({"ref": ref, "libelle": postes.get(ref, {}).get("libelle", ref), "N-1": b, "N": a, "ecart": d,
                        "pct": pct if b else None, "faits": expliquer(ref, a, b, cN, c1, libelles),
                        "verifier": a_verifier(ref), "cause": causes.get(ref)})
    out.sort(key=lambda x: -abs(x["ecart"]) / (base_b if x["ref"][0] in "ABCD" else base_r))
    return out


def construire(D, annees, sortie_xlsx, sortie_pdf=None, causes=None, entite=""):
    """D : calcul de la DEF (def_ge) ; annees : {an: {"v": {code: montant}, "source": texte}}. Rend le résumé."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    causes = causes or {}
    postes, libelles = D.get("postes", {}), D.get("libelles", {})
    ans = sorted(annees)
    cN = comptes_par_poste((D.get("agg") or {}).get("detail"))
    c1 = comptes_par_poste((D.get("agg1") or {}).get("detail")) if D.get("agg1") else None
    var = variations(annees, postes, cN, c1, libelles, {k: v for k, v in causes.items()})
    R = {an: ratios(annees[an]["v"]) for an in ans}
    wb = Workbook()
    feuilles = []
    fill = lambda c: PatternFill("solid", fgColor=c)                                        # noqa: E731

    def page(titre, largeurs):
        ws = wb.active if not feuilles else wb.create_sheet()
        ws.title = titre[:31]
        feuilles.append(ws.title)
        ws.append([f"{entite} — analyse financière {ans[0]}–{ans[-1]}" if len(ans) > 1 else f"{entite} — analyse financière"])
        ws.cell(1, 1).font = Font(name="Calibri", size=10, bold=True, color=BLEU)
        ws.append(["Document interne — ne pas joindre au dépôt"])
        ws.cell(2, 1).font = Font(name="Calibri", size=9, italic=True, color="7F7F7F")
        ws.append([titre])
        ws.cell(3, 1).font = Font(name="Calibri", size=15, bold=True, color=BLEU)
        ws.append([])
        for j, w in enumerate(largeurs, 1):
            ws.column_dimensions[ws.cell(1, j).column_letter].width = w
        ws.page_setup.paperSize, ws.page_setup.orientation = ws.PAPERSIZE_A4, "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddFooter.center.text = "Page &P / &N"
        ws.sheet_view.showGridLines = False
        return ws

    def entetes(ws, cols):
        ws.append(cols)
        for j in range(1, len(cols) + 1):
            c = ws.cell(ws.max_row, j)
            c.font, c.fill = Font(name="Calibri", size=9, bold=True, color="FFFFFF"), fill(BLEU)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    def ligne(ws, vals, formats=None, gras=False):
        ws.append(["—" if v is None else v for v in vals])
        r = ws.max_row
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j)
            c.font = Font(name="Calibri", size=9, bold=gras)
            c.border = Border(bottom=Side(style="hair", color="BFBFBF"))
            c.alignment = Alignment(wrap_text=True, vertical="top",
                                    horizontal="right" if isinstance(v, (int, float)) or v is None else None)
            if isinstance(v, (int, float)):
                c.number_format = (formats or {}).get(j, NUM)

    def var_n(v_n, v_p):
        if v_n is None or v_p is None:
            return None, None
        return v_n - v_p, ((v_n - v_p) / abs(v_p) if v_p else None)

    # ---------------- synthèse
    ws = page("Synthèse", [46] + [16] * len(ans) + [16, 12])
    entetes(ws, ["Chiffre clé"] + [str(a) for a in ans] + ["Variation N/N-1", "%"])
    cles = [("XB", "Chiffre d'affaires"), ("XC", "Valeur ajoutée"), ("XD", "Excédent brut d'exploitation"),
            ("XE", "Résultat d'exploitation"), ("XI", "Résultat net"), ("BZ", "Total du bilan"),
            ("CP", "Capitaux propres"), ("DD", "Dettes financières")]
    for code, lib in cles:
        vals = [annees[a]["v"].get(code) for a in ans]
        d, p = var_n(vals[-1], vals[-2]) if len(vals) > 1 else (None, None)
        ligne(ws, [lib] + vals + [d, p], {len(ans) + 3: PCT})
    tn = [ratios(annees[a]["v"]).get("Trésorerie nette (trésorerie actif − trésorerie passif)") for a in ans]
    d, p = var_n(tn[-1], tn[-2]) if len(tn) > 1 else (None, None)
    ligne(ws, ["Trésorerie nette"] + tn + [d, p], {len(ans) + 3: PCT})
    ws.append([])
    ws.append([f"{len(var)} variation(s) significative(s) entre {ans[-2]} et {ans[-1]} (≥ {SEUIL_PCT:.0%} et ≥ "
               f"{SEUIL_PART:.0%} du total du bilan ou du chiffre d'affaires) : voir « Variations expliquées »."
               if len(ans) > 1 else "Une seule année disponible : pas de comparaison."])
    ws.append([f"Généré le {datetime.datetime.now().strftime('%d/%m/%Y %H:%M')} — à relire avec le comptable ; "
               "les causes des variations sont à documenter par l'entreprise."])

    # ---------------- états comparés
    for titre, refs in (("Bilan comparé", BILAN_ACTIF + BILAN_PASSIF), ("Compte de résultat comparé", RESULTAT)):
        ws = page(titre, [8, 52] + [16] * len(ans) + [16, 10])
        entetes(ws, ["Code", "Poste"] + [str(a) for a in ans] + ["Var. N/N-1", "%"])
        for ref in refs:
            vals = [annees[a]["v"].get(ref) for a in ans]
            if all(v in (None, 0) for v in vals):
                continue
            d, p = var_n(vals[-1], vals[-2]) if len(vals) > 1 else (None, None)
            ligne(ws, [ref, postes.get(ref, {}).get("libelle", ref)] + vals + [d, p], {len(ans) + 4: PCT},
                  gras="total" in postes.get(ref, {}))

    # ---------------- SIG et ratios
    ws = page("Soldes et ratios", [62] + [16] * len(ans))
    entetes(ws, ["Soldes intermédiaires de gestion"] + [str(a) for a in ans])
    for code, lib in SIG:
        ligne(ws, [lib] + [annees[a]["v"].get(code) for a in ans])
    ws.append([])
    entetes(ws, ["Ratio"] + [str(a) for a in ans])
    for nom in R[ans[-1]]:
        fmt = NUM if nom in EN_MONTANT else "0.0" if nom in EN_JOURS else PCT
        ligne(ws, [nom] + [R[a][nom] for a in ans], {j: fmt for j in range(2, len(ans) + 2)})

    # ---------------- variations expliquées
    ws = page("Variations expliquées", [8, 30, 15, 15, 15, 9, 60, 46, 40])
    entetes(ws, ["Code", "Poste", f"{ans[-2]}" if len(ans) > 1 else "N-1", f"{ans[-1]}", "Écart", "%",
                 "Faits chiffrés (comptes)", "Points à vérifier", "Cause (entreprise)"])
    for x in var:
        cause = x["cause"]
        ligne(ws, [x["ref"], x["libelle"], x["N-1"], x["N"], x["ecart"], x["pct"], x["faits"], x["verifier"],
                   (f"{cause['texte']} — {cause['origine']}" if cause else "à documenter (question facultative du formulaire)")],
              {6: PCT})
    if not var:
        ligne(ws, ["", "Aucune variation significative" if len(ans) > 1 else "Une seule année : pas de variation"])

    # ---------------- sources
    ws = page("Sources et méthode", [16, 110])
    entetes(ws, ["Année", "Source des montants"])
    for a in ans:
        ligne(ws, [str(a), annees[a]["source"]])
    ws.append([])
    for t in ("Aucun montant n'est estimé : une donnée absente s'affiche « — » et le ratio qui en dépend n'est pas calculé.",
              "Montants lus sur une DEF scannée : seuls ceux PROUVÉS par les formules SYSCOHADA, ou confirmés par "
              "l'utilisateur, sont repris.",
              f"Variation significative : au moins {SEUIL_PCT:.0%} et au moins {SEUIL_PART:.0%} du total du bilan "
              "(postes de bilan) ou du chiffre d'affaires (postes de résultat).",
              "Faits chiffrés : comptes qui portent l'écart (bilan : clôture − ouverture de la balance) ; pour le résultat, "
              "composition de l'exercice — l'écart par compte demande la balance de l'exercice précédent.",
              "Délais en jours sur 360 jours, montants hors taxes : ordres de grandeur, à rapprocher des balances âgées.",
              "Les causes ne sont jamais déduites : elles sont demandées à l'entreprise (questions facultatives)."):
        ligne(ws, ["", t])
    wb.save(sortie_xlsx)
    pdf_ok = False
    if sortie_pdf:
        import def_ge
        pdf_ok = bool(def_ge.exporter_pdf(sortie_xlsx, sortie_pdf, feuilles))
    return {"annees": ans, "variations": var, "ratios": R, "pdf": pdf_ok}


# ------------------------------------------------------------------------------------------------- années
def annees_depuis(D, N, scans, n1_source):
    """{an: {"v", "source"}} : N et N-1 du calcul, puis les années antérieures des DEF scannées lues (montants
    prouvés seulement ; charges en négatif comme dans le calcul)."""
    out = {N: {"v": dict(D["pN"]), "source": "balance de l'exercice (calcul de la DEF)"}}
    if D.get("p1"):
        out[N - 1] = {"v": dict(D["p1"]), "source": n1_source or D.get("src_n1", "")}

    def norme(code, v):
        return -abs(v) if code[0] == "R" and code not in ("RB", "RD", "RF") else v
    # une DEF scannée de l'exercice X donne X (sa colonne N) et X-1 (sa colonne N-1) ; la colonne N d'une DEF prime
    for priorite in ("n", "n1"):
        for fichier, r in scans:
            ex = r.get("exercice")
            if not ex:
                continue
            an = ex if priorite == "n" else ex - 1
            if an in out or an > N:                       # N et N-1 viennent du calcul ; une année déjà tenue reste
                continue
            vals = {c: norme(c, cols[priorite]["valeur"]) for c, cols in r.get("etats", {}).items()
                    if priorite in cols and cols[priorite]["prouve"] and cols[priorite]["valeur"] is not None}
            if not vals:
                continue
            out[an] = {"v": vals, "source": f"DEF {ex} scannée « {fichier} », colonne {'N' if priorite == 'n' else 'N-1'} "
                                            f"(montants prouvés par les formules : {len(vals)})"}
    return out
