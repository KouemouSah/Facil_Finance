# -*- coding: utf-8 -*-
"""DEF « modèle officiel » : reproduit feuille pour feuille une DEF existante (structure GE déposée : 51 feuilles,
mêmes lignes, mêmes colonnes, même mise en forme), en français ET en espagnol, montants RECALCULÉS.

Le modèle fournit la STRUCTURE — jamais ses montants de l'exercice (il ne contient d'ailleurs aucune formule :
ses valeurs sont collées). Chaque cellule numérique est :
  - recalculée : postes du bilan / compte de résultat / TFT (moteur def_ge, IS intégré, TFT réconcilié, note 3),
    lignes de notes « REF + liste de comptes » (ex. « BI 411-4112-4116+413 ») depuis la balance,
    TF1-TF5 (moteurs fiscal_ge / tf345_ge), fiche de synthèse (note 34) depuis les soldes intermédiaires ;
  - ou reprise de la colonne N-1 du modèle (--colonne-precedente N-1 : modèle du même exercice ;
    N : modèle de l'exercice précédent, sa colonne N devient notre N-1) ;
  - ou mise en FORMULE (totaux de notes, variations, échéances) ;
  - ou laissée VIDE (fond ambre) quand elle n'est pas calculable : identité, effectifs, données extra-comptables.
L'identité de l'entité se saisit UNE fois en page de garde ; tous les en-têtes la reprennent par formule.
Le brouillon hors zone d'impression du modèle est écarté ; la zone d'impression et la mise en page sont posées :
imprimer ou exporter en PDF depuis Excel donne le même document que le PDF produit ici.
Contrôles et couverture : classeur SÉPARÉ (interne, non déposé).

Usage :
  python def_modele.py --modele DEF.xlsx --balance B.xls --exercice 2025 [--grand-livre GL.xls]
       [--profil p.json] [--decisions d.json] [--colonne-precedente N|N-1] [--logo logo.png]
       [--langues fr,es] --sortie prefixe [--pdf]
"""
import argparse
import datetime
import json
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from copy import copy

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import agreger_balance_dsf as agr  # noqa: E402
import def_ge  # noqa: E402

I18N = os.path.join(HERE, "..", "references", "i18n-def.json")
NUM = '#,##0;(#,##0);"–"'
AMBRE = "FFF2CC"
BLEU = "1F3864"
# feuilles dont les chiffres sont extra-comptables : vides, à saisir (fond ambre)
ETATS = ("bilan actif", "bilan passif", "bilan paysage", "compte resultat")
EXTRA_COMPTABLES = ("fiche r1", "fiche r2", "fiche r3", "note 27b", "note 2", "note 35", "note 36", "grille d'analyse des notes")


def norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"\s+", " ", s).strip()


def num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


# --------------------------------------------------------------------------- listes de comptes du modèle
def expr_comptes(txt):
    """« BI 411-4112-4116+413 » -> ("BI", [(+1,"411",False), (-1,"4112",False), (-1,"4116",False), (+1,"413",False)]) ;
    « DA 181 à 183 » -> plage 181, 182, 183 ; « 28181p » -> partiel (True). None si ce n'est pas une liste de comptes."""
    if not isinstance(txt, str):
        return None
    m = re.match(r"^\s*([A-Z]{2})\s+(\d.*)$", txt.strip())
    if not m:
        return None
    ref, reste = m.group(1), m.group(2)
    termes = []
    for sgn, a, p, b in re.findall(r"([+\-,;]?)\s*(\d+)(p?)(?:\s*(?:à|a)\s*(\d+))?", reste):
        s = -1 if sgn == "-" else 1
        if b:
            n = len(a)
            for k in range(int(a), int(b) + 1):
                termes.append((s, str(k).zfill(n), False))
        else:
            termes.append((s, a, bool(p)))
    return (ref, termes) if termes else None


def sens(ref, termes):
    """Présentation positive selon la nature : actif / charges = D - C ; passif / produits / amortissements
    et dépréciations (28, 29, 39, 49, 59) = C - D."""
    if termes and termes[0][1][:2] in ("28", "29", "39", "49", "59"):
        return -1
    return 1 if ref[0] in "ABR" else -1


def valeur_expr(ref, termes, soldes):
    """Comptes de tiers et de trésorerie (classes 4 et 5) éclatés PAR SENS, comme au bilan : une ligne d'actif (A, B)
    ne retient que les soldes débiteurs, une ligne de passif (C, D) que les soldes créditeurs."""
    def garde(a, v):
        if a[:1] in ("4", "5") and ref[0] in "AB":
            return max(v, 0.0)
        if a[:1] in ("4", "5") and ref[0] in "CD":
            return min(v, 0.0)
        return v
    tot = 0.0
    for sg, pre, _ in termes:
        tot += sg * sum(garde(a, v) for a, v in soldes.items() if a.startswith(pre))
    return round(sens(ref, termes) * tot, 2)


def valeur_detail(ref, termes, detail):
    """Ligne de note calculée sur l'AFFECTATION du moteur (compte -> poste, sens compris) : seuls les comptes que le
    bilan / compte de résultat a rattachés à CE poste comptent. La note retombe donc exactement sur son poste.
    detail : [(compte, poste, colonne, valeur présentée)] ; charges (R) présentées en positif dans les notes."""
    tot = 0.0
    for sg, pre, _ in termes:
        tot += sg * sum(v for a, rf, _c, v in detail if rf == ref and a.startswith(pre))
    return round(-tot if ref[0] == "R" else tot, 2)


def est_ref(v):
    return isinstance(v, str) and len(v.strip()) == 2 and v.strip().isalpha() and v.strip().isupper()


def ligne_total(lbl):
    return isinstance(lbl, str) and bool(re.match(r"^(sous[- ]?total|total|totaux)\b", norm(lbl)))


# --------------------------------------------------------------------------- rôles des colonnes
def roles_colonnes(ws, h0, h1, exercice):
    """Rôle de chaque colonne d'après l'en-tête (lignes h0..h1), cellules fusionnées propagées."""
    fusions = {}
    for rg in ws.merged_cells.ranges:
        v = ws.cell(rg.min_row, rg.min_col).value
        for r in range(rg.min_row, rg.max_row + 1):
            for c in range(rg.min_col, rg.max_col + 1):
                fusions[(r, c)] = v
    txt, annees = defaultdict(list), {}
    for r in range(h0, h1 + 1):
        for c in range(1, ws.max_column + 1):
            v = fusions.get((r, c), ws.cell(r, c).value)
            if isinstance(v, (datetime.date, datetime.datetime)):
                annees[c] = v.year
            elif isinstance(v, str) and v.strip():
                txt[c].append(norm(v))
                y = re.search(r"(?:^|\D)(20\d{2})(?:\D|$)", v)
                if y:
                    annees[c] = int(y.group(1))
    roles, nets = {}, []
    for c, parts in txt.items():
        t = " | ".join(parts)
        if "brut" in t and "amort" not in t:
            roles[c] = "brut"
        elif "amort" in t:
            roles[c] = "amort"
        elif re.search(r"\bnet\b", t):
            nets.append(c)
        elif "variation en %" in t or "pourcentage" in t:
            roles[c] = "var_pct"
        elif "variation" in t:
            roles[c] = "var_abs"
        elif "un an au plus" in t:
            roles[c] = "ech1"
        elif "plus de deux ans" in t:
            roles[c] = "ech3"
        elif "plus d'un an" in t or "plus d un an" in t:
            roles[c] = "ech2"
        elif re.search(r"(^|\W)n\s*-\s*1($|\W)", t):
            roles[c] = "n1"
        elif re.search(r"(^|\W)n\s*-\s*[2-9]($|\W)", t):
            roles[c] = "hist"                             # N-2, N-3… : historique repris du modèle
        elif re.search(r"(^|\W)n($|\W)", t):
            roles[c] = "n"
    for c, y in annees.items():                          # colonnes datées (bilan, compte de résultat, TFT)
        if c not in roles and y in (exercice, exercice - 1):
            roles[c] = "n" if y == exercice else "n1"
    for i, c in enumerate(sorted(nets)):
        y = annees.get(c)
        roles[c] = "net_n1" if (y == exercice - 1 or (y is None and i % 2 == 1)) else "net_n"
    return roles


# --------------------------------------------------------------------------- génération
class Generateur:
    def __init__(self, a):
        import openpyxl
        from detection_balance import analyse
        self.a = a
        self.modele = openpyxl.load_workbook(a.modele, data_only=True)
        # MODE GABARIT : le formulaire officiel vierge embarqué (aucune DEF déposée). Ses montants sont à zéro : la
        # structure des totaux vient du fichier annexe (prouvée une fois sur un dossier réel), l'année du gabarit est
        # décalée vers l'exercice, la colonne N-1 des notes reste vide, le bilan N-1 vient de l'ouverture de la balance.
        self.gabarit = bool(getattr(a, "gabarit", False))
        self.structure = json.load(open(a.structure, encoding="utf-8")) if getattr(a, "structure", None) else {}
        self.annee_modele = self.structure.get("annee") if self.gabarit else None
        self.trace = {"totaux": {}, "refs": {}}             # structure constatée : de quoi construire le fichier annexe
        ns = argparse.Namespace(balance=a.balance, exercice=a.exercice, profil=a.profil, decisions=a.decisions,
                                grand_livre=a.grand_livre, tiers=None, def_precedente=None if self.gabarit else a.modele,
                                colonne_precedente=a.colonne_precedente, gabarit=None,
                                n1_externe=getattr(a, "n1_externe", None))
        self.D = def_ge.calculer(ns)
        # profil relu ici pour les COMPLÉMENTS saisis par l'entreprise (NIF des tiers…) : une seule source de vérité,
        # réinjectée à chaque génération dans le classeur ET dans les PDF
        self.profil = json.load(open(a.profil, encoding="utf-8")) if a.profil else {}
        self.det = analyse(a.balance)
        bal, _ = agr.avec_impot(sorted(self.det["soldes"].items()), self.D["fisc"]["impot_du"])
        self.soldes = dict(bal)
        # affectation compte -> poste du moteur ; résultat antérieur non affecté (13, balance non clôturée) présenté
        # en report à nouveau comme au bilan : il rejoint la ligne « CH 121,129 » de la note 14
        rk = (lambda a: "129~" + a if a.startswith("13") else a) if not self.D["agg"]["controles"].get("balance_cloturee") else (lambda a: a)
        self.detail = [(rk(a), rf, col, v) for a, rf, col, v in self.D["agg"]["detail"]]
        self.i18n = json.load(open(I18N, encoding="utf-8"))
        # identité DU MODÈLE : reconnue pour être remplacée par des formules, jamais recopiée
        g = self._feuille("page garde")
        self.ident_modele = {}
        if g is not None:
            for cell in ("D13", "D16", "D17", "D18"):
                if isinstance(g[cell].value, str) and g[cell].value.strip():
                    self.ident_modele[norm(g[cell].value)] = cell
        self.controles, self.a_completer, self.couverture, self._refs_devinees = [], [], Counter(), {}
        self.notes_omises, self.controles_elagage = {}, []     # classeur adapté aux données (par langue)
        self.tf5_lignes = self.tf5_hybride()
        self.notes_absentes = self.notes_hors_formulaire()
        self.controles_init = list(self.controles)

    def notes_hors_formulaire(self):
        """Postes CHIFFRÉS (N ou N-1) dont la note n'a pas de feuille dans le formulaire utilisé (gabarit ou DEF
        déposée) : la note est applicable mais ne peut pas être produite — dit, jamais tu. {note: [(ref, libellé)]}."""
        dispo = {re.match(r"note\s*(\d+)", norm(w.title)).group(1) for w in self.modele.worksheets
                 if re.match(r"note\s*\d+", norm(w.title))}
        manque = defaultdict(list)
        for ref, p in self.D["postes"].items():
            n = str(p.get("note") or "").strip()
            m = re.match(r"(\d+)", n)
            if not m or m.group(1) in dispo or p.get("total"):
                continue
            if abs(self.D["pN"].get(ref) or 0) >= 1 or abs(self.D["p1"].get(ref) or 0) >= 1:
                manque[m.group(1)].append((ref, p.get("libelle", "")))
        for n, postes in sorted(manque.items(), key=lambda x: int(x[0])):
            self.controles.append((f"Note {n} applicable (postes chiffrés) mais ABSENTE du formulaire utilisé", False,
                                   "; ".join(f"{r} {l}" for r, l in postes) + " — montants présents au bilan / compte de "
                                   "résultat ; la note elle-même n'est pas produite (formulaire à compléter)"))
        return dict(manque)

    def _feuille(self, nom):
        return next((w for w in self.modele.worksheets if norm(w.title) == nom), None)

    # ---------------------------------------------------------------- TF5 hybride
    def tf5_hybride(self):
        """Référentiel = TF5 du modèle (nom, adresse, nature, NIF, montant déclaré). Chaque tiers est rapproché d'un
        code tiers du grand livre quand un mot distinctif de son nom apparaît dans les libellés de ce code.
        Modèle du même exercice : montants DÉCLARÉS conservés, rapprochement fourni en contrôle.
        Modèle antérieur : montants du grand livre ; codes du grand livre sans identité ajoutés, à compléter."""
        t = self.D.get("tf345")
        gl_rows = (t or {}).get("tf5", {}).get("lignes", [])
        ref = []
        ws = self._feuille("tf5")
        if ws is not None:
            for r in range(8, ws.max_row + 1):
                nom = ws.cell(r, 1).value
                if isinstance(nom, str) and nom.strip() and not norm(nom).startswith(("nombre", "total")):
                    ref.append({"nom": re.sub(r"\s+", " ", nom.strip()), "adresse": ws.cell(r, 3).value or "",
                                "nature": ws.cell(r, 7).value or "", "nif": ws.cell(r, 11).value or ws.cell(r, 10).value or "",
                                "montant_modele": ws.cell(r, 12).value})
        libs = defaultdict(list)
        if self.a.grand_livre:
            import grand_livre
            gl = grand_livre.lire(self.a.grand_livre)
            cl = grand_livre.pieces_de_cloture(gl)
            for x in gl["lignes"]:
                if x["tiers"] and x["piece"] not in cl:
                    libs[x["tiers"]].append(norm(x["libelle"] or "").upper())
        STOP = {"SERVICES", "SERVICIOS", "GROUP", "GUINEA", "GUINEE", "INDUSTRIAL", "ASSURANCES"}
        par_code = {r["code"]: r for r in gl_rows}
        meme_exercice = self.a.colonne_precedente == "N-1" and not self.gabarit
        lignes, pris = [], set()
        for e in ref:
            toks = [w for w in re.findall(r"[A-Z0-9]{4,}", norm(e["nom"]).upper()) if w not in STOP]
            score = Counter()
            for code, ls in libs.items():
                score[code] = sum(1 for l in ls if any(re.search(r"\b" + re.escape(w), l) for w in toks))
            code = next((c for c, n in score.most_common() if n and c in par_code and c not in pris), None)
            if code:
                pris.add(code)
            g = par_code.get(code)
            if meme_exercice and num(e["montant_modele"]):
                ht, src = e["montant_modele"], "déclaré"
            elif g:
                ht, src = g["ht"], "grand livre"
            else:
                ht, src = None, "à compléter"
            lignes.append({**e, "code": code or "", "ht": ht, "source": src, "ht_gl": g["ht"] if g else None})
            if code:
                det = f"code {code} ({score[code]} libellé(s) portent le nom) — grand livre {g['ht']:,.0f} / retenu " + \
                      (f"{ht:,.0f} ({src})" if num(ht) else "aucun")
            else:
                det = "aucun code tiers du grand livre ne porte ce nom — montant retenu : " + (f"{ht:,.0f} ({src})" if num(ht) else "aucun")
            self.controles.append((f"TF5 {e['nom']} : rapprochement grand livre", bool(code), det))
        nif_saisis = {norm(k): v for k, v in self.profil.get("tiers_nif", {}).items() if v}
        for x in lignes:                                        # NIF complété par l'entreprise : prioritaire
            if nif_saisis.get(norm(x["nom"])):
                x["nif"] = nif_saisis[norm(x["nom"])]
        if not meme_exercice:
            for g in gl_rows:
                if g["code"] not in pris:
                    lignes.append({"nom": "", "adresse": "", "nature": g["nature"], "nif": "", "code": g["code"], "ht": g["ht"],
                                   "source": "grand livre — identité à compléter", "ht_gl": g["ht"]})
        n_id = sum(1 for x in lignes if x["nom"])
        n_nif = sum(1 for x in lignes if x["nif"])
        self.controles.append((f"TF5 : {n_id} bénéficiaires identifiés, {n_nif} avec NIF, sur {len(lignes)}", n_nif == len(lignes),
                               f"{sum(1 for x in lignes if x['code'])} rapprochés d'un code tiers du grand livre ; NIF manquants : "
                               f"{', '.join(x['nom'] for x in lignes if x['nom'] and not x['nif']) or 'aucun'} (art. 27.2)"))
        return lignes

    # ---------------------------------------------------------------- traduction
    def tr(self, s, lang):
        if not isinstance(s, str) or not s.strip():
            return s
        return self.i18n.get(lang, {}).get(s.strip(), s)

    def nom_feuille(self, titre, lang):
        base = "Balance generale" if norm(titre).startswith("bg ") else titre
        nom = self.i18n.get("feuilles_" + lang, {}).get(base, base)
        if lang == "es":
            nom = re.sub(r"^note\s", "nota ", nom)             # « note 11 » -> « nota 11 »
        return nom[:31]

    # ---------------------------------------------------------------- classeur
    def rendre(self, lang, sortie):
        from openpyxl import Workbook
        wb = Workbook()
        wb.remove(wb.active)
        self.noms = {m.title: self.nom_feuille(m.title, lang) for m in self.modele.worksheets}
        g = self.noms.get(next((t for t in self.noms if norm(t) == "page garde"), ""), "page garde")
        self.pg = f"'{g}'"
        self.a_completer, self.couverture, self.meta = [], Counter(), {}
        self.controles = list(self.controles_init)          # contrôles de rendu : une fois par classeur, pas cumulés
        # même liste de feuilles que le modèle ; la balance (annexe, art. 79.2) passe EN FIN de classeur pour que
        # l'impression s'ouvre sur la page de garde
        ordre = sorted(self.modele.worksheets, key=lambda m: norm(m.title).startswith("bg "))
        for mws in ordre:
            ws = wb.create_sheet(self.noms[mws.title])
            if norm(mws.title).startswith("bg "):
                self.feuille_balance(ws, lang)
            else:
                self.feuille(mws, ws, lang)
        wb.active = 1 if len(wb.worksheets) > 1 else 0     # ouverture sur la page de garde
        wb.save(sortie)
        if self.a.pro:
            self.habiller_pro(wb, lang)
            wb.save(sortie.replace("modele-officiel", "modele-pro") if "modele-officiel" in sortie else sortie[:-5] + "_pro.xlsx")
        return [w.title for w in wb.worksheets]

    def habiller_pro(self, wb, lang):
        """MODÈLE PROFESSIONNEL : mêmes feuilles, mêmes cellules, mêmes formules que le modèle officiel (aucun second
        calcul, donc aucune divergence possible), présentation normalisée :
          - une seule typographie, en-têtes de tableau sur bandeau, totaux sur fond clair, filets fins ;
          - lignes de détail SANS MONTANT (N et N-1 nuls) masquées dans les notes — masquées, pas supprimées : les
            totaux les comptent toujours et « Afficher » les restitue ; une note entièrement nulle porte « Néant » ;
          - en-tête de tableau figé à l'écran et répété à l'impression.
        États (bilan, compte de résultat, TFT) et tableaux fiscaux gardent toutes leurs lignes : leur forme est normée."""
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        fin, trait = Side(style="thin", color="D9D9D9"), Side(style="thin", color=BLEU)
        neant = {"fr": "Néant — aucun montant sur l'exercice ni sur le précédent", "es": "Sin importe en el ejercicio ni en el anterior"}[lang]
        for ws in wb.worksheets:
            m = self.meta.get(ws.title)
            if not m or m["t"] in ("page garde", "french version", "sommaire", "statistiques et syntheses") or m["t"] in EXTRA_COMPTABLES:
                continue
            h0, h1, c0, c1, r1 = m["h0"], m["h1"], m["c0"], m["c1"], m["r1"]
            chiffres = [c for c, ro in m["roles"].items() if ro in ("n", "n1", "brut", "amort", "net_n", "net_n1") and c <= c1]
            for r in range(4, r1 + 1):
                for c in range(c0, c1 + 1):
                    cell = ws.cell(r, c)
                    if type(cell).__name__ == "MergedCell":
                        continue
                    gras = bool(cell.font and cell.font.bold)
                    if h0 <= r <= h1 and chiffres:
                        cell.font = Font(name="Calibri", size=8.5, bold=True, color="FFFFFF")
                        cell.fill = PatternFill("solid", fgColor=BLEU)
                        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                        cell.border = Border(bottom=Side(style="thin", color="FFFFFF"))
                    elif r in m["totaux"]:
                        cell.font = Font(name="Calibri", size=9, bold=True, color=BLEU)
                        cell.fill = PatternFill("solid", fgColor="E8EEF8")
                        cell.border = Border(top=trait, bottom=trait)
                    elif r > h1:
                        cell.font = Font(name="Calibri", size=9, bold=gras, color="262626")
                        cell.fill = PatternFill(fill_type=None)
                        cell.border = Border(bottom=fin)
            if r1 >= 4:
                for c in range(c0, c1 + 1):
                    v = ws.cell(4, c).value
                    if isinstance(v, str) and len(v) > 12 and type(ws.cell(4, c)).__name__ != "MergedCell":
                        ws.cell(4, c).font = Font(name="Calibri", size=12, bold=True, color=BLEU)
            if chiffres and h1 >= h0:
                ws.freeze_panes = ws.cell(h1 + 1, 1)
                ws.print_title_rows = f"{h0}:{h1}"
            if m["t"] in ETATS + ("flux tresorerie", "tf1", "tf2", "tf3", "tf4", "tf5", "reintegraciones") or not chiffres:
                continue
            masquees = 0
            for r in m["lignes"]:
                vals = [ws.cell(r, c).value for c in chiffres]
                if all(v is None or (num(v) and abs(v) < 0.5) for v in vals) and any(num(v) for v in vals):
                    ws.row_dimensions[r].hidden = True
                    masquees += 1
            if masquees and masquees == len([r for r in m["lignes"] if any(num(ws.cell(r, c).value) for c in chiffres)]):
                cell = ws.cell(h1 + 1, c0 + 1 if c0 + 1 <= c1 else c0)
                if type(cell).__name__ != "MergedCell" and not cell.value:
                    cell.value = neant
                    cell.font = Font(name="Calibri", size=9, italic=True, color="7F7F7F")

    def zone(self, mws):
        """Zone d'impression du modèle (union des plages) ; à défaut, la zone utilisée."""
        from openpyxl.utils import range_boundaries
        pa = mws.print_area
        if pa:
            parts = pa if isinstance(pa, (list, tuple)) else str(pa).split(",")
            bs = [range_boundaries(x.split("!")[-1].replace("$", "")) for x in parts]
            return min(b[0] for b in bs), min(b[1] for b in bs), max(b[2] for b in bs), max(b[3] for b in bs), \
                [x.split("!")[-1].replace("$", "") for x in parts]
        return 1, 1, mws.max_column, mws.max_row, None

    def entetes(self, mws, r0, r1):
        """(1re, dernière) ligne d'en-tête du tableau : lignes au-dessus de la première ligne de données."""
        for r in range(r0 + 3, min(r1, r0 + 40) + 1):
            if any(expr_comptes(mws.cell(r, c).value) or est_ref(mws.cell(r, c).value) for c in range(1, min(mws.max_column, 6) + 1)):
                return r0 + 4, r - 1          # la ligne 4 (titre, date du titre) n'est pas un en-tête de colonne
        return r0 + 4, r0 + 5

    def feuille(self, mws, ws, lang):
        t = norm(mws.title)
        c0, r0, c1, r1, pa = self.zone(mws)
        h0, h1 = self.entetes(mws, r0, r1)
        roles = roles_colonnes(mws, h0, h1, self.annee_modele or self.a.exercice)
        ref_cols = [c for c in range(1, c1 + 1) if any(norm(mws.cell(r, c).value or "") == "ref" for r in range(h0, h1 + 1))]
        # gabarit : les actionnaires (note 13) sont propres à l'entité -> à saisir, comme les fiches R
        extra = t in EXTRA_COMPTABLES or (self.gabarit and t == "note 13")
        for c in range(1, c1 + 1):
            L = self.L(c)
            if mws.column_dimensions[L].width:
                ws.column_dimensions[L].width = mws.column_dimensions[L].width
            ws.column_dimensions[L].hidden = mws.column_dimensions[L].hidden
        for r in range(1, r1 + 1):
            if mws.row_dimensions[r].height:
                ws.row_dimensions[r].height = mws.row_dimensions[r].height
        for rg in mws.merged_cells.ranges:
            if rg.max_col <= c1 and rg.max_row <= r1:
                ws.merge_cells(str(rg))
        data_rows, libres, lettre = [], [], {}
        self.meta[ws.title] = {"t": t, "h0": h0, "h1": h1, "c0": c0, "c1": c1, "r1": r1, "roles": roles, "totaux": [], "lignes": []}
        for r in range(1, r1 + 1):
            ex = expr_comptes(mws.cell(r, 1).value)
            ref = next((mws.cell(r, c).value.strip() for c in ref_cols if est_ref(mws.cell(r, c).value)), None)
            for c in range(1, c1 + 1):
                m, out = mws.cell(r, c), ws.cell(r, c)
                if m.has_style:
                    out.font, out.fill, out.border = copy(m.font), copy(m.fill), copy(m.border)
                    out.alignment, out.number_format = copy(m.alignment), m.number_format
                self.cellule(mws, ws, t, m, out, r, c, roles, ref_cols, ex, h1, extra, lang)
            if r > h1 and (ex or ref):
                data_rows.append(r)
                lettre[r] = (ex[0] if ex else ref)[0]
            lbl = next((mws.cell(r, c).value for c in (2, 3) if isinstance(mws.cell(r, c).value, str)), None)
            if r > h1:
                est_total = ligne_total(lbl) or (ref and isinstance(self.D["postes"].get(ref, {}).get("total"), (list, str)))
                self.meta[ws.title]["totaux" if est_total else "lignes"].append(r)
            if r > h1 and ligne_total(lbl) and not ex and not ref and not extra and t not in ETATS + ("note 3a", "note 3c"):
                libres = self.total_prouve(mws, ws, r, lbl, roles, data_rows, libres, lettre, c1)
                data_rows = []
        for r in range(1, 7):                               # renvoi « sommaire » hors zone d'impression : navigation conservée
            for c in range(c1 + 1, min(mws.max_column, c1 + 6) + 1):
                if mws.cell(r, c).value == "sommaire":
                    self.cellule(mws, ws, t, mws.cell(r, c), ws.cell(r, c), r, c, roles, ref_cols, None, h1, extra, lang)
        if not extra and t not in ETATS + ("note 3a", "note 3c", "flux tresorerie"):
            self.residus(mws, ws, h1, r1, roles)
        remplir = {"tf1": self.remplir_tf1, "tf2": self.remplir_tf2, "tf3": self.remplir_tf3, "tf4": self.remplir_tf4,
                   "tf5": self.remplir_tf5, "reintegraciones": self.remplir_reint, "note 3a": self.remplir_note3,
                   "note 3c": self.remplir_note3, "note 31": self.remplir_note31, "note 34": self.remplir_note34,
                   "flux tresorerie": self.remplir_tft, "note 12": self.remplir_note12}.get(t)
        if remplir:
            remplir(ws, mws, t, lang)
        if not extra:                                      # variations et échéances APRÈS les remplissages dédiés
            for r in range(h1 + 1, r1 + 1):
                self.formules_ligne(mws, ws, r, roles)
        if t == "french version":
            ws["B6"] = {"fr": "DSF SYSCOHADA   :   VERSION FRANÇAISE", "es": "DSF SYSCOHADA   :   VERSIÓN ESPAÑOLA"}[lang]
        if t == "page garde":
            ws["D11"] = f"31/12/{self.a.exercice}"
            from openpyxl.styles import PatternFill
            ident = self.D.get("identite") or {}
            for cell, cle in (("D13", "raison_sociale"), ("D16", "sigle"), ("D17", "adresse"), ("D18", "nif")):
                if ident.get(cle):                             # identité saisie une fois : tous les en-têtes suivent
                    ws[cell] = ident[cle]
                    ws[cell].fill = PatternFill(fill_type=None)
                    self.a_completer = [x for x in self.a_completer if x != (ws.title, cell)]
        from openpyxl.styles import Alignment
        for c, ro in roles.items():                        # colonnes chiffrées : assez larges pour 15 chiffres, sinon « #### »
            if ro in ("n", "n1", "brut", "amort", "net_n", "net_n1", "var_abs", "ech1", "ech2", "ech3", "hist") and c <= c1:
                L = self.L(c)
                ws.column_dimensions[L].width = max(ws.column_dimensions[L].width or 8.43, 15.5)
                for r in range(h1 + 1, r1 + 1):
                    cell = ws.cell(r, c)
                    if type(cell).__name__ != "MergedCell" and cell.value is not None:
                        a = copy(cell.alignment)
                        cell.alignment = Alignment(horizontal=a.horizontal or "right", vertical=a.vertical, wrap_text=False, shrink_to_fit=True)
        self.entete(ws, t, c0, r0, c1)
        self.mise_en_page(ws, mws, pa, c0, c1, r1, lang)

    @staticmethod
    def L(c):
        from openpyxl.utils import get_column_letter
        return get_column_letter(c)

    def a_saisir(self, ws, out):
        from openpyxl.styles import PatternFill
        out.fill = PatternFill("solid", fgColor=AMBRE)
        self.a_completer.append((ws.title, out.coordinate))

    def date_cloture(self, y):
        """Date du modèle -> date de NOTRE document : même exercice (N-1) inchangée, modèle antérieur (N) + 1 an,
        gabarit : décalée de l'écart entre son année et l'exercice."""
        if self.gabarit:
            return y + self.a.exercice - self.annee_modele
        return y if self.a.colonne_precedente == "N-1" else y + 1

    def cellule(self, mws, ws, t, m, out, r, c, roles, ref_cols, ex, h1, extra, lang):
        v = m.value
        if v is None:
            return
        if isinstance(v, str):
            s = v.strip()
            if s == "sommaire":                                 # renvoi (minuscule) ; les titres « SOMMAIRE » restent des titres
                out.value = {"fr": "Sommaire", "es": "Índice"}[lang]
                cible = next((k for k in self.noms if norm(k) == "sommaire"), None)   # modèle sans sommaire : texte seul
                if cible:
                    out.hyperlink = f"#'{self.noms[cible]}'!A1"
                return
            if norm(s) in self.ident_modele:                                   # identité du modèle -> formule / vide
                cell = self.ident_modele[norm(s)]
                if t == "page garde":
                    out.value = None
                    self.a_saisir(ws, out)
                elif cell == "D16":                        # sigle : facultatif ; à défaut, la raison sociale
                    out.value = f'=IF({self.pg}!D16<>"",{self.pg}!D16,IF({self.pg}!D13="","",{self.pg}!D13))'
                else:
                    out.value = f'=IF({self.pg}!{cell}="","",{self.pg}!{cell})'
                return
            if t in ("fiche r1", "fiche r2", "fiche r3") and self.cellule_identite(t, out.coordinate):
                out.value = None
                self.a_saisir(ws, out)
                return
            d = re.fullmatch(r"(\d{2})/(\d{2})/(\d{4})", s)
            if d:
                out.value = f"{d.group(1)}/{d.group(2)}/{self.date_cloture(int(d.group(3)))}"
                return
            out.value = self.tr(v, lang)
            if out.value.lstrip().startswith("="):       # libellé « = CAPACITE… » : texte, pas une formule
                out.data_type = "s"
            return
        if isinstance(v, (datetime.date, datetime.datetime)):
            out.value = v.replace(year=self.date_cloture(v.year))
            out.number_format = "dd/mm/yyyy"
            return
        if not num(v):
            return
        role = roles.get(c)
        if extra and r > 4:
            self.a_saisir(ws, out)
            return
        if extra:
            out.value = v                                                      # n° de page, durée de l'exercice
            return
        if self.gabarit and f"{t}|{out.coordinate}" in self.structure.get("a_saisir", ()):
            self.a_saisir(ws, out)                     # montant que le modèle portait sans qu'aucun calcul le produise
            return
        if r <= h1 or role is None:
            out.value = v if (r <= h1 or v != 0) else None                     # n° de note, n° de ligne, durée… ; 0 de remplissage écarté
            return
        val = self.valeur(mws, t, r, c, role, ex, ref_cols, v)
        out.value = val
        if num(val):
            out.number_format = "0.0%" if role == "var_pct" else NUM
            self.couverture["N-1 reprise du modèle" if role in ("n1", "net_n1", "hist") and t not in (
                "bilan actif", "bilan passif", "bilan paysage", "compte resultat") else "calculée"] += 1
        elif role not in ("var_abs", "var_pct", "ech1", "ech2", "ech3"):
            self.couverture["non calculée (vide)"] += 1

    def cellule_identite(self, t, coord):
        """Cellules d'IDENTITÉ des fiches R1-R3 (à saisir) ; les libellés sont conservés."""
        zones = {
            "fiche r1": {"G2", "J3", "E4", "F5", "E13", "K15", "C17", "C19", "E19", "I19", "K19", "C21", "C23", "B48", "J34", "K34", "L34"}
                        | {f"{x}{y}" for x in "IK" for y in range(39, 48)},
            "fiche r2": {"C2", "G3", "C4", "C5"} | {f"{x}{y}" for x in "DEFG" for y in range(9, 20)},
            "fiche r3": {"C2", "E3", "C4", "C5"} | {f"{x}{y}" for x in "BCDEF" for y in list(range(9, 22)) + list(range(26, 41))},
        }
        return coord in zones.get(t, set())

    def valeur(self, mws, t, r, c, role, ex, ref_cols, v_modele):
        rc = max([k for k in ref_cols if k < c], default=None)                 # REF la plus proche à gauche (bilan paysage)
        ref = mws.cell(r, rc).value.strip() if rc and est_ref(mws.cell(r, rc).value) else None
        pN, p1, agg = self.D["pN"], self.D["p1"], self.D["agg"]["par_ref"]
        if t in ETATS and not ref and rc:
            ref = self.ref_sans_code(mws, r, rc)
        if t in ETATS:
            if not ref or ref not in pN:
                return None
            return {"brut": agg[ref].get("brut", pN[ref]), "amort": agg[ref].get("amort", 0.0), "net_n": pN[ref], "n": pN[ref],
                    "net_n1": p1.get(ref), "n1": p1.get(ref)}.get(role)
        if t == "flux tresorerie":
            if not ref:
                return None
            if role == "n":
                return (self.D["tft"] or {}).get(ref)
            return self.n1_modele(mws, r, c, v_modele) if role == "n1" else None
        if role == "n" and ex:
            return valeur_detail(ex[0], ex[1], self.detail)
        if role in ("n1", "hist"):
            return self.n1_modele(mws, r, c, v_modele)
        return None

    def residus(self, mws, ws, h1, r1, roles):
        """Comptes qu'un poste contient mais qu'AUCUNE ligne de la note ne cite (compte ouvert hors nomenclature,
        sous-compte non prévu) : versés à la ligne « Autres… / … divers » du même poste, comme le fait la DEF déposée,
        et tracés dans les contrôles. Sans ligne « autres », rien n'est déplacé : le contrôle signale l'écart."""
        cn = next((c for c, ro in sorted(roles.items()) if ro == "n"), None)
        if not cn:
            return
        lignes = defaultdict(list)
        for r in range(h1 + 1, r1 + 1):
            ex = expr_comptes(mws.cell(r, 1).value)
            if ex:
                lignes[ex[0]].append((r, ex[1], norm(mws.cell(r, 2).value or "")))
        def commun(a, p):
            k = 0
            while k < min(len(a), len(p)) and a[k] == p[k]:
                k += 1
            return k

        for ref, ls in lignes.items():
            pre = [p for _, termes, _ in ls for sg, p, _ in termes if sg > 0]
            hors = [(a, v) for a, rf, col, v in self.detail if rf == ref and col != "amort"
                    and not any(a.startswith(p) for _, termes, _ in ls for _, p, _ in termes) and abs(v) >= 0.5]
            par_ligne = defaultdict(list)
            for a, v in hors:
                # ligne dont un préfixe (positif) partage le plus long début avec le compte ; à égalité, la ligne « autres »
                cands = [(max((commun(a, p) for sg, p, _ in termes if sg > 0), default=0),
                          lb.startswith("autre") or " divers" in lb, r) for r, termes, lb in ls]
                best = max(c[0] for c in cands)
                ex_aequo = [c for c in cands if c[0] == best]
                if best >= 3:        # compte de même sous-classe (ex. 1960 -> ligne 1961) : première ligne de la sous-classe
                    cible = ex_aequo[0][2]
                elif best == 2:      # même classe seulement : la ligne « autres » de la classe, sinon la dernière
                    cible = next((c[2] for c in reversed(ex_aequo) if c[1]), ex_aequo[-1][2])
                else:
                    cible = next((c[2] for c in reversed(cands) if c[1]), None)
                par_ligne[cible].append((a, v))
            for cible, lst in par_ligne.items():
                res = round(sum(-v if ref[0] == "R" else v for _, v in lst), 2)
                comptes = ", ".join(f"{a.replace('129~', '')} ({v:,.0f})" for a, v in sorted(lst, key=lambda x: -abs(x[1]))[:6])
                if cible and num(ws.cell(cible, cn).value):
                    ws.cell(cible, cn).value = round(ws.cell(cible, cn).value + res, 2)
                    self.controles.append((f"{mws.title} : comptes du poste {ref} hors nomenclature de la note", True,
                                           f"{res:,.0f} présentés ligne {cible} « {str(mws.cell(cible, 2).value).strip()[:40]} » "
                                           f"(préfixe le plus proche) — {comptes}"))
                else:
                    self.controles.append((f"{mws.title} : comptes du poste {ref} hors nomenclature de la note", False,
                                           f"{res:,.0f} non présentés (aucune ligne proche) — {comptes}"))

    def ref_sans_code(self, mws, r, rc):
        """Ligne d'état sans REF dans le modèle (ex. « TOTAL DETTES FINANCIERES » sans DD) : poste retrouvé par son
        libellé, sinon par la valeur N-1 que le modèle y porte (unique parmi les postes) ; sinon aucun."""
        if (mws.title, r) in self._refs_devinees:
            return self._refs_devinees[(mws.title, r)]
        lbl = norm(mws.cell(r, rc + 1).value or mws.cell(r, rc + 2).value or "")
        ref = None
        cle = f"{norm(mws.title)}|{r}"
        if self.gabarit and cle in self.structure.get("refs", {}):      # prouvé sur le dossier réel, figé
            ref = self.structure["refs"][cle]
        elif lbl:
            ref = next((k for k, p in self.D["postes"].items() if norm(p.get("libelle", ""))[:25] == lbl[:25]), None)
        if ref is None:
            vals = [mws.cell(r, c).value for c in range(rc + 1, mws.max_column + 1) if num(mws.cell(r, c).value)]
            cands = {k for k, v in self.D["p1"].items() if any(abs(v - x) < 1 and abs(x) >= 1 for x in vals)}
            ref = cands.pop() if len(cands) == 1 else None
        self._refs_devinees[(mws.title, r)] = ref
        self.trace["refs"][cle] = ref
        self.controles.append((f"{mws.title} ligne {r} « {lbl[:40]} » sans REF dans le modèle", ref is not None,
                               f"rattachée au poste {ref}" if ref else "non rattachée : cellules laissées vides"))
        return ref

    def n1_modele(self, mws, r, c, v):
        if self.gabarit:                                   # aucun N-1 dans un formulaire vierge : vide, jamais 0
            return None
        if self.a.colonne_precedente == "N-1":
            return v
        g = mws.cell(r, c - 1).value
        return g if num(g) else None

    def total_prouve(self, mws, ws, r, lbl, roles, data_rows, libres, lettre, c1):
        """Formule d'une ligne de total. Plusieurs structures sont possibles (somme des lignes depuis le dernier total,
        somme des sous-totaux, total net = total précédent − dépréciations, solde produits − charges) : on retient
        celle que les MONTANTS DU MODÈLE prouvent (égalité sur toutes les colonnes chiffrées, à 1 franc près) ;
        à défaut de preuve (modèle à zéro), l'ordre de préférence suit le libellé.
        libres = sous-totaux pas encore repris par un total supérieur : [(ligne, signe)]."""
        n = norm(lbl)
        cols = [c for c, ro in roles.items() if ro in ("n", "n1", "brut", "amort", "net_n", "net_n1", "hist") and c <= c1]
        cands = {}
        if data_rows:
            cands["lignes"] = [(x, 1) for x in data_rows]
        if libres:
            cands["sous_totaux"] = [(x, 1) for x, _ in libres] + [(x, 1) for x in data_rows]
            cands["solde"] = list(libres) + [(x, -1 if lettre.get(x) == "R" else 1) for x in data_rows]
            if data_rows:
                cands["net"] = [(libres[-1][0], 1)] + [(x, -1) for x in data_rows]
        if not cands:
            return libres
        ordre = (["net", "lignes", "sous_totaux", "solde"] if "net" in n else
                 ["solde", "sous_totaux", "lignes"] if "produits - charges" in n or "produits-charges" in n else
                 ["sous_totaux", "solde", "lignes", "net"] if libres and not data_rows else
                 ["lignes", "sous_totaux", "net", "solde"])
        ordre = [k for k in ordre if k in cands]

        def prouve(k):
            vals = [(mws.cell(r, c).value, [mws.cell(x, c).value for x, _ in cands[k]], [s for _, s in cands[k]]) for c in cols]
            vals = [(t, xs, ss) for t, xs, ss in vals if num(t)]
            if not vals or all(abs(t) < 1 for t, _, _ in vals):
                return None                                   # rien à prouver
            return all(abs(t - sum(s * (x if num(x) else 0) for x, s in zip(xs, ss))) <= 1 for t, xs, ss in vals)

        preuves = {k: prouve(k) for k in ordre}
        choix = next((k for k in ordre if preuves[k]), None)
        cle = f"{norm(mws.title)}|{r}"
        fige = self.structure.get("totaux", {}).get(cle) if self.gabarit else None
        if choix is None and fige in cands:                # gabarit à zéro : la structure prouvée sur le dossier réel
            choix = fige
        self.trace["totaux"][cle] = choix or ordre[0]
        if choix is None:
            choix = ordre[0]
            if any(p is False for p in preuves.values()) and not any(preuves.values()):
                self.controles.append((f"{ws.title} ligne {r} « {lbl.strip()[:40]} » : structure du total non prouvée par le modèle",
                                       False, f"formule retenue par défaut : {choix}"))
        for c in cols:
            L = self.L(c)
            f = "=" + "".join(("+" if s > 0 else "-") + f"N({L}{x})" for x, s in cands[choix]).lstrip("+")
            ws.cell(r, c).value = f
            ws.cell(r, c).number_format = NUM
        signe = -1 if data_rows and all(lettre.get(x) == "R" for x in data_rows) else 1
        if choix == "lignes":
            return libres + [(r, signe)]
        if choix == "net":
            return libres[:-1] + [(r, libres[-1][1])]
        return [(r, 1)]

    def formules_ligne(self, mws, ws, r, roles):
        col = {}
        for c, v in sorted(roles.items()):
            col.setdefault(v, c)
        cn, c1 = col.get("net_n") or col.get("n"), col.get("net_n1") or col.get("n1")
        if not cn or ws.cell(r, cn).value is None:
            return
        N = f"{self.L(cn)}{r}"
        if c1:
            P = f"{self.L(c1)}{r}"
            if "var_abs" in col and num(mws.cell(r, col["var_abs"]).value):
                ws.cell(r, col["var_abs"]).value = f"=N({N})-N({P})"
                ws.cell(r, col["var_abs"]).number_format = NUM
            if "var_pct" in col and num(mws.cell(r, col["var_pct"]).value):
                ws.cell(r, col["var_pct"]).value = f'=IF(N({P})=0,"",(N({N})-N({P}))/{P})'   # convention du modèle : / N-1
                ws.cell(r, col["var_pct"]).number_format = "0.0%"
        if "ech1" in col and num(mws.cell(r, col["ech1"]).value):   # échéancier : < 1 an PAR DÉFAUT, à ventiler
            autres = "".join(f"-N({self.L(col[k])}{r})" for k in ("ech2", "ech3") if k in col)
            ws.cell(r, col["ech1"]).value = f"={N}{autres}"
            ws.cell(r, col["ech1"]).number_format = NUM
            for k in ("ech2", "ech3"):
                if k in col:
                    ws.cell(r, col[k]).value = 0              # à ventiler par l'entreprise ; « < 1 an » se recalcule
                    ws.cell(r, col[k]).number_format = NUM

    # ---------------------------------------------------------------- remplissages dédiés
    @staticmethod
    def vider(ws, cols, rows):
        for r in rows:
            for c in cols:
                if type(ws.cell(r, c)).__name__ != "MergedCell":
                    ws.cell(r, c).value = None

    def ecrire(self, ws, coord, v, fmt=NUM):
        from openpyxl.styles import Alignment
        ws[coord] = v
        ws[coord].number_format = fmt
        if (ws.title, coord) in self.a_completer and v is not None:     # calculé finalement : plus « à saisir »
            from openpyxl.styles import PatternFill
            self.a_completer.remove((ws.title, coord))
            ws[coord].fill = PatternFill(fill_type=None)
        a = copy(ws[coord].alignment)                    # jamais « #### » : le chiffre se réduit à la largeur de la colonne
        ws[coord].alignment = Alignment(horizontal=a.horizontal or "right", vertical=a.vertical, wrap_text=False, shrink_to_fit=True)
        self.a_completer = [x for x in self.a_completer if x != (ws.title, coord)]

    def remplir_note3(self, ws, mws, t, lang):
        n3, postes = self.D["n3"], self.D["postes"]
        cols = "CDEFGHI" if t == "note 3a" else "CDEFG"
        lignes, sous, groupes, faits = [], [], [], {}
        for r in range(6, mws.max_row + 1):
            a, lbl = mws.cell(r, 1).value, norm(mws.cell(r, 2).value or "")
            ref = a.strip()[:2] if isinstance(a, str) and re.match(r"^[A-Z]{2}\b", a.strip()) else None
            if ref is None and lbl:                            # ligne sans REF dans le modèle (ex. AM en note 3C) : par libellé
                ref = next((k for k in n3 if norm(postes.get(k, {}).get("libelle", ""))[:14] == lbl[:14]), None)
            if lbl.startswith(("sous total", "sous-total")) or (lbl.startswith("total") and ref not in n3):
                parts = lignes if lbl.startswith("sous") else (sous or lignes)
                for k in cols:
                    if parts:
                        self.ecrire(ws, f"{k}{r}", "=" + "+".join(f"N({k}{x})" for x in parts))
                if lbl.startswith("sous"):
                    sous.append(r)
                    lignes = []
                continue
            if ref not in n3:
                comp = postes.get(ref, {}).get("total")
                if isinstance(comp, list):                     # poste de regroupement (AD, AI, AQ) : ses composants
                    groupes.append((r, comp))
                continue
            if ref in faits:                                   # sous-ligne du même poste (ex. « AP 251 », « AP 252 »)
                continue
            faits[ref] = r
            lignes.append(r)
            L = n3[ref]
            if t == "note 3a":        # C ouverture · D acquisitions · E virements + · F réévaluation · G cessions · H virements − · I clôture
                vals = {"C": L["b_ouv"], "D": L["b_aug"], "E": max(L["b_vir"], 0), "F": 0, "G": L["b_dim"], "H": max(-L["b_vir"], 0),
                        "I": f"=C{r}+D{r}+E{r}+F{r}-G{r}-H{r}"}
            else:                     # C cumul ouverture · D dotations · E virements · F diminutions · G clôture
                vals = {"C": L["a_ouv"], "D": L["a_dot"], "E": max(L["a_vir"], 0), "F": L["a_rep"] + max(-L["a_vir"], 0),
                        "G": f"=C{r}+D{r}+E{r}-F{r}"}
            for k, v in vals.items():
                self.ecrire(ws, f"{k}{r}", v)
        for r, comp in groupes:
            enfants = [faits[x] for x in comp if x in faits]
            for k in cols:
                if enfants:
                    self.ecrire(ws, f"{k}{r}", "=" + "+".join(f"N({k}{x})" for x in enfants))

    def remplir_note12(self, ws, mws, t, lang):
        """Écarts de conversion : le montant du poste (BU actif, DV passif) porté sur la ligne de détail qui suit la
        ligne « BU 478 » / « DV 479 », colonne « variation en valeur absolue » (la devise et les cours restent à saisir)."""
        pN = self.D["pN"]
        for r in range(6, mws.max_row + 1):
            a = mws.cell(r, 1).value
            ref = a.strip()[:2] if isinstance(a, str) and a.strip()[:2] in ("BU", "DV") else None
            if ref and num(mws.cell(r + 1, 7).value):
                self.ecrire(ws, f"G{r + 1}", round(pN.get(ref, 0.0), 2))

    def remplir_tft(self, ws, mws, t, lang):
        """Lignes du TFT sans REF : variation du BF (somme FB..FE) et contrôle trésorerie (bilan BT - DT)."""
        pos = {mws.cell(r, 2).value.strip(): r for r in range(5, mws.max_row + 1) if est_ref(mws.cell(r, 2).value)}
        for r in range(5, mws.max_row + 1):
            n = norm(mws.cell(r, 3).value or "")
            if n.startswith("variation du bf"):
                rows = [pos[k] for k in ("FB", "FC", "FD", "FE") if k in pos]
                for c in "FG":
                    self.ecrire(ws, f"{c}{r}", "=" + "+".join(f"N({c}{x})" for x in rows))
            elif n.startswith("controle"):
                pN = self.D["pN"]
                self.ecrire(ws, f"F{r}", round(pN.get("BT", 0) - pN.get("DT", 0), 2))
                if num(mws[f"G{r}"].value) and self.a.colonne_precedente == "N-1":
                    self.ecrire(ws, f"G{r}", mws[f"G{r}"].value)

    def remplir_note31(self, ws, mws, t, lang):
        pN = self.D["pN"]
        cibles = {"capital social": pN.get("CA"), "chiffre d'affaires": pN.get("XB"), "resultat des activites ordinaires": pN.get("XG"),
                  "participation des trav": -pN.get("RQ", 0), "impot sur le resultat": -pN.get("RS", 0), "resultat net": pN.get("XI")}
        for r in range(7, mws.max_row + 1):
            n = norm(mws.cell(r, 2).value or "")
            for k, v in cibles.items():
                if n.startswith(k) and v is not None:
                    self.ecrire(ws, f"C{r}", v)
            if self.a.colonne_precedente == "N" and not self.gabarit:   # modèle antérieur : N..N-3 deviennent N-1..N-4
                for src, dst in zip("CDEF", "DEFG"):
                    if num(mws[f"{src}{r}"].value):
                        self.ecrire(ws, f"{dst}{r}", mws[f"{src}{r}"].value)

    def remplir_note34(self, ws, mws, t, lang):
        """Fiche de synthèse : soldes intermédiaires et éléments de CAF depuis le compte de résultat recalculé."""
        pN = self.D["pN"]
        g = lambda *k: round(sum(pN.get(x, 0.0) for x in k), 2)
        cpt = lambda ref, pre: round(sum(v for a, rf, _c, v in self.detail if rf == ref and a.startswith(pre)), 2)
        change_p, change_c = cpt("TK", "776"), cpt("RM", "676")          # change isolé des revenus / frais financiers
        sig = [("+ revenus financiers", round(g("TK") - change_p, 2)), ("+ gains de change", change_p),
               ("- frais financiers", round(g("RM") - change_c, 2)), ("- pertes de change", change_c),
               ("chiffre d'affaires", g("XB")), ("marge commerciale", g("XA")), ("valeur ajoutee", g("XC")),
               ("excedent brut d'exploitation", g("XD")), ("resultat d'exploitation", g("XE")), ("resultat financier", g("XF")),
               ("resultat des activites ordinaires", g("XG")), ("resultat hors activites ordinaires", g("XH")), ("resultat net", g("XI")),
               ("ebe", g("XD")), ("+ produits hao", g("TN", "TO")), ("- participation", g("RQ")), ("- impots sur les resultats", g("RS"))]
        lignes = {}
        for r in range(9, mws.max_row + 1):
            n = norm(mws.cell(r, 2).value or "")
            lignes.setdefault(n[:30], r)
            for k, v in sig:
                if n.startswith(k):
                    self.ecrire(ws, f"D{r}", v)
                    break
        # structure financière, trésorerie, endettement : postes du bilan / TFT, puis FORMULES entre lignes
        tft = self.D["tft"] or {}
        postes = [("capitaux propres", g("CP")), ("+ dettes financieres", g("DD")), ("- actif immobilise", -g("AZ")),
                  ("actif circulant d'exploitation", g("BK", "BU")), ("- passif circulant d'exploitation", -g("DP", "DV")),
                  ("actif circulant hao", g("BA")), ("- passif circulant hao", -g("DH")),
                  ("flux de tresorerie des activites op", tft.get("ZB")), ("-flux de tresorerie des activites d", tft.get("ZC")),
                  ("+flux de tresorerie des activites d", tft.get("ZF")), ("endettement financier brut", g("DA", "DB", "DT")),
                  ("-tresorerie actif", -g("BT"))]
        formules = [("= capacite d'autofinancement d'expl", "ebe", None), ("= capacite d'autofinancement glob", "= capacite d'autofinancement d'expl", None),
                    ("= ressources stables", "capitaux propres", None), ("= fonds de roulement", "= ressources stables", None),
                    ("= besoin de financement d'exploitat", "actif circulant d'exploitation", None),
                    ("= besoin de financement hao", "actif circulant hao", None), ("= variation de la tresorerie nett", "flux de tresorerie des activites op", None),
                    ("= endettement financier net", "endettement financier brut", None)]
        pos = {k: r for k, r in ((norm(mws.cell(r, 2).value or ""), r) for r in range(9, mws.max_row + 1)) if k}

        def ligne(prefixe):
            return next((r for k, r in pos.items() if k.startswith(prefixe)), None)
        for k, v in postes:
            r = ligne(k)
            if r and v is not None:
                self.ecrire(ws, f"D{r}", round(v, 2))
        for k, debut, _ in formules:                        # somme des lignes depuis « debut » jusqu'à la ligne précédente
            r, r0 = ligne(k), ligne(debut)
            if r and r0 and r0 < r:
                for c in "DE":
                    self.ecrire(ws, f"{c}{r}", "=" + "+".join(f"N({c}{x})" for x in range(r0, r) if x not in (ligne("= capacite d'autofinancement d'expl"),) or r0 == x))
        r, r1_, r2_ = ligne("= autofinancement"), ligne("= capacite d'autofinancement glob"), ligne("- distributions de dividendes")
        if r and r1_:
            for c in "DE":
                self.ecrire(ws, f"{c}{r}", f"=N({c}{r1_})" + (f"+N({c}{r2_})" if r2_ else ""))
            if r2_:
                ws[f"D{r2_}"] = None
                self.a_saisir(ws, ws[f"D{r2_}"])             # dividendes décidés : donnée d'assemblée, à saisir
        bg, fr_, bh = ligne("besoin de financement global"), ligne("= fonds de roulement"), ligne("= besoin de financement hao")
        be = ligne("= besoin de financement d'exploitat")
        if bg and be:
            for c in "DE":
                self.ecrire(ws, f"{c}{bg}", f"=N({c}{be})" + (f"+N({c}{bh})" if bh else ""))
        tn = ligne("tresorerie nette (5)")
        if tn and fr_ and bg:
            for c in "DE":
                self.ecrire(ws, f"{c}{tn}", f"=N({c}{fr_})-N({c}{bg})")
        ct = ligne("controle : tresorerie nette")
        if ct:
            self.ecrire(ws, f"D{ct}", round(g("BT") - g("DT"), 2))
        rf = ligne("rentabilite financiere")
        cp = ligne("capitaux propres")
        rn = ligne("resultat net")
        if rf and cp and rn:
            for c in "DE":
                self.ecrire(ws, f"{c}{rf}", f'=IF(N({c}{cp})=0,"",N({c}{rn})/{c}{cp})', "0.0%")

    def remplir_tf1(self, ws, mws, t, lang):
        f = self.D["fisc"]["tf1"]
        for r in range(8, mws.max_row + 1):                 # tout le tableau : la ligne 26 (résultat net) est en bas
            n = mws.cell(r, 6).value
            self.vider(ws, (7, 8), [r])
            if num(n) and int(n) in f:
                v = f[int(n)]
                perte = int(n) in (1, 22) and v < 0
                self.ecrire(ws, f"{'G' if perte else 'H'}{r}", -v if perte else v)

    def remplir_tf2(self, ws, mws, t, lang):
        f = self.D["fisc"]["tf2"]
        for r in range(10, 32):
            n = mws.cell(r, 5).value
            self.vider(ws, (6,), [r])
            if num(n) and int(n) in f:
                self.ecrire(ws, f"F{r}", f[int(n)], "0.0%" if int(n) in (4, 13) else NUM)
        self.ecrire(ws, "F31", "=F18")

    def remplir_reint(self, ws, mws, t, lang):
        ordre = [str(mws.cell(k, 1).value) for k in range(3, mws.max_row + 1) if mws.cell(k, 1).value]
        self.vider(ws, (1, 2, 3), range(3, max(mws.max_row, 20) + 1))
        r = 3
        det = [x for x in self.D["fisc"]["reintegrations_detail"] if x.get("sens") != "déduction"]
        det.sort(key=lambda x: ordre.index(str(x.get("compte"))) if str(x.get("compte")) in ordre else len(ordre))
        for x in det:
            ws.cell(r, 1).value, ws.cell(r, 2).value = x.get("compte", ""), self.tr(x.get("libelle", ""), lang)
            self.ecrire(ws, f"C{r}", x["montant"])
            r += 1
        ws.cell(r, 1).value = "TOTAL"
        self.ecrire(ws, f"C{r}", f"=SUM(C3:C{r - 1})")

    def remplir_tf3(self, ws, mws, t, lang):
        if not self.D["tf345"]:
            return
        r3 = self.D["tf345"]["tf3"]
        col = 4 if r3["nature"] == "prestacion_servicios" else 5
        C = self.L(col)
        self.vider(ws, range(3, 8), range(10, 25))
        for x in r3["mois"]:
            self.ecrire(ws, f"{C}{9 + x['mois']}", x["base"])
        self.ecrire(ws, f"{C}22", f"=SUM({C}10:{C}21)")
        self.ecrire(ws, f"{C}23", r3["taux"], "0%")
        self.ecrire(ws, f"{C}24", f"=ROUND({C}22*{C}23,0)")
        pays = list(r3["paiements"])
        if r3["interpretation"] == "dans_paiements":
            # une ligne par compte d'IVA déductible (comme la DEF déposée), si les débits de la balance le justifient
            mv = self.det.get("mouvements") or {}
            ded = {a: md for a, (md, _) in mv.items() if a.startswith("445") and md}
            lib = {"fr": "IVA déductible — compte ", "es": "IVA deducible — cuenta "}[lang]
            if ded and abs(sum(ded.values()) - r3["iva_deductible"]) < 1:
                pays += [{"date": "", "piece": lib + a, "montant": v} for a, v in sorted(ded.items())]
            else:
                pays.append({"date": "", "piece": lib.split(" — ")[0], "montant": r3["iva_deductible"]})
        self.vider(ws, (3, 4, 5), range(29, 41))
        for i, p in enumerate(pays[:12]):
            ws.cell(29 + i, 3).value, ws.cell(29 + i, 4).value = p["date"], p["piece"]
            self.ecrire(ws, f"E{29 + i}", p["montant"])
        self.ecrire(ws, "E41", "=SUM(E29:E40)")
        self.ecrire(ws, "E43", f"=MAX({C}24-E41,0)")
        self.ecrire(ws, "E45", f"=MAX(E41-{C}24,0)")

    def remplir_tf4(self, ws, mws, t, lang):
        if not self.D["tf345"]:
            return
        cats = self.D["tf345"]["tf4"]["categories"]
        self.vider(ws, range(4, 13), range(11, 39))
        for rr, k in ((11, "national"), (13, "non_national")):
            v = cats[k]
            self.ecrire(ws, f"D{rr}", v["remunerations"])
            self.ecrire(ws, f"J{rr}", v["irpf_retenu"])
            self.ecrire(ws, f"K{rr}", v["irpf_reverse"])
            self.ecrire(ws, f"L{rr}", f"=J{rr}-K{rr}")
        for c in "DJKL":
            self.ecrire(ws, f"{c}39", f"=SUM({c}10:{c}38)")

    def remplir_tf5(self, ws, mws, t, lang):
        tot = next((r for r in range(22, mws.max_row + 1) if num(mws.cell(r, 12).value)), 36)
        self.vider(ws, range(1, 13), range(9, tot))
        rr = 9
        for x in self.tf5_lignes[: tot - 9]:
            ws.cell(rr, 1).value = x["nom"] or f"[{x['code']}]"
            ws.cell(rr, 3).value = x["adresse"]
            ws.cell(rr, 7).value = self.tr(x["nature"], lang)
            ws.cell(rr, 11).value = x["nif"] or None
            if num(x["ht"]):
                self.ecrire(ws, f"L{rr}", x["ht"])
            else:
                self.a_saisir(ws, ws.cell(rr, 12))
            if not x["nom"]:
                self.a_saisir(ws, ws.cell(rr, 1))
            if not x["nif"]:
                self.a_saisir(ws, ws.cell(rr, 11))
            rr += 1
        if len(self.tf5_lignes) > tot - 9:
            self.controles.append(("TF5 : place insuffisante dans le modèle", False,
                                   f"{len(self.tf5_lignes) - (tot - 9)} bénéficiaire(s) non imprimé(s) — voir le classeur des contrôles"))
        self.ecrire(ws, f"L{tot}", f"=SUM(L9:L{tot - 1})")

    # ---------------------------------------------------------------- balance (annexe art. 79.2)
    def feuille_balance(self, ws, lang):
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        hd = {"fr": ["Compte", "Libellé", "Ouverture débit", "Ouverture crédit", "Mouvements débit", "Mouvements crédit", "Solde débit", "Solde crédit"],
              "es": ["Cuenta", "Denominación", "Apertura debe", "Apertura haber", "Movimientos debe", "Movimientos haber", "Saldo deudor", "Saldo acreedor"]}[lang]
        ws["A1"] = f'=IF({self.pg}!D16<>"",{self.pg}!D16,IF({self.pg}!D13="","",{self.pg}!D13))'   # sigle, sinon raison sociale
        ws["A2"] = {"fr": "BALANCE GÉNÉRALE AU 31/12/", "es": "BALANCE GENERAL DE COMPROBACIÓN AL 31/12/"}[lang] + str(self.a.exercice)
        ws["A3"] = f'="NIF : "&{self.pg}!D18'
        ws["A1"].font = Font(size=12, bold=True, color=BLEU)
        ws["A2"].font = Font(size=11, bold=True, color=BLEU)
        for j, h in enumerate(hd, 1):
            c = ws.cell(5, j, h)
            c.font, c.fill = Font(bold=True, color="FFFFFF", size=9), PatternFill("solid", fgColor=BLEU)
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ouv, mv = self.det.get("soldes_ouverture") or {}, self.det.get("mouvements") or {}
        fin = Side(style="thin", color="D9D9D9")
        r = 6
        for a in sorted(self.det["soldes"]):
            o, (md, mk), so = ouv.get(a, 0.0), mv.get(a, (0.0, 0.0)), self.det["soldes"][a]
            for j, v in enumerate([a, self.det["libelles"].get(a, ""), max(o, 0), max(-o, 0), md, mk, max(so, 0), max(-so, 0)], 1):
                c = ws.cell(r, j, v)
                c.font, c.border = Font(size=8), Border(bottom=fin)
                if j > 2:
                    c.number_format = NUM
            r += 1
        ws.cell(r, 1, "TOTAL")
        for j in range(3, 9):
            ws.cell(r, j, f"=SUM({self.L(j)}6:{self.L(j)}{r - 1})").number_format = NUM
        for j in range(1, 9):
            ws.cell(r, j).font, ws.cell(r, j).fill = Font(bold=True, size=9), PatternFill("solid", fgColor="D9E1F2")
        for j, w in enumerate([11, 42, 15, 15, 15, 15, 15, 15], 1):
            ws.column_dimensions[self.L(j)].width = w
        ws.sheet_view.showGridLines = False
        ws.print_title_rows = "5:5"
        ws.print_area = f"A1:H{r}"
        self.pied(ws, lang, "portrait")

    # ---------------------------------------------------------------- en-tête, mise en page
    def entete(self, ws, t, c0, r0, c1):
        """En-tête : dénomination en évidence, filet sous l'en-tête, logo à droite (si fourni)."""
        from openpyxl.styles import Border, Font, Side
        if t not in ("page garde", "french version", "sommaire", "statistiques et syntheses"):
            for r in (1, 2, 3):
                for c in range(c0, c1 + 1):
                    v = ws.cell(r, c).value
                    if isinstance(v, str) and v.startswith("=IF("):
                        ws.cell(r, c).font = Font(name="Calibri", size=12 if r == 1 else 9, bold=r == 1, color=BLEU)
            for c in range(c0, c1 + 1):
                b = copy(ws.cell(3, c).border)
                ws.cell(3, c).border = Border(left=b.left, right=b.right, top=b.top, bottom=Side(style="medium", color=BLEU))
        if t == "page garde" and self.a.logo and os.path.exists(self.a.logo):   # logo : page de présentation SEULEMENT
            from openpyxl.drawing.image import Image
            im = Image(self.a.logo)
            h = 75
            im.width, im.height = int(im.width * h / im.height), h
            larg, c = 0.0, c1                                # ancrage : colonnes de droite de la zone d'impression
            while c > c0 and larg * 7 < im.width:
                larg += ws.column_dimensions[self.L(c)].width or 8.43
                c -= 1
            # ancré DANS la zone d'impression (la page de garde commence en ligne 4 dans le modèle GE)
            ws.add_image(im, f"{self.L(max(c + 1, c0))}{r0 + 1}")

    def pied(self, ws, lang, orientation):
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = orientation
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddFooter.left.text = {"fr": f"DEF {self.a.exercice} — version française", "es": f"DEF {self.a.exercice} — versión en español"}[lang]
        ws.oddFooter.center.text = "&P / &N"
        ws.oddFooter.right.text = "&A"
        for p in (ws.oddFooter.left, ws.oddFooter.center, ws.oddFooter.right):
            p.size, p.color = 7, "7F7F7F"

    def mise_en_page(self, ws, mws, pa, c0, c1, r1, lang):
        from openpyxl.worksheet.page import PageMargins
        ws.sheet_view.showGridLines = False
        ws.print_area = pa if pa else f"{self.L(c0)}1:{self.L(c1)}{r1}"
        o = mws.page_setup.orientation
        larg = sum((ws.column_dimensions[self.L(c)].width or 8.43) for c in range(c0, c1 + 1))
        self.pied(ws, lang, o if o in ("portrait", "landscape") else ("landscape" if larg > 110 else "portrait"))
        if r1 <= 70 and (not pa or len(pa) == 1):
            ws.page_setup.fitToHeight = 1                  # une feuille courte tient sur UNE page (pas de page orpheline)
        ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.45, bottom=0.55, header=0.2, footer=0.25)
        ws.print_options.horizontalCentered = True

    # ---------------------------------------------------------------- contrôles (classeur séparé, interne)
    def controles_classeur(self, sortie):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        wb = Workbook()
        ws = wb.active
        ws.title = "Controles"
        ws.append(["Contrôle / élément", "Statut", "Détail"])
        for x in self.D["controles"]:
            ws.append([x["controle"], "OK" if x["ok"] else "ÉCART", x["detail"]])
        for nom, ok, det in self.controles:
            ws.append([nom, "OK" if ok else "À VOIR", det])
        ws.append([])
        ws.append(["COUVERTURE DES CELLULES CHIFFRÉES DU MODÈLE", "", ""])
        for k, n in self.couverture.items():
            ws.append([k, n, ""])
        ws.append(["à saisir (identité, effectifs, extra-comptable)", len(self.a_completer), ""])
        for f, n in Counter(f for f, _ in self.a_completer).most_common():
            ws.append([f"   {f}", n, ""])
        ws.append([])
        ws.append(["AVERTISSEMENTS", "", ""])
        for w in self.D["avertissements"]:
            ws.append([w, "", ""])
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=BLEU)
        for c, w in zip("ABC", (70, 14, 110)):
            ws.column_dimensions[c].width = w
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        wb.save(sortie)


def exporter_pdf(xlsx, pdf):
    """Export PDF du classeur ENTIER par Excel (= « Fichier > Exporter > PDF » avec les zones d'impression posées)."""
    import subprocess
    ps = rf"""
$x = New-Object -ComObject Excel.Application; $x.Visible = $false; $x.DisplayAlerts = $false
try {{ $wb = $x.Workbooks.Open('{os.path.abspath(xlsx)}', 0, $true); $x.CalculateFull()
  $wb.ExportAsFixedFormat(0, '{os.path.abspath(pdf)}', 0, $true, $false); $wb.Close($false); 'OK' }} finally {{ $x.Quit() }}
"""
    if os.environ.get("FACIL_PDF", "").lower() != "natif":
        try:
            cp = subprocess.run(["powershell", "-NoProfile", "-Command", ps], capture_output=True, text=True, timeout=900)
            if "OK" in cp.stdout and os.path.exists(pdf):
                return True
        except Exception:
            pass
    try:                                                  # sans Excel : formules évaluées par le plugin, rendu natif
        import tempfile
        import formules_xl
        import rendu_pdf
        val = os.path.join(tempfile.mkdtemp(), "valeurs.xlsx")
        formules_xl.evaluer_classeur(xlsx, val)
        return rendu_pdf.classeur_vers_pdf(val, pdf)
    except Exception:
        return False


def main():
    p = argparse.ArgumentParser(description="DEF sur le modèle officiel GE (FR + ES), classeur + PDF.")
    p.add_argument("--modele", required=True, help="DEF de référence (structure)")
    p.add_argument("--balance", required=True)
    p.add_argument("--exercice", type=int, required=True)
    p.add_argument("--grand-livre")
    p.add_argument("--profil")
    p.add_argument("--decisions")
    p.add_argument("--colonne-precedente", choices=["N", "N-1"], default="N",
                   help="colonne du modèle qui porte NOTRE N-1 : N (modèle de l'exercice précédent) ou N-1 (modèle du même exercice)")
    p.add_argument("--logo", help="logo de l'entité : page de garde du classeur et couverture du rapport PDF, nulle part ailleurs")
    p.add_argument("--langues", default="fr,es")
    p.add_argument("--sortie", required=True, help="préfixe des fichiers produits")
    p.add_argument("--pdf", action="store_true",
                   help="RAPPORT PDF du plugin (mise en page du rapport, logo en couverture) + PDF des contrôles internes")
    p.add_argument("--pdf-excel", action="store_true",
                   help="en plus : PDF des classeurs tels qu'Excel les imprime (le même que « Fichier > Exporter » dans Excel)")
    p.add_argument("--pro", action="store_true", help="produire AUSSI le modèle professionnel (mêmes cellules, présentation normalisée)")
    a = p.parse_args()
    res = produire(a, journal=print)
    for x in res["gen"].D["controles"]:
        print(f"  {'OK   ' if x['ok'] else 'ÉCART'} {x['controle']}")
    for nom, ok, det in res["gen"].controles:
        print(f"  {'OK    ' if ok else 'À VOIR'} {nom} — {det}")


def chemins_prefixe(sortie):
    """Nommage historique : tous les fichiers à côté du préfixe --sortie."""
    def chemin(nature, lang=None):
        L = (lang or "").upper()
        return {"xlsx": f"{sortie}_{L}.xlsx", "pro_xlsx": f"{sortie}_{L}_pro.xlsx", "excel_pdf": f"{sortie}_{L}.pdf",
                "rapport_pdf": f"{sortie}_rapport_{L}.pdf", "balance_pdf": f"{sortie}_annexe-balance_{L}.pdf",
                "controles_xlsx": f"{sortie}_controles.xlsx", "controles_pdf": f"{sortie}_controles.pdf",
                "rapport_abrege_pdf": f"{sortie}_rapport.pdf", "sn_xlsx": f"{sortie}_SYSCOHADA_{L}.xlsx",
                "sn_pdf": f"{sortie}_SYSCOHADA_{L}.pdf"}[nature]
    return chemin


PRESENTATION = {"fr": ("Présentation", "présentée", "néant — non produite (aucun montant N ni N-1)"),
                "es": ("Presentación", "presentada", "sin importe — no presentada (ni N ni N-1)")}


def notes_neant(g, x_evalue, lang):
    """Notes du classeur sans AUCUN montant (N et N-1), selon la règle MÊME du rapport PDF (rapport_def.utile) ;
    jamais les notes extra-comptables, ni une note qui attend une saisie de l'entreprise (fond ambre)."""
    import openpyxl
    import rapport_def as rd
    wb = openpyxl.load_workbook(x_evalue, data_only=True)
    a_saisir = {s for s, _ in g.a_completer}
    out = []
    for titre_src, nom in g.noms.items():
        t = norm(titre_src)
        if not t.startswith("note") or t in rd.EXTRA or t in EXTRA_COMPTABLES or nom in a_saisir or nom not in wb.sheetnames:
            continue
        if not rd.utile(rd.extraire(wb[nom], g.meta[nom], t, lang), texte=(t == "note 2")):
            out.append(nom)
    return out


def elaguer_notes(g, x, lang):
    """Retire du classeur les notes néant (sauf si une autre feuille y fait référence) et le dit dans la grille
    d'analyse des notes (colonne « Présentation »). Retourne les feuilles retirées."""
    import tempfile
    import openpyxl
    from openpyxl.utils import get_column_letter
    import formules_xl
    ev = os.path.join(tempfile.mkdtemp(), "eval.xlsx")
    formules_xl.evaluer_classeur(x, ev)
    candidates = notes_neant(g, ev, lang)
    wb = openpyxl.load_workbook(x)
    formules = " ".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
                        if isinstance(c.value, str) and c.value.startswith("="))
    retirees = [n for n in candidates if f"'{n}'!" not in formules and f"{n}!" not in formules]
    for n in retirees:
        del wb[n]
    L = PRESENTATION[lang]
    grille = next((ws for ws in wb.worksheets if norm(ws.title).startswith(("grille", "cuadro de analisis", "tabla de analisis"))), None)
    if grille is not None:
        cle = lambda s: re.sub(r"[^a-z0-9]", "", norm(s)).replace("nota", "note")          # noqa: E731
        presentes = {cle(ws.title) for ws in wb.worksheets}
        omises = {cle(n) for n in retirees}
        col = grille.max_column + 1
        entete = next((r for r in range(1, 15) if norm(str(grille.cell(r, 1).value or "")).startswith(("notes", "notas"))), 6)
        grille.cell(entete, col, L[0]).font = copy(grille.cell(entete, 2).font)
        grille.column_dimensions[get_column_letter(col)].width = 34
        for r in range(entete + 1, grille.max_row + 1):
            k = cle(str(grille.cell(r, 1).value or ""))
            if k in omises:
                grille.cell(r, col, L[2])
            elif k in presentes:
                grille.cell(r, col, L[1])
    wb.save(x)
    g.controles_elagage.append((f"[{lang}] classeur adapté aux données : {len(retirees)} note(s) sans montant omise(s), "
                        "signalée(s) dans la grille d'analyse des notes", True, ", ".join(retirees) or "aucune"))
    return retirees


def produire(a, chemin=None, gen=None, journal=lambda *_: None):
    """Production complète (classeurs FR/ES, rapports PDF, annexe balance, contrôle). `chemin(nature, lang)` décide
    où va chaque fichier (le dossier de travail classe XLS et PDF à part). Retourne {"gen", "fichiers": {nature: [..]}}."""
    chemin = chemin or chemins_prefixe(a.sortie)
    g = gen or Generateur(a)
    fichiers, rapports = defaultdict(list), 0
    for lang in a.langues.split(","):
        x = chemin("xlsx", lang)
        feuilles = g.rendre(lang, x)
        fichiers["xlsx"].append(x)
        journal(f"[{lang}] {x} : {len(feuilles)} feuilles, {len(g.a_completer)} cellules à saisir, couverture {dict(g.couverture)}")
        if a.pdf:
            # RAPPORT complet : classeur recalculé, relu en valeurs, remis en page (rapport_def)
            import rapport_def
            xr = rapport_def.recalculer(x)
            if xr:
                pdf = chemin("rapport_pdf", lang)
                m = rapport_def.construire(g, lang, xr, a.logo, pdf)
                # « balance » seul est ambigu en espagnol (« balance activo » = bilan) : la balance GÉNÉRALE
                bal = next((n for n in feuilles if norm(n).startswith("balance genera")), None)
                bp = chemin("balance_pdf", lang)
                mb = def_ge.exporter_pdf(xr, bp, [bal]) if bal else None
                if m:
                    fichiers["rapport_pdf"].append(pdf)
                if mb:
                    fichiers["balance_pdf"].append(bp)
                journal(f"[{lang}] rapport PDF : {pdf if m else 'NON PRODUIT'} · annexe balance : {'produite' if mb else 'NON PRODUITE'}")
                rapports += bool(m)
            else:
                journal(f"[{lang}] rapport complet NON PRODUIT : Excel indisponible pour recalculer le classeur")
        # le CLASSEUR s'adapte aux données, APRÈS le rapport (qui lit le classeur complet) : les notes sans aucun montant
        # sont omises et signalées dans la grille d'analyse (SYSCOHADA révisé) ; choix « conserver » : les 51 feuilles
        if (g.profil.get("interpretations") or {}).get("notes_neant") != "conserver":
            omises = elaguer_notes(g, x, lang)
            g.notes_omises[lang] = omises
            journal(f"[{lang}] notes sans montant omises du classeur : {len(omises)}")
        sorties = [x]
        if getattr(a, "pro", False):
            sorties.append(chemin("pro_xlsx", lang))
            fichiers["xlsx"].append(sorties[-1])
        if getattr(a, "pdf_excel", False):
            for s in sorties:
                if exporter_pdf(s, s[:-5] + ".pdf"):
                    fichiers["excel_pdf"].append(s[:-5] + ".pdf")
    if getattr(a, "sn", False) and fichiers["xlsx"]:
        # MODÈLE SYSCOHADA SYSTÈME NORMAL : même calcul, présentation sur le formulaire officiel OHADA (def_sn) ;
        # ses contrôles (note = poste, formules du formulaire, corrections) rejoignent le rapport de contrôle
        import tempfile
        import openpyxl
        import formules_xl
        import def_sn
        x_fr = next((x for x in fichiers["xlsx"] if "_FR" in os.path.basename(x).upper()), fichiers["xlsx"][0])
        ev = os.path.join(tempfile.mkdtemp(), "ge_eval.xlsx")
        formules_xl.evaluer_classeur(x_fr, ev)
        conserver = (g.profil.get("interpretations") or {}).get("notes_neant") == "conserver"
        sn = def_sn.ProducteurSN(g, openpyxl.load_workbook(ev, data_only=True)).produire(chemin("sn_xlsx", "fr"), conserver)
        fichiers["sn_xlsx"].append(sn["xlsx"])
        g.controles_elagage += [(f"[SYSCOHADA] {n}", ok, det) for n, ok, det in sn["controles"]]
        g.controles_elagage.append((f"[SYSCOHADA] liasse adaptée aux données : {sn['feuilles']} feuilles, "
                                    f"{len(sn['omises'])} omises (grille d'analyse)", True,
                                    "; ".join(f"{t} — {m}" for t, m in sn["omises"])[:900]))
        g.liasse_sn = sn
        journal(f"[SYSCOHADA] {sn['xlsx']} : {sn['feuilles']} feuilles, {len(sn['omises'])} omises")
        if a.pdf:
            wb_ev = openpyxl.load_workbook(sn["evalue"], data_only=True)
            p_sn = chemin("sn_pdf", "fr")
            if def_ge.exporter_pdf(sn["evalue"], p_sn, [w.title for w in wb_ev.worksheets]):
                fichiers["sn_pdf"].append(p_sn)
    import controle_def
    g.controles += g.controles_elagage
    ctrl, ctrl_pdf = chemin("controles_xlsx"), chemin("controles_pdf")
    m2 = controle_def.construire(g, ctrl, ctrl_pdf if a.pdf else None)   # rapport A4 par sections
    fichiers["controles_xlsx"].append(ctrl)
    if a.pdf and m2:
        fichiers["controles_pdf"].append(ctrl_pdf)
    journal(f"Rapport de contrôle (interne, jamais déposé) : {ctrl}" + (f" + {ctrl_pdf}" if a.pdf and m2 else ""))
    if a.pdf and not rapports:
        # repli sans Excel : rapport abrégé de def_ge (mêmes chiffres, même TF5), exporté par LibreOffice si présent
        import tempfile
        tmp = os.path.join(tempfile.mkdtemp(), "rapport.xlsx")
        g.D["tf5_retenu"] = g.tf5_lignes
        feuilles_rapport = def_ge.ecrire_xlsx(g.D, tmp, a.logo)
        ab = chemin("rapport_abrege_pdf")
        if def_ge.exporter_pdf(tmp, ab, feuilles_rapport):
            fichiers["rapport_pdf"].append(ab)
        journal(f"Rapport ABRÉGÉ (repli) : {ab}")
    return {"gen": g, "fichiers": dict(fichiers)}


if __name__ == "__main__":
    main()
