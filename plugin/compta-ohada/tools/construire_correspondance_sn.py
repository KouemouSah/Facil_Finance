# -*- coding: utf-8 -*-
"""Construit la CORRESPONDANCE « ligne du modèle SYSCOHADA système normal -> comptes » à partir du format GE.

Le format GE porte, en colonne A de ses notes, l'expression de comptes de chaque ligne (« BI 411-4112+413 »), prouvée
sur un dossier réel (0 écart inexpliqué). Le modèle SN (formulaire officiel, datas/DSF_Normal_MODELE-PAR-DEFAUT.xlsx)
a les mêmes notes, aux libellés près : chaque ligne de données du SN reçoit l'expression de la ligne GE de même
note dont le libellé est le plus proche (score >= SEUIL). Ce qui ne se rapproche pas est LISTÉ, pour être écrit à la
main dans references/correspondance-sn-manuelle.json (qui prime toujours sur l'automatique).
Sortie : skills/dsf-syscohada/references/correspondance-sn.json
Usage : python tools/construire_correspondance_sn.py [--seuil 0.72]
"""
import argparse
import difflib
import json
import os
import re
import sys
import warnings

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "skills", "dsf-syscohada", "scripts"))
sys.path.insert(0, os.path.join(ROOT, "tools"))
import def_modele as dm  # noqa: E402
from comparer_modeles import code, plat  # noqa: E402

SN = os.path.join(ROOT, "datas", "DSF_Normal_MODELE-PAR-DEFAUT.xlsx")
GE = os.path.join(ROOT, "skills", "dsf-syscohada", "templates", "def_ge_gabarit.xlsx")
SORTIE = os.path.join(ROOT, "skills", "dsf-syscohada", "references", "correspondance-sn.json")
ROLES_DONNEES = ("n", "net_n", "brut")


def entete(ws):
    for r in range(1, 25):
        txt = " ".join(dm.norm(str(ws.cell(r, c).value or "")) for c in range(1, ws.max_column + 1))
        if re.search(r"annee n|exercice n|\bbrut\b|\bnet\b|montant|n-1|annee", txt):
            return r
    return None


def libelle(ws, r, cmax=4):
    for c in range(1, cmax + 1):
        v = ws.cell(r, c).value
        if isinstance(v, str) and len(plat(v)) > 2 and not dm.expr_comptes(v) and not dm.est_ref(v):
            return v.strip()
    return None


CONSIGNE = re.compile(r"^\s*([•·●\-–*]|\(\d\)|n\.?b\b|comm+entaire|libell[eé]s?\s*$|nature des indications|indiquer|justifier|"
                      r"commenter|d[ée]tailler|d[ée]crire|pour (chaque|les)|le cas [ée]ch[ée]ant)", re.I)


def lignes_donnees_sn(ws):
    """Lignes de DONNÉES d'une note SN : libellé présent, cellule N (ou brut) non calculée par une formule."""
    h = entete(ws)
    if not h:
        return None, {}, []
    roles = dm.roles_colonnes(ws, h, h + 2, 2025)
    cn = next((c for c, ro in sorted(roles.items()) if ro in ROLES_DONNEES), None)
    out = []
    if cn:
        for r in range(h + 1, ws.max_row + 1):
            lb = libelle(ws, r, cn - 1 if cn > 1 else 1)
            v = ws.cell(r, cn).value
            # consignes et commentaires du formulaire : du texte à rédiger, jamais une ligne de montant
            if lb and not CONSIGNE.match(lb) and not (isinstance(v, str) and v.startswith("=")) and not dm.ligne_total(lb):
                out.append((r, lb))
    return cn, roles, out


def lignes_ge(ws):
    out = []
    for r in range(1, ws.max_row + 1):
        ex = ws.cell(r, 1).value
        if dm.expr_comptes(ex):
            lb = libelle(ws, r)
            if lb:
                out.append((r, ex.strip(), lb))
    return out


def main():
    import openpyxl
    p = argparse.ArgumentParser()
    p.add_argument("--seuil", type=float, default=0.72)
    a = p.parse_args()
    sn, ge = openpyxl.load_workbook(SN), openpyxl.load_workbook(GE)
    ge_par_code = {code(w.title): w for w in ge.worksheets}
    feuilles, non_rapprochees, stats = {}, {}, {"lignes_sn": 0, "rapprochees": 0}
    for ws in sn.worksheets:
        k = code(ws.title)
        if not k.startswith("note"):
            continue
        cn, roles, lignes = lignes_donnees_sn(ws)
        if not lignes:
            continue
        stats["lignes_sn"] += len(lignes)
        cands = lignes_ge(ge_par_code[k]) if k in ge_par_code else []
        # rapprochement UN POUR UN : chaque ligne GE sert une seule ligne SN (sinon son montant compterait deux fois) ;
        # les paires les plus ressemblantes d'abord
        paires = sorted(((difflib.SequenceMatcher(None, plat(lb), plat(gl)).ratio(), r, lb, gi, ex, gl)
                         for r, lb in lignes for gi, (_, ex, gl) in enumerate(cands)), reverse=True)
        pris, ge_pris = {}, set()
        for sc, r, lb, gi, ex, gl in paires:
            if sc < a.seuil:
                break
            if str(r) in pris or gi in ge_pris:
                continue
            pris[str(r)] = {"expr": ex, "libelle_sn": lb, "libelle_ge": gl, "score": round(sc, 2)}
            ge_pris.add(gi)
        for r, lb in lignes:
            if str(r) not in pris:
                non_rapprochees.setdefault(k, []).append({"ligne": r, "libelle": lb})
        stats["rapprochees"] += len(pris)
        feuilles[k] = {"feuille": ws.title, "colonne_n": cn, "roles": {str(c): ro for c, ro in roles.items()},
                       "lignes": pris}
    json.dump({"origine": "format GE (gabarit) rapproché par libellé ; manuel : correspondance-sn-manuelle.json",
               "seuil": a.seuil, "feuilles": feuilles, "non_rapprochees": non_rapprochees},
              open(SORTIE, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"lignes de données SN : {stats['lignes_sn']} · rapprochées du format GE : {stats['rapprochees']}")
    for k, ls in non_rapprochees.items():
        print(f"  {k:12s} {len(ls):3d} à écrire : " + " | ".join(x["libelle"][:34] for x in ls[:6]))
    print(f"  → {os.path.relpath(SORTIE, ROOT)}")


if __name__ == "__main__":
    main()
