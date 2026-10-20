# -*- coding: utf-8 -*-
"""Extrait le squelette de la liasse DSF (Bilan + Compte de résultat + index des Notes)
depuis le modèle fourni dans datas/, et écrit une référence markdown GROUNDED.

But : fournir la structure réelle (REF | libellé | note | formule) sur laquelle
s'appuie le mapping comptes→DSF et le skill d'assemblage. Ne réinvente rien :
lit le fichier officiel.

Usage : python tools/extract_dsf_skeleton.py [chemin_liasse.xlsx]
"""
import os
import re
import sys
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT_LIASSE = os.path.join(ROOT, "datas", "DSF_Normal_VEROUILLE-09-03-2021.xlsx")
OUT = os.path.join(ROOT, "skills", "dsf-syscohada", "references", "squelette-dsf.md")


def clean(v):
    if v is None:
        return ""
    return re.sub(r"\s+", " ", str(v).replace("\xa0", " ")).strip()


def is_ref(a):
    return bool(a) and isinstance(a, str) and 1 <= len(a.strip()) <= 3 and a.strip().isupper() and a.strip() != "REF"


def extract_statement(ws):
    rows = []
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if is_ref(a):
            lib = clean(ws.cell(r, 2).value)
            note = clean(ws.cell(r, 4).value)
            e = ws.cell(r, 5).value
            formula = clean(e) if isinstance(e, str) and e.startswith("=") else ""
            rows.append((clean(a), lib, note, formula))
    return rows


def main():
    liasse = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_LIASSE
    wb = load_workbook(liasse, data_only=False)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    L = []
    L.append("# Squelette de la liasse DSF (extrait du modèle officiel)")
    L.append("")
    L.append(f"> Source : `{os.path.basename(liasse)}` — {len(wb.sheetnames)} feuilles. "
             "Extraction automatique (grounded), ne pas éditer à la main : régénérer via "
             "`tools/extract_dsf_skeleton.py`.")
    L.append("")
    L.append("**Principe** : les Notes sont les points de saisie ; le Bilan et le Compte de "
             "résultat **agrègent les Notes par formules** (colonne « Formule »). Le mapping "
             "comptes→DSF (`mapping-comptes-dsf.md`) alimente les Notes.")
    L.append("")

    bilan = next((s for s in wb.sheetnames if "BILAN" in s.upper()), None)
    cr = next((s for s in wb.sheetnames if "RESULTAT" in s.upper()), None)

    for title, sheet in [("BILAN", bilan), ("COMPTE DE RÉSULTAT", cr)]:
        if not sheet:
            continue
        rows = extract_statement(wb[sheet])
        L.append(f"## {title} (feuille « {sheet} ») — {len(rows)} postes")
        L.append("")
        L.append("| REF | Libellé | Note | Formule (source) |")
        L.append("|---|---|---|---|")
        for ref, lib, note, formula in rows:
            L.append(f"| {ref} | {lib} | {note} | `{formula}` |")
        L.append("")

    # index des notes
    L.append("## Index des feuilles / Notes")
    L.append("")
    L.append("| # | Feuille | Titre (le cas échéant) |")
    L.append("|---|---|---|")
    for i, s in enumerate(wb.sheetnames, 1):
        ws = wb[s]
        titre = ""
        for r in range(1, 9):
            for c in range(1, 4):
                v = clean(ws.cell(r, c).value)
                if v and ("NOTE" in v.upper() and len(v) < 60) or (v and any(k in v.upper() for k in ("BILAN", "RESULTAT", "FLUX", "GARDE", "STATIST"))):
                    titre = v
                    break
            if titre:
                break
        L.append(f"| {i} | {s.strip()} | {titre} |")
    L.append("")
    wb.close()
    with open(OUT, "w", encoding="utf-8") as f:
        f.write("\n".join(L))
    print("Écrit :", os.path.relpath(OUT, ROOT))
    print(f"Postes Bilan + CR extraits, {len(wb.sheetnames)} feuilles indexées.")


if __name__ == "__main__":
    main()
