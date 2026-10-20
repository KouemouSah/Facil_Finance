# -*- coding: utf-8 -*-
"""Compare deux modèles de liasse : le SYSCOHADA révisé SYSTÈME NORMAL (formulaire officiel embarqué, DGI Cameroun)
et le FORMAT GE (gabarit tiré de la DEF déposée). Feuille par feuille (par code de note), puis ligne par ligne sur
les notes communes (libellés normalisés). Sortie : JSON (pour la documentation) et résumé à l'écran.
Usage : python tools/comparer_modeles.py [--json sortie.json]
"""
import argparse
import json
import os
import re
import sys
import unicodedata
import warnings

warnings.filterwarnings("ignore")
sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SN = os.path.join(ROOT, "datas", "DSF_Normal_MODELE-PAR-DEFAUT.xlsx")
GE = os.path.join(ROOT, "skills", "dsf-syscohada", "templates", "def_ge_gabarit.xlsx")


def plat(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def code(titre):
    """Code de feuille comparable : « NOTE  3C » -> note 3c ; « CO1-NOTE 3C » -> c1 note 3c ; sinon titre aplati."""
    t = plat(titre).replace("co1", "c1")
    m = re.match(r"(c\d )?note (\d+ ?[a-z]?( bis)?)", t)
    if m:
        return ((m.group(1) or "") + "note " + m.group(2).replace(" ", "")).strip()
    return {"bilan paysage": "bilan", "compte de resultat": "compte resultat", "compte resultat": "compte resultat",
            "tableau des flux de tresorerie": "flux tresorerie", "flux tresorerie": "flux tresorerie",
            "page de garde": "page garde"}.get(t, t)


def libelles(ws):
    out = []
    for r in range(1, ws.max_row + 1):
        for c in range(1, 4):
            v = ws.cell(r, c).value
            # ni un code de poste (« AD »), ni une expression de comptes (« BI 411-4112+413 ») : le LIBELLÉ
            if isinstance(v, str) and len(plat(v)) > 3 and not re.fullmatch(r"[A-Z]{2}", v.strip()) \
                    and not re.match(r"^[A-Z]{2}\b[\s\d+\-,~]*\d", v.strip()):
                out.append(plat(v)[:60])
                break
    return out


def main():
    import openpyxl
    p = argparse.ArgumentParser()
    p.add_argument("--json")
    a = p.parse_args()
    sn, ge = openpyxl.load_workbook(SN), openpyxl.load_workbook(GE)
    csn = {code(w.title): w for w in sn.worksheets}
    cge = {code(w.title): w for w in ge.worksheets}
    communs = sorted(set(csn) & set(cge))
    seul_sn = sorted(set(csn) - set(cge))
    seul_ge = sorted(set(cge) - set(csn))
    lignes = {}
    for k in communs:
        a_, b_ = libelles(csn[k]), libelles(cge[k])
        sa, sb = set(a_), set(b_)
        lignes[k] = {"sn": len(sa), "ge": len(sb), "identiques": len(sa & sb),
                     "seulement_sn": sorted(sa - sb)[:8], "seulement_ge": sorted(sb - sa)[:8]}
    res = {"feuilles_sn": len(sn.worksheets), "feuilles_ge": len(ge.worksheets), "communes": communs,
           "seulement_sn": seul_sn, "seulement_ge": seul_ge, "lignes": lignes}
    print(f"SN : {len(sn.worksheets)} feuilles · GE : {len(ge.worksheets)} feuilles · communes : {len(communs)}")
    print("seulement SN :", ", ".join(seul_sn))
    print("seulement GE :", ", ".join(seul_ge))
    for k, v in lignes.items():
        print(f"  {k:20s} libellés SN {v['sn']:3d} · GE {v['ge']:3d} · identiques {v['identiques']:3d}")
    if a.json:
        json.dump(res, open(a.json, "w", encoding="utf-8"), ensure_ascii=False, indent=1)


if __name__ == "__main__":
    main()
