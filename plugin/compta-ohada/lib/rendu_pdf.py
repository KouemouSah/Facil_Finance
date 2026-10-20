# -*- coding: utf-8 -*-
"""Rendu PDF NATIF d'un classeur mis en page (PyMuPDF) — pour se passer d'Excel.

Chaque feuille demandée devient une ou plusieurs pages A4 (orientation de la feuille) :
  - zone d'impression (sinon zone utilisée), lignes masquées ignorées ;
  - cellules fusionnées (rowspan / colspan), largeurs de colonnes proportionnelles ;
  - fonds, gras, italique, taille et couleur de police, alignements, bordures haute / basse, texte riche ;
  - formats : montants « 1 234 567 », négatifs entre parenthèses, zéro « – » ; pourcentages « 12,5 % » ; dates ;
  - images (logos) à leur ligne d'ancrage ;
  - « ajuster à une page en hauteur » : la police est réduite jusqu'à ce que la feuille tienne sur une page ;
  - pied de page de la feuille (gauche / centre / droite, &P et &N = page n / N du document).
Usage : classeur_vers_pdf(xlsx, pdf, feuilles=None) -> True / False
Source unique : compta-ohada/lib/rendu_pdf.py (copies synchronisées par tools/sync_lib.py — ne pas éditer les copies).
"""
import datetime
import html
import os
import re
import tempfile

A4 = (595.28, 841.89)
MARGE = 34.0          # 12 mm


# couleurs du thème Office par défaut (indice « theme » d'openpyxl) : sans elles, un chiffre blanc sur bandeau bleu
# (couleur de thème 0) serait rendu noir, donc invisible
THEME = {0: "FFFFFF", 1: "000000", 2: "E7E6E6", 3: "44546A", 4: "4472C4", 5: "ED7D31", 6: "A5A5A5", 7: "FFC000",
         8: "5B9BD5", 9: "70AD47"}


def _rgb(c):
    if c is None:
        return None
    try:
        if getattr(c, "type", None) == "theme" and c.theme is not None:
            return "#" + THEME.get(int(c.theme), "000000")
        v = c.rgb
    except Exception:
        return None
    if isinstance(v, str) and len(v) >= 6 and v not in ("00000000",):
        return "#" + v[-6:]
    return None


def _format(v, fmt):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "VRAI" if v else "FAUX"
    if isinstance(v, (datetime.date, datetime.datetime)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, (int, float)):
        f = fmt or "General"
        if "%" in f:
            dec = len(re.search(r"0\.(0+)%", f).group(1)) if re.search(r"0\.(0+)%", f) else 0
            return f"{v * 100:,.{dec}f} %".replace(",", " ").replace(".", ",")
        if "#,##0" in f or f.startswith("0"):
            if abs(v) < 0.5 and '"–"' in f:
                return "–"
            dec = len(re.search(r"0\.(0+)", f).group(1)) if re.search(r"#,##0\.(0+)|^0\.(0+)", f) else 0
            s = f"{abs(v):,.{dec}f}".replace(",", " ").replace(".", ",")
            return f"({s})" if v < 0 and "(" in f else ("-" + s if v < 0 else s)
        if float(v).is_integer():
            return f"{int(v)}"
        return f"{v:.2f}".replace(".", ",")
    return str(v)


def _texte_riche(v):
    """Texte riche openpyxl -> HTML (gras / italique par segment)."""
    if type(v).__name__ != "CellRichText":
        return None
    out = []
    for b in v:
        if isinstance(b, str):
            out.append(html.escape(b))
        else:
            t = html.escape(b.text)
            f = b.font
            if getattr(f, "b", False):
                t = f"<b>{t}</b>"
            if getattr(f, "i", False):
                t = f"<i>{t}</i>"
            out.append(t)
    return "".join(out)


def _zone(ws):
    from openpyxl.utils import range_boundaries
    pa = ws.print_area
    if pa:
        parts = pa if isinstance(pa, (list, tuple)) else str(pa).split(",")
        bs = [range_boundaries(p.split("!")[-1].replace("$", "")) for p in parts]
        return min(b[0] for b in bs), min(b[1] for b in bs), max(b[2] for b in bs), max(b[3] for b in bs)
    return 1, 1, ws.max_column, ws.max_row


def feuille_html(ws, dossier_images, taille=9.0, largeur_px=527.0):
    """HTML d'une feuille ; les images sont copiées dans dossier_images. Largeurs ABSOLUES par cellule : le moteur
    HTML de PyMuPDF ignore les largeurs en pourcentage et celles des <col>, il respecte celles des cellules."""
    from openpyxl.utils import get_column_letter
    c0, r0, c1, r1 = _zone(ws)
    larg = [(ws.column_dimensions[get_column_letter(c)].width or 8.43) for c in range(c0, c1 + 1)]
    cache = [c for c in range(c0, c1 + 1) if ws.column_dimensions[get_column_letter(c)].hidden]
    tot = sum(w for i, w in enumerate(larg) if (c0 + i) not in cache) or 1
    px = {c0 + i: (0 if (c0 + i) in cache else w / tot * largeur_px) for i, w in enumerate(larg)}
    # lignes réellement émises : ni masquées, et non vides OU prises dans une fusion verticale (sinon la fusion
    # déborderait sur la ligne suivante et décalerait le tableau)
    dans_fusion_v = set()
    for rg in ws.merged_cells.ranges:
        if rg.max_row > rg.min_row:
            dans_fusion_v.update(range(rg.min_row, rg.max_row + 1))
    emises_set = {r for r in range(r0, r1 + 1) if not ws.row_dimensions[r].hidden}
    fusion, couvert = {}, set()
    for rg in ws.merged_cells.ranges:
        if rg.min_row > r1 or rg.min_col > c1:
            continue
        rs = sum(1 for r in range(rg.min_row, min(rg.max_row, r1) + 1) if r in emises_set) or 1
        fusion[(rg.min_row, rg.min_col)] = (rs, min(rg.max_col, c1) - rg.min_col + 1)
        for r in range(rg.min_row, rg.max_row + 1):
            for c in range(rg.min_col, rg.max_col + 1):
                if (r, c) != (rg.min_row, rg.min_col):
                    couvert.add((r, c))
    images = {}
    for k, im in enumerate(getattr(ws, "_images", [])):
        try:
            ligne = im.anchor._from.row + 1
            nom = f"img_{id(ws)}_{k}.png"
            chemin = os.path.join(dossier_images, nom)
            if not os.path.exists(chemin):          # les données d'une image openpyxl ne se lisent QU'UNE fois :
                data = im._data() if callable(getattr(im, "_data", None)) else None   # on garde le fichier
                if data:
                    open(chemin, "wb").write(data)
            if os.path.exists(chemin):
                ext = getattr(im.anchor, "ext", None)           # taille voulue = celle de l'ancrage (EMU), sinon l'image
                h = (ext.height / 9525) if ext is not None and getattr(ext, "height", None) else (getattr(im, "height", 60) or 60)
                images.setdefault(max(ligne, r0), []).append((nom, h * 0.75 * taille / 9.0))
        except Exception:
            continue
    nvis = c1 - c0 + 1 - len(cache)
    lignes = []
    for r in range(r0, r1 + 1):
        for nom, h in images.get(r, []):
            lignes.append(f'<tr><td colspan="{nvis}" style="width:{largeur_px - 6:.0f}px;text-align:center;border:none">'
                          f'<img src="{nom}" style="height:{h:.0f}px"></td></tr>')
        if ws.row_dimensions[r].hidden:
            continue
        hauteur = ws.row_dimensions[r].height or 15
        if r not in dans_fusion_v and all(ws.cell(r, c).value in (None, "") for c in range(c0, c1 + 1)):
            # ligne vide hors fusion : un espaceur à la hauteur d'Excel (la mise en page du modèle en dépend)
            lignes.append(f'<tr><td colspan="{nvis}" style="width:{largeur_px - 6:.0f}px;height:{hauteur * 0.85 * taille / 9.0:.0f}px"></td></tr>')
            continue
        cells, vide, deborde = [], True, set()
        for c in range(c0, c1 + 1):
            if c in cache or (r, c) in couvert or c in deborde:
                continue
            cell = ws.cell(r, c)
            rs, cs = fusion.get((r, c), (1, 1))
            v = cell.value
            al0 = cell.alignment
            if (r, c) not in fusion and isinstance(v, str) and v and not (al0 is not None and (al0.wrap_text or al0.horizontal in ("right", "center"))):
                # comme Excel : un texte non renvoyé à la ligne DÉBORDE sur les cellules vides qui le suivent
                k = c + 1
                while k <= c1 and ws.cell(r, k).value in (None, "") and (r, k) not in couvert and (r, k) not in fusion and k not in cache:
                    deborde.add(k)
                    k += 1
                cs = 1 + len([x for x in deborde if c < x < k])
            contenu = _texte_riche(v)
            if contenu is None:
                contenu = html.escape(_format(v, cell.number_format)).replace("\n", "<br>")
            if contenu:
                vide = False
            st = []
            f = cell.font
            if f is not None:
                if f.b:
                    st.append("font-weight:bold")
                if f.i:
                    st.append("font-style:italic")
                if f.sz:
                    st.append(f"font-size:{float(f.sz) * taille / 9.5:.1f}pt")
                col = _rgb(f.color)
                if col and col != "#000000":
                    st.append(f"color:{col}")
            fl = cell.fill
            if fl is not None and fl.fill_type == "solid":
                fond = _rgb(fl.fgColor)
                if fond:
                    st.append(f"background-color:{fond}")
            al = cell.alignment
            h = al.horizontal if al is not None else None
            if h in ("right", "center", "left"):
                st.append(f"text-align:{h}")
            elif isinstance(v, (int, float)) and not isinstance(v, bool):
                st.append("text-align:right")
            b = cell.border
            for cote in ("top", "bottom"):
                s = getattr(b, cote, None) if b is not None else None
                if s is not None and s.style:
                    ep = {"thin": 0.5, "hair": 0.3, "medium": 1.2, "thick": 2, "double": 1.5}.get(s.style, 0.5)
                    st.append(f"border-{cote}:{ep}pt solid {_rgb(s.color) or '#999999'}")
            attrs = (f' rowspan="{rs}"' if rs > 1 else "") + (f' colspan="{cs}"' if cs > 1 else "")
            w = sum(px.get(x, 0) for x in range(c, c + cs))
            st.insert(0, f"width:{max(w - 6, 4):.0f}px")          # largeur absolue (padding 3 px de chaque côté)
            # comme Excel « ajuster » : un montant (ou une cellule marquée shrink_to_fit) ne déborde jamais de sa
            # colonne — la police se réduit (≈ 0,56 × corps par caractère), jamais sous 4,5 pt
            brut = re.sub(r"<[^>]+>", "", contenu)
            num_ou_ajuste = isinstance(v, (int, float)) and not isinstance(v, bool) or (al is not None and al.shrink_to_fit)
            if brut and num_ou_ajuste and "<br>" not in contenu:
                corps = float(f.sz) * taille / 9.5 if f is not None and f.sz else taille
                besoin = (w - 7) / (len(brut) * 0.56)
                if besoin < corps:
                    st = [x for x in st if not x.startswith("font-size")] + [f"font-size:{max(besoin, 4.5):.1f}pt"]
            cells.append(f'<td{attrs} style="{";".join(st)}">{contenu}</td>')
        lignes.append(f"<tr>{''.join(cells)}</tr>")
    for r in sorted(k for k in images if k > r1):          # image ancrée sous la dernière ligne remplie : en fin de feuille
        for nom, h in images[r]:
            lignes.append(f'<tr><td colspan="{nvis}" style="width:{largeur_px - 6:.0f}px;text-align:center;border:none">'
                          f'<img src="{nom}" style="height:{h:.0f}px"></td></tr>')
    return f'<table>{"".join(lignes)}</table>', tot


CSS = """
body { font-family: sans-serif; font-size: %(t).1fpt; color: #1a1a1a; }
table { width: 100%%; border-collapse: collapse; table-layout: fixed; }
td { padding: %(p).2fpx 3px; vertical-align: top; border: none; line-height: 1.15; }
"""


def _pied(ws):
    try:
        f = ws.oddFooter
        return tuple((getattr(f, p).text or "") for p in ("left", "center", "right"))
    except Exception:
        return ("", "", "")


def classeur_vers_pdf(xlsx, pdf, feuilles=None):
    try:
        try:
            import pymupdf as fitz            # nom actuel de PyMuPDF
        except ImportError:
            import fitz                       # ancien nom (versions < 1.24.3)
        import openpyxl
    except ImportError:
        return False
    wb = openpyxl.load_workbook(xlsx, rich_text=True)          # texte riche : gras / italique par segment
    noms = feuilles or [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
    tmpdir = tempfile.mkdtemp(prefix="rendu_")
    brut = os.path.join(tmpdir, "mise_en_page.pdf")     # mise en page, puis pieds de page -> PDF final (jamais
    writer = fitz.DocumentWriter(brut)                  # de remplacement d'un fichier encore ouvert, Windows le refuse)
    pieds = []
    for nom in noms:
        if nom not in wb.sheetnames:
            continue
        ws = wb[nom]
        paysage = ws.page_setup.orientation == "landscape"
        page = fitz.Rect(0, 0, *(A4[::-1] if paysage else A4))
        zone = page + (MARGE, MARGE, -MARGE, -MARGE - 14)
        une_page = bool(ws.page_setup.fitToHeight == 1 and ws.sheet_properties.pageSetUpPr.fitToPage)
        utile = zone.width * 0.985                        # marge de sécurité : bordures et arrondis de largeur
        corps, largeur = feuille_html(ws, tmpdir, 9.0, utile)
        arch = fitz.Archive(tmpdir)                       # APRÈS l'écriture des images : l'archive fige son contenu
        # police de base : Excel « ajuste à la largeur » en réduisant l'échelle ; ici, 1 caractère ≈ 5,6 pt à 9 pt
        base = max(5.5, min(9.0, 9.0 * (zone.width / 5.6) / max(largeur * 1.0, 1) * 0.9))
        # Excel réduit jusqu'à 10 % pour tenir sur une page ; ici jusqu'à 30 % (au-delà, illisible)
        tailles = [base] + ([base * k for k in (0.92, 0.85, 0.78, 0.7, 0.62, 0.55, 0.48, 0.42, 0.36, 0.3)] if une_page else [])
        choisie = tailles[-1]
        for t in tailles:                         # la plus grande police qui tient (sur une page si demandé)
            corps_t, _ = feuille_html(ws, tmpdir, t, utile)
            s = fitz.Story(html=corps_t, user_css=CSS % {"t": t, "p": 1.5 * t / 9}, archive=arch)
            reste, _ = s.place(zone)
            if not une_page or not reste:
                choisie = t
                break
        corps_t, _ = feuille_html(ws, tmpdir, choisie, utile)
        story = fitz.Story(html=corps_t, user_css=CSS % {"t": choisie, "p": 1.5 * choisie / 9}, archive=arch)
        reste = True
        while reste:
            dev = writer.begin_page(page)
            reste, _ = story.place(zone)
            story.draw(dev)
            writer.end_page()
            pieds.append((_pied(ws), nom))
    writer.close()
    # pieds de page : &P / &N sur le document entier, &A = nom de la feuille
    doc = fitz.open(brut)
    n = doc.page_count
    for i, page in enumerate(doc):
        (g, c, d), nom = pieds[i] if i < len(pieds) else (("", "", ""), "")
        y = page.rect.height - MARGE + 8
        for texte, x, align in ((g, MARGE, 0), (c, page.rect.width / 2, 1), (d, page.rect.width - MARGE, 2)):
            t = texte.replace("&P", str(i + 1)).replace("&N", str(n)).replace("&A", nom)
            t = re.sub(r"&[A-Za-z\"0-9,\-]+", "", t)
            if not t:
                continue
            w = fitz.get_text_length(t, fontsize=7)
            xx = x if align == 0 else (x - w / 2 if align == 1 else x - w)
            page.insert_text((xx, y), t, fontsize=7, color=(0.5, 0.5, 0.5))
    if os.path.exists(pdf):
        os.remove(pdf)
    doc.save(pdf, garbage=3, deflate=True)
    doc.close()
    return os.path.exists(pdf)
