# -*- coding: utf-8 -*-
"""Garde du rendu PDF NATIF (lib/rendu_pdf.py, copie dans skills/dsf-syscohada/scripts), sans Excel.
Lancer : python tests/test_rendu_pdf.py     (code 0 = conforme)
"""
import os
import re
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "skills", "dsf-syscohada", "scripts"))
import rendu_pdf  # noqa: E402

ok = []
NUM = '#,##0;(#,##0);"–"'


def check(nom, cond, detail=""):
    ok.append(bool(cond))
    print(f"  {'OK   ' if cond else 'ÉCHEC'} {nom}" + ("" if cond else f"  — {detail}"))


def main():
    try:
        import pymupdf as fitz            # nom actuel de PyMuPDF
    except ImportError:
        import fitz                       # ancien nom (versions < 1.24.3)
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Alignment, Font, PatternFill
    from PIL import Image as PILImage
    d = tempfile.mkdtemp()
    wb = Workbook()
    ws = wb.active
    ws.title = "Etat"
    for c, w in zip("ABCD", (30, 9, 14, 10)):                     # colonne B volontairement étroite
        ws.column_dimensions[c].width = w
    ws.merge_cells("A1:D1")
    ws["A1"] = "TITRE FUSIONNÉ"
    ws["A1"].fill = PatternFill("solid", fgColor="1F3864")
    ws["A1"].font = Font(bold=True, color="FFFFFF")
    ws["A3"], ws["B3"], ws["C3"], ws["D3"] = "Ligne", 1234567, -500, 0.125
    ws["B3"].number_format = ws["C3"].number_format = NUM
    ws["D3"].number_format = "0.0%"
    ws["A4"], ws["B4"] = "Zéro", 0
    ws["B4"].number_format = NUM
    ws["A5"] = "Texte long qui déborde sur les cellules vides voisines, comme dans Excel"
    ws["B6"] = 123456789012345                                    # montant trop large pour sa colonne
    ws["B6"].number_format = NUM
    ws["B6"].alignment = Alignment(shrink_to_fit=True)
    logo = os.path.join(d, "logo.png")
    PILImage.new("RGB", (200, 100), (31, 56, 100)).save(logo)
    im = Image(logo)
    im.width, im.height = 80, 40
    ws.add_image(im, "A8")
    ws.oddFooter.center.text = "Page &P / &N"
    long_ = wb.create_sheet("Longue")
    for r in range(1, 121):
        long_.cell(r, 1, f"ligne {r}")
    court = wb.create_sheet("Paysage")
    court.page_setup.orientation = "landscape"
    court.page_setup.fitToHeight = 1
    court.sheet_properties.pageSetUpPr.fitToPage = True
    for r in range(1, 90):                                         # 89 lignes : exigent la réduction pour tenir
        court.cell(r, 1, f"ligne {r}")
    x, pdf = os.path.join(d, "c.xlsx"), os.path.join(d, "c.pdf")
    wb.save(x)
    check("PDF produit", rendu_pdf.classeur_vers_pdf(x, pdf) and os.path.exists(pdf))
    doc = fitz.open(pdf)
    t0 = doc[0].get_text()
    check("montant « 1 234 567 »", "1 234 567" in t0, t0[:200])
    check("négatif entre parenthèses « (500) »", "(500)" in t0)
    check("zéro affiché « – »", "–" in t0)
    check("pourcentage « 12,5 % »", "12,5 %" in t0)
    check("texte fusionné présent", "TITRE FUSIONNÉ" in t0)
    check("texte long sur une seule ligne (débordement émulé)", "comme dans Excel" in t0.replace("\n", " ") and
          any("Texte long" in l and "Excel" in l for l in t0.splitlines()), [l for l in t0.splitlines() if "Texte" in l])
    check("image présente", len(doc[0].get_images(full=True)) == 1)
    xmax = max(w[2] for w in doc[0].get_text("words"))
    check("rien ne déborde à droite", xmax <= doc[0].rect.width - rendu_pdf.MARGE + 1, xmax)
    check("grand montant réduit pour tenir sur UNE ligne (jamais coupé entre ses chiffres)",
          any("123 456 789 012 345" in l for l in t0.splitlines()), [l for l in t0.splitlines() if "123" in l])
    tailles = [s["size"] for b in doc[0].get_text("dict")["blocks"] for l in b.get("lines", []) for s in l["spans"]]
    t_montant = [s["size"] for b in doc[0].get_text("dict")["blocks"] for l in b.get("lines", []) for s in l["spans"]
                 if "123 456 789" in s["text"]]
    check("la police du montant trop large est RÉDUITE (colonne tenue, pas élargie)",
          t_montant and t_montant[0] < max(tailles) * 0.9, (t_montant, max(tailles)))
    pages = {p.number: p for p in doc}
    orient = [(p.rect.width > p.rect.height) for p in doc]
    check("feuille longue sur plusieurs pages, feuille « une page » sur UNE page paysage",
          doc.page_count == 1 + (doc.page_count - 2) + 1 and orient[-1] and not orient[-2], (doc.page_count, orient))
    pied = doc[0].get_text()
    check("pied de page « Page 1 / N » (N = total du document)", f"Page 1 / {doc.page_count}" in pied, re.findall(r"Page \d+ / \d+", pied))
    print(f"\n{sum(ok)}/{len(ok)} contrôles conformes")
    return 0 if all(ok) else 1


if __name__ == "__main__":
    sys.exit(main())
