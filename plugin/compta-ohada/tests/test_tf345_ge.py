# -*- coding: utf-8 -*-
"""Garde du lecteur de grand livre (lib/grand_livre.py) et de TF3 / TF4 / TF5 (lib/tf345_ge.py).

Grand livre SYNTHÉTIQUE écrit sous deux présentations : « par compte » (xlsx : en-têtes de compte, report
à nouveau, totaux, soldes progressifs, dates série) et « à plat » (csv : compte sur chaque ligne, dates
jj/mm/aaaa, montants texte « 1 234,56 »). Mêmes écritures -> mêmes TF. Valeurs ATTENDUES calculées à la main.
Lancer : python tests/test_tf345_ge.py     (code 0 = conforme)
"""
import csv
import datetime
import os
import re
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "lib"))
import grand_livre as g  # noqa: E402
import tf345_ge as t  # noqa: E402

# (date, journal, pièce, libellé, [(compte, tiers, débit, crédit)])
E = [
    ("2025-01-10", "VTE", "V1", "FACTURE TRAVAUX JANVIER", [("41110000", "", 1_150_000, 0), ("70510000", "", 0, 1_000_000),
                                                           ("44330000", "", 0, 150_000)]),
    ("2025-01-31", "ODP", "P0", "PAIE JANVIER", [("66110000", "", 400_000, 0), ("66210000", "", 900_000, 0),
                                                ("44720000", "", 0, 40_000), ("44721000", "", 0, 120_000), ("42200000", "", 0, 1_140_000)]),
    ("2025-02-15", "VTE", "V2", "FACTURE TRAVAUX ETAT", [("41141000", "", 2_000_000, 0), ("70510000", "", 0, 2_000_000)]),
    ("2025-03-05", "ACH", "A1", "LOYER BUREAU SCI ALPHA MARS", [("62220000", "F100", 500_000, 0), ("44520000", "", 75_000, 0),
                                                               ("40110000", "", 0, 575_000)]),
    ("2025-04-05", "ACH", "A2", "LOYER BUREAU SCI ALPHA AVRIL", [("62220000", "F100", 500_000, 0), ("44520000", "", 75_000, 0),
                                                                ("40110000", "", 0, 575_000)]),
    ("2025-04-20", "ACH", "S1", "SOUS TRAITANCE BETA CONSTRUCT", [("62100000", "F200", 1_000_000, 0), ("40110000", "", 0, 1_000_000)]),
    ("2025-05-12", "CAI", "C1", "PETITS FRAIS DIVERS", [("62800000", "", 30_000, 0), ("57110000", "", 0, 30_000)]),
    ("2025-06-15", "BQ", "B1", "PAIEMENT IVA MAI", [("44330000", "", 60_000, 0), ("52110000", "", 0, 60_000)]),
    ("2025-07-15", "BQ", "B2", "REVERSEMENT IRPF LOCAUX", [("44720000", "", 40_000, 0), ("52110000", "", 0, 40_000)]),
    ("2025-12-31", "ODG", "CL", "CLOTURE 2025", [("70510000", "", 3_000_000, 0), ("66110000", "", 0, 400_000),
                                                ("66210000", "", 0, 900_000), ("62100000", "", 0, 1_000_000),
                                                ("62220000", "", 0, 1_000_000), ("62800000", "", 0, 30_000), ("13900000", "", 330_000, 0)]),
]
OUV = {"52110000": 5_000_000, "10100000": -5_000_000}
LIB = {"44720000": "IMPOT SUR SALAIRE IRPF LOCAUX", "44721000": "IMPOT SUR SALAIRE IRPF EXPATRIES",
       "62220000": "LOCATIONS BUREAUX", "62100000": "SOUS-TRAITANCE", "62800000": "DIVERS"}
SERIE = datetime.date(1899, 12, 30)


def par_compte(path):
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active
    ws.append(["", "Grand-livre général", 46267.0, 0.6561])          # en-tête d'édition daté (piège)
    ws.append(["Date", "Jal", "Pièce", "Libellé", "Tiers", "", "", "", "", "Débit", "Crédit", "Solde D", "Solde C"])
    comptes = sorted({c for e in E for c, *_ in e[4]} | set(OUV))
    for c in comptes:
        ws.append([c, LIB.get(c, "COMPTE " + c)])
        s = OUV.get(c, 0.0)
        if c in OUV:
            ws.append(["Report à nouveau", "", max(s, 0), max(-s, 0), max(s, 0), max(-s, 0)])
        for date, jal, pce, lib, ls in E:
            for cc, tiers, d, cr in ls:
                if cc == c:
                    s += d - cr
                    serie = (datetime.date.fromisoformat(date) - SERIE).days
                    ws.append([serie, jal, pce, lib, tiers, "SIG", "", "", "D", d, cr, max(s, 0), max(-s, 0)])
        ws.append([c, "Total compte au", 46022, 0, 0, max(s, 0), max(-s, 0)])   # ligne de synthèse datée (piège)
    wb.save(path)


def a_plat(path):
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Date", "Compte", "Journal", "Pièce", "Libellé", "Tiers", "Débit", "Crédit"])
        for date, jal, pce, lib, ls in E:
            dd = datetime.date.fromisoformat(date).strftime("%d/%m/%Y")
            for cc, tiers, d, cr in ls:
                fr = lambda x: f"{x:,.2f}".replace(",", " ").replace(".", ",") if x else ""
                w.writerow([dd, cc, jal, pce, lib, tiers, fr(d), fr(cr)])


ok = []


def check(nom, cond, detail=""):
    ok.append(bool(cond))
    print(f"  {'OK   ' if cond else 'ÉCHEC'} {nom}" + ("" if cond else f"  — {detail}"))


def verifier(nom, path, avec_ouverture):
    print(f"== {nom}")
    gl = g.lire(path)
    n_attendu = sum(len(e[4]) for e in E)
    check(f"{n_attendu} lignes d'écriture (en-tête d'édition et totaux datés écartés)", len(gl["lignes"]) == n_attendu,
          (len(gl["lignes"]), gl["detection"]))
    check("pièce de clôture détectée : CL", g.pieces_de_cloture(gl) == ["CL"], g.pieces_de_cloture(gl))
    if avec_ouverture:
        check("reports à nouveau lus", gl["ouvertures"] == {"52110000": 5_000_000, "10100000": -5_000_000}, gl["ouvertures"])
        check("solde progressif vérifié", (gl["detection"].get("solde_progressif_verifie") or 0) >= 0.99, gl["detection"])
    r = t.calculer(gl, libelles=LIB)
    r3, r4, r5 = r["tf3"], r["tf4"], r["tf5"]
    # TF3 : IVA 150 000 (janvier) -> base 1 000 000 ; déductible 150 000 ; paiement B1 60 000
    #       net = 150 000 - 150 000 - 60 000 = -60 000 -> 29 = 0, crédit 30 = 60 000 ; CA 3 000 000 -> sans IVA 2 000 000
    check("TF3 base 1 000 000, IVA 150 000", r3["base_totale"] == 1_000_000 and r3["iva_theorique"] == 150_000, r3["base_totale"])
    check("TF3 déductible 150 000 (ligne 15 bis)", r3["iva_deductible"] == 150_000 and r3["ligne_15bis"] == 150_000)
    check("TF3 paiement au Trésor B1 = 60 000 (ligne 28)", r3["ligne_28"] == 60_000 and [p["piece"] for p in r3["paiements"]] == ["B1"],
          r3["paiements"])
    check("TF3 à payer 0, crédit 60 000", r3["ligne_29"] == 0 and r3["ligne_30"] == 60_000, (r3["ligne_29"], r3["ligne_30"]))
    check("TF3 CA sans IVA 2 000 000 -> question (art. 155)", r3["ca_sans_iva"] == 2_000_000 and r3["questions"], r3["ca_sans_iva"])
    # TF4 : national 400 000 / IRPF 40 000 reversé 40 000 ; non national 900 000 / IRPF 120 000 non reversé
    n, x = r4["categories"]["national"], r4["categories"]["non_national"]
    check("TF4 national : 400 000, IRPF 40 000, reversé 40 000 (B2)",
          (n["remunerations"], n["irpf_retenu"], n["irpf_reverse"], n["cartas"]) == (400_000, 40_000, 40_000, ["B2"]), n)
    check("TF4 non national (libellé « EXPATRIES ») : 900 000, IRPF 120 000, reversé 0",
          (x["remunerations"], x["irpf_retenu"], x["irpf_reverse"]) == (900_000, 120_000, 0), x)
    check("TF4 non-reversement signalé", any("non national" in w and "AUCUN reversement" in w for w in r4["avertissements"]))
    # TF5 : F100 HT 1 000 000 + IVA 150 000 ; F200 HT 1 000 000 ; sans tiers 30 000 (628)
    lg = {z["code"]: z for z in r5["lignes"]}
    f1, f2 = lg.get("F100", {}), lg.get("F200", {})          # sans plantage si un tiers manque
    check("TF5 F100 : HT 1 000 000, IVA 150 000, TTC 1 150 000",
          (f1.get("ht"), f1.get("iva"), f1.get("ttc")) == (1_000_000, 150_000, 1_150_000), f1)
    check("TF5 F200 : HT 1 000 000, IVA 0", (f2.get("ht"), f2.get("iva")) == (1_000_000, 0), f2)
    check("TF5 nom suggéré contient ALPHA / BETA", "ALPHA" in f1.get("nom_suggere", "") and "BETA" in f2.get("nom_suggere", ""),
          (f1.get("nom_suggere"), f2.get("nom_suggere")))
    check("TF5 sans code tiers : 628 = 30 000 + rappel art. 27.4",
          r5["sans_tiers"] == [{"compte": "62800000", "nature": "DIVERS", "ht": 30_000, "nb_lignes": 1}]
          and any("27.4" in w for w in r5["avertissements"]), r5["sans_tiers"])
    return gl, r


d = tempfile.mkdtemp()
p1, p2 = os.path.join(d, "gl.xlsx"), os.path.join(d, "gl.csv")
par_compte(p1); a_plat(p2)
gl, r = verifier("grand livre « par compte » (xlsx)", p1, True)
verifier("grand livre « à plat » (csv, dates texte, montants texte)", p2, False)

print("== profil : IVA déductible présentée dans les « pagos » (pratique du dossier pilote)")
r2 = t.calculer(gl, libelles=LIB, profil={"interpretations": {"tf3_iva_deductible": "dans_paiements"}})["tf3"]
check("ligne 28 = 60 000 + 150 000 = 210 000 ; 15 bis = 0 ; crédit inchangé 60 000",
      (r2["ligne_28"], r2["ligne_15bis"], r2["ligne_30"]) == (210_000, 0, 60_000), (r2["ligne_28"], r2["ligne_15bis"]))
check("écart de présentation signalé", any("dans_paiements" in w for w in r2["avertissements"]))
r3 = t.calculer(gl, libelles=LIB, profil={"faits": {"iva_ca_sans_tva": "exonere"}})["tf3"]
check("fait « exonéré » : 2 000 000 en base exonerada, plus de question", r3["base_exoneree"] == 2_000_000 and not r3["questions"])

print("== fichier des tiers : noms confirmés")
tp = os.path.join(d, "tiers.csv")
with open(tp, "w", encoding="utf-8", newline="") as f:
    csv.writer(f, delimiter=";").writerows([["Code", "Nom", "Adresse", "NIF"], ["F100", "SCI ALPHA SL", "Malabo", "123-AB"]])
r5 = t.calculer(gl, libelles=LIB, tiers_ref=t.lire_tiers(tp))["tf5"]
f100 = next((z for z in r5["lignes"] if z["code"] == "F100"), {})
check("F100 : nom et NIF du fichier, plus de suggestion",
      (f100.get("nom"), f100.get("nif"), f100.get("nom_suggere")) == ("SCI ALPHA SL", "123-AB", ""), f100)

print("== Excel : formules vivantes, taux lu dans Parametros")
px = os.path.join(d, "tf345.xlsx")
t.ecrire_xlsx(r, px)
from openpyxl import load_workbook  # noqa: E402
wb = load_workbook(px)
f = [c.value for n in ("TF3", "TF4", "TF5") for row in wb[n].iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
check("formules présentes (≥ 25)", len(f) >= 25, len(f))
check("aucun taux littéral (0.15)", not any(re.search(r"0\.15\b|15%", x) for x in f))
check("taux IVA lu dans Parametros", sum("Parametros!" in x for x in f) >= 12)
check("traçabilité : une ligne par montant utilisé", wb["Tracabilite"].max_row - 1 == len(r["tf3"]["trace"]) + len(r["tf4"]["trace"]) + len(r["tf5"]["trace"]))

print("== reproductibilité : même fichier -> même résultat, quelle que soit la graine de hachage")
import json as _json  # noqa: E402
import subprocess  # noqa: E402
prog = ("import sys, json; sys.path.insert(0, r'%s'); import grand_livre as g, tf345_ge as t; "
        "r = t.calculer(g.lire(r'%s'), libelles=%r); "
        "print(json.dumps({k: {kk: vv for kk, vv in v.items() if kk != 'trace'} for k, v in r.items() if k.startswith('tf')}, "
        "sort_keys=True, default=list))") % (os.path.join(ROOT, "lib"), p1, LIB)
sorties = set()
for seed in ("1", "2", "3", "4", "5", "6"):
    env = dict(os.environ, PYTHONHASHSEED=seed)
    cp = subprocess.run([sys.executable, "-c", prog], capture_output=True, text=True, encoding="utf-8", env=env)
    sorties.add(cp.stdout)
check("TF3/TF4/TF5 identiques sous 6 graines de hachage", len(sorties) == 1 and "" not in sorties, f"{len(sorties)} résultats différents")

print(f"\n{sum(ok)}/{len(ok)} contrôles conformes")
sys.exit(0 if all(ok) else 1)
