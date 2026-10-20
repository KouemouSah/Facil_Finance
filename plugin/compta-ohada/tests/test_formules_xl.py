# -*- coding: utf-8 -*-
"""Garde de l'évaluateur de formules (skills/dsf-syscohada/scripts/formules_xl.py), sans Excel.
Chaque attendu est la valeur qu'Excel donne, calculée à la main (commentaire). Lancer : python tests/test_formules_xl.py
"""
import os
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "skills", "dsf-syscohada", "scripts"))
import formules_xl as fx  # noqa: E402

ok = []


def check(nom, cond, detail=""):
    ok.append(bool(cond))
    print(f"  {'OK   ' if cond else 'ÉCHEC'} {nom}" + ("" if cond else f"  — {detail}"))


def main():
    from openpyxl import Workbook
    import openpyxl
    wb = Workbook()
    pg = wb.active
    pg.title = "page garde"
    pg["D16"] = "SOMA"
    pg["D18"] = "NIF-9"
    ws = wb.create_sheet("note 7")
    vals = {"C7": 100, "D7": 80, "C8": 50, "D8": None, "C9": "texte", "E1": 2.5, "E2": -2.5, "E3": 1234.5678}
    for k, v in vals.items():
        ws[k] = v
    f = {
        "C10": ("=N(C7)+N(C8)-N(D8)", 150),                    # 100 + 50 - 0
        "C11": ('=IF(N(D8)=0,"",(N(C7)-N(D8))/D8)', ""),      # D8 vide -> "" (pas de division)
        "C12": ('=IF(N(D7)=0,"",(N(C7)-N(D7))/D7)', 0.25),     # (100-80)/80
        "C13": ("=SUM(C7:C9)", 150),                           # le texte est ignoré par SUM
        "C14": ("=ROUND(E1,0)", 3), "C15": ("=ROUND(E2,0)", -3),   # au plus loin de zéro, pas « au pair »
        "C16": ("=ROUND(E3,2)", 1234.57),
        "C17": ("=MAX(C7-D7,0)", 20), "C18": ("=MAX(D7-C7,0)", 0),
        "C19": ("=C7/D8", "#DIV/0!"), "C20": ("=C9+1", "#VALUE!"),
        "C21": ("=IF('page garde'!D16=\"\",\"\",'page garde'!D16)", "SOMA"),
        "C22": ("=\"NIF : \"&'page garde'!D18", "NIF : NIF-9"),
        "C23": ("=C10+C13", 300),                              # formule qui lit des formules
        "C24": ("=-C7+2^3", -92), "C25": ("=50%*C7", 50),
        "C26": ("=ABS(D7-C7)", 20), "C27": ("=IFERROR(C7/D8,0)", 0),
        "C28": ("=IF(AND(C7>0,NOT(D7>100)),1,0)", 1),
        "C29": ("=N(C9)+1", 1),                                # N(texte) = 0, jamais le texte
        "C31": ("=IF(C7>0,1,FOO(C7))", 1),                     # la branche non retenue n'est PAS évaluée
    }
    for k, (form, _) in f.items():
        ws[k] = form
    ws["C30"] = "=FOO(C7)"                                     # hors périmètre : signalée, jamais devinée
    d = tempfile.mkdtemp()
    src, dst = os.path.join(d, "f.xlsx"), os.path.join(d, "v.xlsx")
    wb.save(src)
    r = fx.evaluer_classeur(src, dst)
    v = openpyxl.load_workbook(dst)["note 7"]
    for k, (form, att) in f.items():
        got = v[k].value
        # un texte vide enregistré dans un classeur se relit « cellule vide » (None) : équivalent pour le rapport
        bon = abs(got - att) < 1e-9 if isinstance(att, (int, float)) and isinstance(got, (int, float)) else \
            (got == att or (att == "" and got is None))
        check(f"{form} = {att!r}", bon, got)
    check("fonction hors périmètre signalée (une seule)", len(r["non_supportees"]) == 1 and "C30" in r["non_supportees"][0], r)
    check("toutes les formules comptées", r["formules"] == len(f) + 1, r["formules"])
    check("mise en forme conservée : le texte reste du texte", v["C22"].data_type == "s")
    print(f"\n{sum(ok)}/{len(ok)} contrôles conformes")
    return 0 if all(ok) else 1


if __name__ == "__main__":
    sys.exit(main())
