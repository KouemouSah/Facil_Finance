# -*- coding: utf-8 -*-
"""Garde du générateur de DEF « modèle officiel » (skills/dsf-syscohada/scripts/def_modele.py).

Même balance construite à la main que test_def_ge.py ; MODÈLE construit ici : une DEF 2024 d'une AUTRE entité
(identité « ANCIENNE SA »), avec des montants volontairement faux (999, 5) pour prouver qu'aucun montant du
modèle ne fuit dans la colonne N, et qu'il devient bien la colonne N-1 (--colonne-precedente N).
Lancer : python tests/test_def_modele.py     (code 0 = conforme ; ni Excel ni le PDF ne sont requis)
"""
import argparse
import csv
import json
import os
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "skills", "dsf-syscohada", "scripts"))
import def_modele as dm  # noqa: E402

B = {"10130000": (-1_000_000, 0, 0), "12100000": (200_000, 0, 0),
     "24100000": (600_000, 100_000, 0), "28410000": (-200_000, 0, 70_000),
     "52110000": (500_000, 1_000_000, 500_000), "40110000": (-100_000, 400_000, 450_000),
     "70510000": (0, 0, 1_000_000), "60400000": (0, 450_000, 0), "68130000": (0, 70_000, 0)}
LIB = {"10130000": "CAPITAL", "12100000": "REPORT A NOUVEAU", "24100000": "MATERIEL", "28410000": "AMORT MATERIEL",
       "52110000": "BANQUE", "40110000": "FOURNISSEURS", "70510000": "TRAVAUX", "60400000": "ACHATS MATIERES",
       "68130000": "DOTATIONS AMORT"}
ok = []
# un libellé réellement présent dans le dictionnaire de traduction (dérivé de la source, pas recopié)
LIBELLE_TRADUIT = next(k for k in json.load(open(dm.I18N, encoding="utf-8"))["es"] if k.startswith("Fournisseurs"))


def check(nom, cond, detail=""):
    ok.append(bool(cond))
    print(f"  {'OK   ' if cond else 'ÉCHEC'} {nom}" + ("" if cond else f"  — {detail}"))


def modele(path):
    from openpyxl import Workbook
    wb = Workbook()
    g = wb.active
    g.title = "page garde"
    for c, v in {"C11": "EXERCICE CLOS LE", "D11": "31/12/2024", "C13": "DENOMINATION SOCIALE", "D13": "ANCIENNE SA",
                 "C16": "SIGLE USUEL", "D16": "ANC", "C18": "N° IDENTIFICATION FISCALE", "D18": "NIF-OLD"}.items():
        g[c] = v
    g.print_area = "B4:H20"
    wb.create_sheet("SOMMAIRE")["A2"] = "SOMMAIRE"
    n = wb.create_sheet("note 17")
    for c, v in {"B1": "ANC", "B2": "31/12/2024", "B3": "NIF-OLD", "B4": "NOTE 17 FOURNISSEURS D'EXPLOITATION", "J4": "sommaire",
                 "B5": "Libellés", "C5": "Année", "D5": "Année", "E5": "Variation en %", "C6": "N", "D6": "N-1",
                 "A7": "DJ 401", "B7": LIBELLE_TRADUIT, "C7": 999, "D7": 111, "E7": 7.0,
                 "A8": "DJ 408", "B8": "Fournisseurs factures non parvenues", "C8": 5, "D8": 0, "E8": 0,
                 "A9": "DJ 409", "B9": "Ligne sans montant", "C9": 0, "D9": 0, "E9": 0,
                 "B10": "TOTAL FOURNISSEURS", "C10": 1004, "D10": 111, "E10": 8.05,
                 "B11": "= CAPACITE TEST (libellé, pas une formule)", "M20": 123456}.items():
        n[c] = v
    n["B11"].data_type = "s"                    # comme dans la DEF déposée : texte commençant par « = »
    n.print_area = "B1:E12"
    t = wb.create_sheet("TF5")
    for c, v in {"A7": "Nombre del beneficio", "C7": "Dirección", "G7": "Naturaleza de la remuneración", "L7": "Importe Pagado",
                 "A9": "ALPHA SARL", "C9": "MALABO", "G9": "REGIES", "K9": "NIF-A", "L9": 500,
                 "A10": "BETA", "C10": "BATA", "G10": "LOCATION", "L10": 300, "L22": 800}.items():
        t[c] = v
    t.print_area = "A1:L22"
    wb.save(path)


def main():
    d = tempfile.mkdtemp()
    p = os.path.join(d, "balance.csv")
    with open(p, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Compte", "Intitulé", "Ouverture D", "Ouverture C", "Mvt D", "Mvt C", "Solde D", "Solde C"])
        for a, (o, md, mk) in B.items():
            s = o + md - mk
            w.writerow([a, LIB[a], max(o, 0), max(-o, 0), md, mk, max(s, 0), max(-s, 0)])
    m = os.path.join(d, "modele.xlsx")
    modele(m)
    prof = os.path.join(d, "profil.json")
    json.dump({"faits": {"nouvelle_societe": False}}, open(prof, "w", encoding="utf-8"))
    from PIL import Image as PILImage
    logo = os.path.join(d, "logo.png")
    PILImage.new("RGB", (120, 60), (31, 56, 100)).save(logo)
    a = argparse.Namespace(modele=m, balance=p, exercice=2025, grand_livre=None, profil=prof, decisions=None,
                           colonne_precedente="N", logo=logo, langues="fr,es", sortie=os.path.join(d, "DEF"), pdf=False, pro=True)
    g = dm.Generateur(a)
    import openpyxl
    print("== Listes de comptes du modèle")
    check("« DA 181 à 183 » : plage développée", [x[1] for x in dm.expr_comptes("DA 181 à 183")[1]] == ["181", "182", "183"])
    check("« BI 411-4112+413 » : signes", [x[0] for x in dm.expr_comptes("BI 411-4112+413")[1]] == [1, -1, 1])
    check("« AK 2831, 2832p » : partiel repéré", dm.expr_comptes("AK 2831, 2832p")[1][1] == (1, "2832", True))
    check("libellé sans compte : pas une liste", dm.expr_comptes("Fournisseurs") is None)

    x = os.path.join(d, "DEF_FR.xlsx")
    noms = g.rendre("fr", x)
    wb = openpyxl.load_workbook(x)
    print("== Structure")
    check("mêmes feuilles que le modèle", noms == ["page garde", "SOMMAIRE", "note 17", "TF5"], noms)
    n = wb["note 17"]
    print("== Identité : jamais recopiée, saisie une fois en page de garde, reprise par formule")
    pg = wb["page garde"]
    check("dénomination / sigle / NIF du modèle effacés en page de garde", all(pg[c].value is None for c in ("D13", "D16", "D18")),
          [pg[c].value for c in ("D13", "D16", "D18")])
    check("date de clôture = exercice généré", pg["D11"].value == "31/12/2025", pg["D11"].value)
    check("en-tête de note = formule vers la page de garde", str(n["B1"].value).startswith("=IF('page garde'!D16"), n["B1"].value)
    check("NIF de l'en-tête = formule vers la page de garde", "'page garde'!D18" in str(n["B3"].value), n["B3"].value)
    check("aucune trace de l'identité du modèle", not any("ANCIENNE" in str(c.value) or "NIF-OLD" in str(c.value)
                                                          for ws in wb.worksheets for row in ws.iter_rows() for c in row))
    print("== Montants : N recalculé, N-1 = colonne N du modèle, totaux et variations en formules")
    check("N ligne 401 = 150 000 (balance, sens passif)", n["C7"].value == 150_000, n["C7"].value)
    check("N ligne 408 = 0", n["C8"].value == 0, n["C8"].value)
    check("N-1 = colonne N du modèle (999)", n["D7"].value == 999, n["D7"].value)
    check("aucun montant N du modèle en colonne N", n["C7"].value != 999 and n["C8"].value != 5)
    check("total = formule des lignes (structure prouvée par les montants du modèle)", n["C10"].value == "=N(C7)+N(C8)+N(C9)", n["C10"].value)
    check("variation en % = formule", str(n["E7"].value).startswith("=IF(N(D7)=0"), n["E7"].value)
    check("brouillon hors zone d'impression écarté", n["M20"].value is None, n["M20"].value)
    check("libellé « = … » écrit en TEXTE", n["B11"].data_type == "s" and str(n["B11"].value).startswith("="), n["B11"].data_type)
    check("renvoi « sommaire » = lien interne", n["J4"].hyperlink is not None)
    print("== Logo de l'entité : page de présentation du classeur, couverture du rapport, nulle part ailleurs")
    check("classeur : logo sur la page de garde", len(wb["page garde"]._images) == 1, len(wb["page garde"]._images))
    check("classeur : aucune autre feuille ne porte le logo", all(not w._images for w in wb.worksheets if w.title != "page garde"),
          [w.title for w in wb.worksheets if w._images])
    import def_ge
    rx = os.path.join(d, "rapport.xlsx")
    fr = def_ge.ecrire_xlsx(g.D, rx, logo)
    wr = openpyxl.load_workbook(rx)
    check("rapport : logo sur la couverture seulement", len(wr["Couverture"]._images) == 1
          and all(not w._images for w in wr.worksheets if w.title != "Couverture"), [w.title for w in wr.worksheets if w._images])
    check("rapport : l'annexe « Controles » n'est pas une feuille du rapport déposé", "Controles" not in fr, fr)
    print("== Rapport : extraction des tableaux d'une feuille (sans Excel)")
    import rapport_def as rd
    from openpyxl import Workbook
    wx = Workbook().active
    for c, v in {"B4": "NOTE 12 ECARTS DE CONVERSION", "B5": "Libellés", "C5": "Devises", "G5": "Variation",
                 "A7": "BU 478", "B7": "Ecarts actif", "B8": "Various", "G8": 1857.67, "B9": "Commentaire :",
                 "B10": "Texte libre du commentaire", "B12": "TRANSFERTS DE CHARGES", "B13": "Libellés", "C13": "Année",
                 "C14": "N", "A15": "TI 781", "B15": "Transferts de charges", "C15": 132_100_047}.items():
        wx[c] = v
    tabs = rd.extraire(wx, {"c0": 1, "c1": 7, "r1": 16, "h0": 5, "h1": 6, "totaux": []}, "note 12")
    check("deux tableaux séparés par le commentaire", len(tabs) == 2, [t[0] for t in tabs])
    t2 = tabs[1] if len(tabs) > 1 else ("", [], [])
    check("second tableau : titre, en-têtes et montant repris", t2[0] == "TRANSFERTS DE CHARGES"
          and any(v == 132_100_047 for _, vals, _, _ in t2[2] for v in vals), t2)
    check("texte du commentaire non repris", not any("Texte libre" in str(v) for t_ in tabs for _, vals, _, _ in t_[2] for v in vals))
    check("note de texte (note 2) réputée utile, note à zéro réputée néant",
          rd.utile([("", [], [("section", ["A - DECLARATION DE CONFORMITE"], [None], 5)])], texte=True)
          and not rd.utile([("", [], [("base", ["x", 0], [None, None], 5)])]))
    print("== Rapport de contrôle A4 : sections, statuts écrits, sources déclarées")
    import controle_def
    cx = os.path.join(d, "controles.xlsx")
    controle_def.construire(g, cx)
    wc = openpyxl.load_workbook(cx)
    check("couverture + 7 sections", wc.sheetnames[0] == "Rapport de contrôle" and len(wc.sheetnames) == 8
          and wc.sheetnames[-1].startswith("7."), wc.sheetnames)
    txt_cov = " ".join(str(c.value) for row in wc["Rapport de contrôle"].iter_rows() for c in row if c.value)
    check("verdict écrit sur la couverture", "Verdict :" in txt_cov, txt_cov[:200])
    txt_src = " ".join(str(c.value) for row in wc[wc.sheetnames[-1]].iter_rows() for c in row if c.value)
    check("sources déclarées : balance et modèle nommés", "balance.csv" in txt_src and "modele.xlsx" in txt_src, txt_src[:300])
    check("montants au format français", controle_def.fr(1234567) == "1 234 567", controle_def.fr(1234567))
    check("anomalie non bloquante (valeur nette négative) : À VOIR, pas ÉCART",
          not any(b in "aucune valeur nette négative à l'actif" for b in controle_def.BLOQUANTS)
          and any(b in "total actif = total passif" for b in controle_def.BLOQUANTS))
    print("== TF5 hybride : identité des tiers reprise du modèle, montant non inventé")
    t = wb["TF5"]
    check("tiers du modèle conservés (nom, NIF)", (t["A9"].value, t["K9"].value, t["A10"].value) == ("ALPHA SARL", "NIF-A", "BETA"))
    check("modèle d'un autre exercice, sans grand livre : montant laissé à saisir", t["L9"].value is None, t["L9"].value)
    check("NIF manquant signalé (fond ambre)", t["K10"].fill.fgColor.rgb.endswith(dm.AMBRE), t["K10"].fill.fgColor.rgb)
    print("== Version espagnole : même classeur, autre langue")
    xe = os.path.join(d, "DEF_ES.xlsx")
    noms_es = g.rendre("es", xe)
    we = openpyxl.load_workbook(xe)
    check("feuilles traduites, même nombre", len(noms_es) == len(noms) and "portada" in noms_es, noms_es)
    ne = we[noms_es[2]]
    check("formules d'en-tête vers la page de garde ESPAGNOLE", "'portada'!D16" in str(ne["B1"].value), ne["B1"].value)
    check("mêmes montants qu'en français", (ne["C7"].value, ne["D7"].value, ne["C10"].value) == (n["C7"].value, n["D7"].value, n["C10"].value))
    check("libellé traduit", ne["B7"].value != n["B7"].value, ne["B7"].value)
    check("TF5 : entête espagnole traduite en FRANÇAIS dans la version FR", t["A7"].value == "Nom du bénéficiaire", t["A7"].value)
    print("== Modèle professionnel : mêmes cellules, présentation seule")
    wp = openpyxl.load_workbook(os.path.join(d, "DEF_FR_pro.xlsx"))
    np_ = wp["note 17"]
    check("mêmes valeurs et formules que l'officiel", all(np_.cell(r, c).value == n.cell(r, c).value for r in range(1, 12) for c in range(1, 6)))
    check("ligne sans montant N ni N-1 masquée, jamais supprimée", np_.row_dimensions[9].hidden and np_["C9"].value == 0)
    check("ligne à N nul mais N-1 non nul : visible", not np_.row_dimensions[8].hidden)
    check("en-tête figé", np_.freeze_panes == "A7", np_.freeze_panes)
    print("== Parité des traductions (i18n-def.json)")
    i18n = json.load(open(dm.I18N, encoding="utf-8"))
    fe = list(i18n["feuilles_es"].values())
    check("noms de feuilles ES : ≤ 31 caractères, uniques", all(len(v) <= 31 for v in fe) and len(set(fe)) == len(fe))
    check("aucune traduction vide", all(v.strip() for dic in (i18n["es"], i18n["fr"]) for v in dic.values()))
    check("noms propres non traduits (GAMMA LOGISTICS)", i18n["es"].get("GAMMA LOGISTICS", "GAMMA LOGISTICS") == "GAMMA LOGISTICS")
    print(f"\n{sum(ok)}/{len(ok)} contrôles conformes")
    return 0 if all(ok) else 1


if __name__ == "__main__":
    sys.exit(main())
