# -*- coding: utf-8 -*-
"""Falsification de tests/test_gabarit.py : une mutation par règle protégée du mode gabarit, appliquée à une COPIE du
plugin ; la garde doit échouer sur un contrôle métier (pas planter).
Lancer : python tests/falsification_gabarit.py   -> code 0 si toutes sont détectées.
"""
import os
import shutil
import subprocess
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
G = os.path.join("skills", "dsf-syscohada", "scripts", "def_modele.py")
E = os.path.join("skills", "dsf-syscohada", "scripts", "espace_def.py")
MUT = [
    ("N-1 des notes à 0 au lieu de vide", G,
     "        if self.gabarit:                                   # aucun N-1 dans un formulaire vierge : vide, jamais 0\n            return None",
     "        if self.gabarit:\n            return v"),
    ("année du gabarit non décalée", G,
     "            return y + self.a.exercice - self.annee_modele", "            return y"),
    ("rôles des colonnes lus avec l'exercice au lieu de l'année du gabarit", G,
     "roles = roles_colonnes(mws, h0, h1, self.annee_modele or self.a.exercice)",
     "roles = roles_colonnes(mws, h0, h1, self.a.exercice)"),
    ("gabarit ignoré : liasse sans DEF déposée", E,
     "                     gabarit=not mod,", "                     gabarit=False,"),
    ("classeur non adapté (51 feuilles toujours)", G,
     "            omises = elaguer_notes(g, x, lang)\n", "            omises = []\n"),
    ("notes omises non signalées dans la grille", G,
     "            if k in omises:\n                grille.cell(r, col, L[2])", "            if False:\n                grille.cell(r, col, L[2])"),
    ("notes extra-comptables omises elles aussi", G,
     "or t in rd.EXTRA or t in EXTRA_COMPTABLES or nom in a_saisir", "or nom in a_saisir"),
    ("choix « conserver » ignoré", G,
     ".get(\"notes_neant\") != \"conserver\":", ".get(\"notes_neant\") != \"jamais\":"),
    ("note applicable absente du formulaire passée sous silence", G,
     "            if not m or m.group(1) in dispo or p.get(\"total\"):", "            if True:"),
    ("question de la DEF précédente oubliée", E,
     "        qs.insert(0, {\"id\": \"sources.def_precedente\"", "        (lambda *_: None)(0, {\"id\": \"sources.def_precedente\""),
]


def main():
    aveugles = 0
    for name, rel, a, b in MUT:
        d = tempfile.mkdtemp()
        dst = os.path.join(d, "compta-ohada")
        shutil.copytree(ROOT, dst, ignore=shutil.ignore_patterns("dist", "datas", "__pycache__"))
        path = os.path.join(dst, rel)
        src = open(path, encoding="utf-8").read()
        assert src.count(a) == 1, f"motif introuvable ou multiple : {name}"
        open(path, "w", encoding="utf-8").write(src.replace(a, b))
        cp = subprocess.run([sys.executable, os.path.join(dst, "tests", "test_gabarit.py")],
                            capture_output=True, text=True, encoding="utf-8")
        detectee = cp.returncode != 0 and "ÉCHEC" in cp.stdout
        aveugles += not detectee
        bilan = next((l.strip() for l in cp.stdout.splitlines() if "conformes" in l), "PLANTAGE : " + cp.stderr.strip()[-90:])
        print(f"  {name:62s} {'détectée' if detectee else '!! NON DÉTECTÉE'}   {bilan}")
        shutil.rmtree(d, ignore_errors=True)
    # et le gabarit lui-même : un montant réinjecté doit être vu
    d = tempfile.mkdtemp()
    dst = os.path.join(d, "compta-ohada")
    shutil.copytree(ROOT, dst, ignore=shutil.ignore_patterns("dist", "datas", "__pycache__"))
    import openpyxl
    gp = os.path.join(dst, "skills", "dsf-syscohada", "templates", "def_ge_gabarit.xlsx")
    wb = openpyxl.load_workbook(gp)
    wb["note 28"]["C9"] = 378_000_000
    wb.save(gp)
    cp = subprocess.run([sys.executable, os.path.join(dst, "tests", "test_gabarit.py")], capture_output=True, text=True, encoding="utf-8")
    detectee = cp.returncode != 0 and "ÉCHEC" in cp.stdout
    aveugles += not detectee
    print(f"  {'montant resté dans le gabarit':62s} {'détectée' if detectee else '!! NON DÉTECTÉE'}")
    shutil.rmtree(d, ignore_errors=True)
    sys.exit(1 if aveugles else 0)


if __name__ == "__main__":
    main()
