# -*- coding: utf-8 -*-
"""DEF « scannée » de VÉRITÉ CONNUE, pour éprouver ocr_def sans aucune donnée client.

DEF de test (balance synthétique de test_def_modele, gabarit officiel GE) → bilan paysage et compte de résultat
imprimés en PDF par le rendu NATIF du plugin (aucun Excel requis) → rastérisés à 200 dpi en niveaux de gris, bilan
tourné de 90° comme sur les DEF réelles numérisées, légère inclinaison, bruit, JPEG. La vérité (montants par code
de ligne) est lue dans le classeur évalué.
"""
import glob
import json
import os
import re
import sys
import tempfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tests"))
FEUILLES = ["balance apaisado", "cuenta resultados"]


def fabriquer(dossier=None):
    """(chemin du PDF scanné, vérité {code: [montants de la ligne]}, exercice)."""
    import cv2
    import numpy as np
    import openpyxl
    import pymupdf
    import test_reprise_validation as trv
    ed = trv.ed
    import def_ge
    import formules_xl
    dossier = dossier or tempfile.mkdtemp()
    reel = ed.generer
    ed.generer = lambda d, **k: reel(d, **{**k, "pdf": False})
    try:
        d, _ = trv.balance_seule()
        ed.repondre(d, "interp.modele_def", "ge")
        ed.generer(d)
    finally:
        ed.generer = reel
    xlsx = glob.glob(os.path.join(d, "02_xls", "*_ES_projet.xlsx"))[0]
    # Excel IMPRIME les traits verticaux des colonnes (DEF réelles numérisées) ; le rendu natif du plugin ne trace que
    # les bordures déclarées : on les déclare toutes, pour une page imprimée comme celles des DEF déposées
    # et, comme les DEF réelles imprimées (constaté sur les deux scans) : pas d'aplat, texte noir — un total en blanc
    # sur fond sombre ne se lit pas une fois numérisé
    from copy import copy
    from openpyxl.styles import Border, PatternFill, Side
    wb_q = openpyxl.load_workbook(xlsx)
    fin = Side(style="thin", color="000000")
    for f in FEUILLES:
        ws = wb_q[f]
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
            for c in row:
                if type(c).__name__ != "MergedCell":
                    c.border = Border(left=fin, right=fin, top=fin, bottom=fin)
                    c.fill = PatternFill(fill_type=None)
                    police = copy(c.font)
                    police.color = "000000"
                    c.font = police
    quadrille = os.path.join(dossier, "quadrille.xlsx")
    wb_q.save(quadrille)
    xlsx_imp = quadrille
    avant = os.environ.get("FACIL_PDF")
    os.environ["FACIL_PDF"] = "natif"                     # rendu du plugin : identique sur tout poste
    try:
        pdf = os.path.join(dossier, "imprimee.pdf")
        def_ge.exporter_pdf(xlsx_imp, pdf, FEUILLES)
    finally:
        if avant is None:
            os.environ.pop("FACIL_PDF", None)
        else:
            os.environ["FACIL_PDF"] = avant
    val = os.path.join(dossier, "valeurs.xlsx")
    formules_xl.evaluer_classeur(xlsx, val)
    wb = openpyxl.load_workbook(val, data_only=True)
    verite = {}
    for f in FEUILLES:
        for row in wb[f].iter_rows(values_only=True):
            cells = list(row)
            for i, c in enumerate(cells):
                if isinstance(c, str) and re.fullmatch(r"[A-Z]{2}", c.strip()):
                    suite = []
                    for x in cells[i + 1:]:
                        if isinstance(x, str) and re.fullmatch(r"[A-Z]{2}", x.strip()):
                            break
                        if isinstance(x, (int, float)) and not isinstance(x, bool):
                            suite.append(round(x))
                    verite.setdefault(c.strip(), suite)
    src, out = pymupdf.open(pdf), pymupdf.open()
    rng = np.random.default_rng(7)
    for page in src:
        pix = page.get_pixmap(dpi=200)
        img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, pix.n)[:, :, :3].copy()
        g = cv2.cvtColor(img, cv2.COLOR_RGB2GRAY)
        img = cv2.cvtColor(np.where(g < 235, g, 255).astype(np.uint8), cv2.COLOR_GRAY2RGB)   # imprimante noir et blanc
        if img.shape[1] > img.shape[0]:
            img = np.ascontiguousarray(np.rot90(img, 1))  # paysage numérisé debout
        h, w = img.shape[:2]
        M = cv2.getRotationMatrix2D((w / 2, h / 2), rng.uniform(-0.5, 0.5), 1.0)
        img = cv2.warpAffine(img, M, (w, h), borderValue=(255, 255, 255))
        img = np.clip(img.astype(int) + rng.normal(0, 9, img.shape), 0, 255).astype(np.uint8)
        ok, jpg = cv2.imencode(".jpg", img, [cv2.IMWRITE_JPEG_QUALITY, 70])
        p = out.new_page(width=595, height=842)
        p.insert_image(p.rect, stream=jpg.tobytes())
    scan = os.path.join(dossier, "DEF_test_scannee.pdf")
    out.save(scan)
    return scan, verite, int(ed.lire_json(os.path.join(d, "manifeste.json"), {})["exercice"])


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    scan, verite, ex = fabriquer(sys.argv[1] if len(sys.argv) > 1 else None)
    json.dump(verite, open(scan + ".verite.json", "w", encoding="utf-8"))
    print(scan, len(verite), "codes ; exercice", ex)
