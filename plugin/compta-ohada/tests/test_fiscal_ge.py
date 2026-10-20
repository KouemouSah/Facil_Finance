# -*- coding: utf-8 -*-
"""Garde du moteur fiscal GE (lib/fiscal_ge.py) : TF1 / TF2.
Valeurs ATTENDUES calculées à la main (commentaires), jamais par le code testé.
Lancer : python tests/test_fiscal_ge.py     (code 0 = conforme)
"""
import os
import re
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "lib"))
import fiscal_ge as f  # noqa: E402

ok = []


def check(nom, cond, detail=""):
    ok.append(bool(cond))
    print(f"  {'OK   ' if cond else 'ÉCHEC'} {nom}" + ("" if cond else f"  — {detail}"))


def E(res, ca, n=2025, autres=0.0, clo=None):
    return {"exercice": n, "date_cloture": clo or f"{n}-12-31", "resultat_avant_is": res, "ca": ca, "autres_produits": autres}


print("== A. réintégration, déficits (dont un périmé), CMF versée")
# RF = 1 000 000 + 50 000 = 1 050 000 ; déficit 2021 imputé 200 000 (2024-2021 = 3 ≤ 5) ;
# déficit 2018 périmé (6 > 5) ; imposable 850 000 ; IS 25 % = 212 500 ; CMF 1,5 % × 10 M = 150 000 ;
# dû = max = 212 500 ; versé 100 000 -> à payer 112 500
r = f.calculer(E(1_000_000, 10_000_000, n=2024), {
    "reintegrations": [{"ligne": "multas", "montant": 50_000}],
    "deficits": [{"exercice": 2021, "montant": 200_000}, {"exercice": 2018, "montant": 300_000}],
    "cmf_versee": [{"date": "2024-07-15", "montant": 100_000, "piece": "Q-1"}]})
check("RF (22) = 1 050 000", r["tf1"][22] == 1_050_000, r["tf1"][22])
check("déficit imputé (24) = 200 000", r["tf1"][24] == 200_000, r["tf1"][24])
check("résultat net (26) = 850 000", r["tf1"][26] == 850_000, r["tf1"][26])
check("IS (5) = 212 500", r["tf2"][5] == 212_500, r["tf2"][5])
check("CMF (14) = 150 000", r["tf2"][14] == 150_000, r["tf2"][14])
check("IS à payer (7) = 112 500", r["tf2"][7] == 112_500, r["tf2"][7])
check("déficit 2018 signalé périmé (art. 73.2)", any("2018" in w and "73.2" in w for w in r["avertissements"]))

print("== B. CMF supérieure à l'IS ; CMF versée non renseignée")
# IS = 25 % × 100 000 = 25 000 ; CMF = 1,5 % × 50 M = 750 000 -> dû 750 000 ; versé 0 (défaut visible)
r = f.calculer(E(100_000, 50_000_000), {})
check("IS à payer = CMF = 750 000", r["tf2"][7] == 750_000 and r["impot_du"] == 750_000, r["tf2"])
check("hypothèse « CMF versée = 0 » signalée, non bloquante", any("inconnue" in w for w in r["avertissements"]))

print("== B bis. CMF VERSÉE lue dans les fichiers (règle du porteur : ce que la comptabilité enregistre est versé)")
# dû 750 000 (cas B). Grand livre, compte 441, exercice 2025 :
#   15/07 ACOMPTE CMF 1ER 2025 débit 300 000 -> retenu (libellé explicite)
#   20/01 VERSEMENT IS         débit 200 000 -> exclu : janvier = 2e acompte 2024 (art. 80.2)
#   28/02 REGUL CONTROL FISCAL 2023  débit 90 000 -> exclu : exercice antérieur
#   10/11 VIREMENT TRESOR      débit 50 000 -> retenu, à confirmer (pas de libellé CMF)
#   31/12 IMPOT 2025           CRÉDIT 750 000 -> ignoré (ce n'est pas un versement)
# versé = 350 000 ; à payer = 750 000 - 350 000 = 400 000
GL = [{"date": "2025-07-15", "compte": "44100000", "piece": "BQ1", "libelle": "ACOMPTE CMF 1ER 2025", "debit": 300_000, "credit": 0},
      {"date": "2025-01-20", "compte": "44100000", "piece": "BQ2", "libelle": "VERSEMENT IS", "debit": 200_000, "credit": 0},
      {"date": "2025-02-28", "compte": "44100000", "piece": "OD1", "libelle": "REGUL CONTROL FISCAL 2023", "debit": 90_000, "credit": 0},
      {"date": "2025-11-10", "compte": "44100000", "piece": "BQ3", "libelle": "VIREMENT TRESOR", "debit": 50_000, "credit": 0},
      {"date": "2025-12-31", "compte": "44100000", "piece": "OD2", "libelle": "IMPOT 2025", "debit": 0, "credit": 750_000}]
r = f.calculer(E(100_000, 50_000_000), {}, gl_lignes=GL)
check("TF2 ligne 6 = 350 000 (juillet + novembre)", r["tf2"][6] == 350_000, r["tf2"][6])
check("TF2 ligne 7 = 400 000", r["tf2"][7] == 400_000, r["tf2"][7])
check("janvier exclu comme 2e acompte de N-1, redressement 2023 exclu",
      sorted(x["piece"] for x in r["versements_exclus"]) == ["BQ2", "OD1"], r["versements_exclus"])
check("versement sans libellé CMF : retenu mais « à confirmer »", any("à confirmer" in w for w in r["avertissements"]))
check("quittance non exigée pour le calcul", any("non exigées pour le calcul" in w for w in r["avertissements"]))
# balance seule : débit de l'exercice au 441 = 120 000 -> retenu
r = f.calculer(E(100_000, 50_000_000), {}, mouvements={"44100000": (120_000, 750_000), "52100000": (5, 0)})
check("balance seule : débits du 441 = versés (120 000)", r["tf2"][6] == 120_000 and r["tf2"][7] == 630_000, r["tf2"])
# aucune écriture de versement -> 0, dit explicitement
r = f.calculer(E(100_000, 50_000_000), {}, mouvements={"44100000": (0, 750_000)})
check("aucun débit du 441 : ligne 6 = 0, constat écrit", r["tf2"][6] == 0 and any("Aucun versement" in w for w in r["avertissements"]))
print("== B ter. IS d'un exercice antérieur passé EN CHARGE, reconnu par la structure de la pièce (art. 62.3.a)")
# pièce RAPPEL : 441 crédit 1 000 000 (IS 2023) ; 6471 « RECARGO IS » 300 000 ; 6480 « CONTROLE 2023 » 700 000 + 150 000 (autre taxe)
#   -> principal d'IS en charge = 1 000 000 - 300 000 = 700 000 (le libellé du 6480 ne dit PAS « IS »)
# pièce IS 2025 : 891 débit / 441 crédit -> exclue (l'IS de l'exercice passe par le 89, rien à réintégrer)
GL2 = [{"date": "2025-02-28", "compte": "44100000", "piece": "RAP", "libelle": "IS 23 CONTROLE", "debit": 0, "credit": 1_000_000},
       {"date": "2025-02-28", "compte": "44780000", "piece": "RAP", "libelle": "AUTRE TAXE", "debit": 0, "credit": 150_000},
       {"date": "2025-02-28", "compte": "64710000", "piece": "RAP", "libelle": "RECARGO IS 2023", "debit": 300_000, "credit": 0},
       {"date": "2025-02-28", "compte": "64800000", "piece": "RAP", "libelle": "CONTROLE 2023", "debit": 700_000, "credit": 0},
       {"date": "2025-02-28", "compte": "64800000", "piece": "RAP", "libelle": "CONTROLE 2023", "debit": 150_000, "credit": 0},
       {"date": "2025-12-31", "compte": "89110000", "piece": "IS25", "libelle": "IMPOT 2025", "debit": 400_000, "credit": 0},
       {"date": "2025-12-31", "compte": "44100000", "piece": "IS25", "libelle": "IMPOT 2025", "debit": 0, "credit": 400_000}]
r = f.calculer(E(10_000_000, 1_000_000), {}, gl_lignes=GL2)
pr = [p for p in r["propositions"] if p.get("source", "").startswith("grand livre")]
check("proposition : 64800000 pour 700 000 (crédit IS − majorations d'IS)", [(p["compte"], p["montant"]) for p in pr] == [("64800000", 700_000)], pr)
check("jamais appliquée sans acceptation (TF1 ligne 9 = 0)", r["tf1"][9] == 0, r["tf1"][9])
r = f.calculer(E(10_000_000, 1_000_000), {"propositions_acceptees": ["64800000"]}, gl_lignes=GL2)
# RF = 10 000 000 + 700 000 = 10 700 000 ; IS 25 % = 2 675 000
check("acceptée : TF1 ligne 9 = 700 000, IS = 2 675 000", r["tf1"][9] == 700_000 and r["tf2"][5] == 2_675_000, (r["tf1"][9], r["tf2"][5]))

# la saisie de l'utilisateur prime sur les fichiers
r = f.calculer(E(100_000, 50_000_000), {"cmf_versee": [{"montant": 10_000}]}, mouvements={"44100000": (120_000, 0)})
check("saisie utilisateur prioritaire sur les fichiers", r["tf2"][6] == 10_000, r["tf2"][6])

print("== C. arrondi de la base (art. 75.1)")
# 1 234 567,89 -> base 1 234 000 -> IS 308 500
r = f.calculer(E(1_234_567.89, 1_000_000), {})
check("base arrondie 1 234 000, IS 308 500", r["base_is_arrondie"] == 1_234_000 and r["tf2"][5] == 308_500,
      (r["base_is_arrondie"], r["tf2"][5]))

print("== D. perte")
# perte 500 000 : IS 0 ; CMF 150 000 due ; déficit 2025 reportable jusqu'en 2030
r = f.calculer(E(-500_000, 10_000_000), {})
check("IS 0, à payer = CMF 150 000", r["tf2"][5] == 0 and r["tf2"][7] == 150_000, r["tf2"])
check("déficit reportable 500 000 jusqu'en 2030",
      r["deficits_reportables"] == [{"exercice": 2025, "montant": 500_000, "imputable_jusqu_en": 2030}], r["deficits_reportables"])

print("== E. société nouvelle (art. 75.4, 77.4)")
r = f.calculer(E(-500_000, 10_000_000), {"nouvelle_societe": True})
check("perte : CMF non applicable, rien à payer", r["tf2"][14] == 0 and r["tf2"][7] == 0, r["tf2"])
r = f.calculer(E(1_000_000, 10_000_000), {"nouvelle_societe": True})
check("bénéfice : IS 250 000 bonifié de 50 % = 125 000", r["tf2"][5] == 125_000 and r["tf2"][7] == 125_000, r["tf2"])

print("== F. propositions : le libellé doit confirmer la nature")
soldes = {"64710000": 1000, "65820000": 500, "65830000": 300, "65829000": 200, "62580000": 400, "89100000": 999}
lib = {"64710000": "PENALITES FISCALES", "65820000": "AUTRES CHARGES DIVERSES",
       "65830000": "DONS AUX ASSOCIATIONS", "65829000": "ABANDON DE CREANCE", "62580000": "ASSURANCE MALADIE CADRES"}
r = f.calculer(E(10_000, 1_000_000), {"propositions_acceptees": "toutes"}, soldes, libelles=lib)
typ = {p["compte"]: p["type"] for p in r["propositions"]}
check("647 pénalités -> réintégration", typ.get("64710000") == "reintegration", typ)
check("6582 « AUTRES CHARGES DIVERSES » -> question (nomenclature contredite)", typ.get("65820000") == "question", typ)
check("« ABANDON » ne passe pas pour un « don »", typ.get("65829000") == "question", typ)
check("« toutes » n'applique que les réintégrations confirmées : multas 1 000, dons 300",
      r["tf1"][10] == 1000 and r["tf1"][11] == 300 and r["tf1"][12] == 0, r["tf1"])
check("assurance maladie en attente de réponse, non appliquée",
      any(p["compte"] == "62580000" and p["statut"] == "en attente de réponse" for p in r["propositions"]))
r = f.calculer(E(10_000, 1_000_000), {}, soldes, libelles=lib)
check("sans décision : aucune proposition appliquée", r["tf1"][13] == 0, r["tf1"][13])
r = f.calculer(E(10_000, 1_000_000), {"propositions_acceptees": ["62580000"]}, soldes, libelles=lib)
check("question acceptée explicitement -> divers 400", r["tf1"][12] == 400, r["tf1"][12])

print("== G. dons : plafond 0,25 % du CA signalé (art. 64.2)")
r = f.calculer(E(1_000_000, 10_000_000), {"reintegrations": [{"ligne": "donaciones", "montant": 30_000}]})
check("déductible possible 25 000 (plafond), économie ≈ 6 250",
      any("25,000" in w and "6,250" in w for w in r["avertissements"]), r["avertissements"])

print("== H. échéances (art. 79.1.b, 80.2)")
r = f.calculer(E(1, 1), {})
check("clôture 31/12/2025 -> DEF au 30/06/2026", r["echeances"][0]["date"] == "2026-06-30", r["echeances"][0])
r = f.calculer(E(1, 1, clo="2025-08-31"), {})
check("clôture 31/08/2025 -> DEF au 28/02/2026", r["echeances"][0]["date"] == "2026-02-28", r["echeances"][0])

print("== I. Excel : formules vivantes, aucun taux écrit en dur")
p = os.path.join(tempfile.mkdtemp(), "tf.xlsx")
f.ecrire_xlsx(f.calculer(E(1_000_000, 10_000_000), {}), p)
from openpyxl import load_workbook  # noqa: E402
wb = load_workbook(p)
formules = [c.value for ws in (wb["TF1"], wb["TF2"]) for row in ws.iter_rows() for c in row
            if isinstance(c.value, str) and c.value.startswith("=")]
check("TF1/TF2 portent des formules (≥ 12)", len(formules) >= 12, len(formules))
check("aucun taux littéral (0.25 / 0.015) dans une formule",
      not any(re.search(r"0\.25|0\.015|\b25%|1000\b", x) for x in formules), [x for x in formules if "0." in x])
check("taux lus dans l'onglet Parametros", sum("Parametros!" in x for x in formules) >= 3)

print("== J. profil d'entreprise : interprétation appliquée, écart CHIFFRÉ")
# imposable 1 234 567,89 : strict -> base 1 234 000, IS 308 500 ; sans arrondis -> IS 308 641,97 (écart +141,97)
prof = {"interpretations": {"arrondi_base_is": False, "arrondi_is_franc": False},
        "justifications": {"arrondi_base_is": "pratique historique"}}
r = f.calculer(E(1_234_567.89, 1_000_000), {}, profil=prof)
check("IS selon le profil = 308 641,97", r["tf2"][5] == 308_641.97, r["tf2"][5])
check("écart à la lecture stricte chiffré (+141,97) avec sa justification",
      any("INTERPRÉTATION" in w and "+141.97" in w and "pratique historique" in w for w in r["avertissements"]),
      [w for w in r["avertissements"] if "INTERPR" in w])
r = f.calculer(E(1_234_567.89, 1_000_000), {})
check("sans profil : aucune mention d'interprétation", not any("INTERPRÉTATION" in w for w in r["avertissements"]))
p2 = os.path.join(tempfile.mkdtemp(), "tf2.xlsx")
f.ecrire_xlsx(f.calculer(E(1_234_567.89, 1_000_000), {}, profil=prof), p2)
wbp = load_workbook(p2)["Parametros"]
flags = {r[0].value: r[1].value for r in wbp.iter_rows(min_row=2) if str(r[0].value).startswith("interpretation.")}
check("interprétations retenues visibles dans Parametros (0 = non)",
      flags == {"interpretation.arrondi_base_is": 0, "interpretation.arrondi_is_franc": 0}, flags)

print("== K. questions d'entrée : faits inconnus, indices tirés des libellés")
soldes = {"40110002": -500_000, "41141000": 800_000, "60420000": 100}
lib = {"40110002": "FOURNISSEURS ETRANGERS", "41141000": "CLIENTS ETATS ET COLLECTIVITES"}
r = f.calculer(E(1_000_000, 10_000_000), {}, soldes, libelles=lib)
q = {x["fait"]: x for x in r["questions"]}
check("3 questions posées (nouvelle société, non-résidents, marchés publics)", set(q) == {
    "nouvelle_societe", "paiements_non_residents", "marches_publics"}, list(q))
ind = lambda k: [i["compte"] for i in q.get(k, {}).get("indices", [])]   # sans plantage si la question manque
check("indice non-résidents = 40110002 « FOURNISSEURS ETRANGERS »", ind("paiements_non_residents") == ["40110002"],
      q.get("paiements_non_residents"))
check("indice marchés publics = 41141000", ind("marches_publics") == ["41141000"])
r = f.calculer(E(1_000_000, 10_000_000), {"facturation_non_residents": 2_000_000},
               soldes, libelles=lib, profil={"faits": {"paiements_non_residents": True, "nouvelle_societe": False,
                                                       "marches_publics": True}})
check("faits renseignés -> plus de question", r["questions"] == [], r["questions"])
check("non-résidents : obligation annexe 10 % × 2 000 000 = 200 000 (art. 76), hors TF2",
      r["obligations_annexes"] and r["obligations_annexes"][0]["montant"] == 200_000 and r["tf2"][7] == 250_000,
      (r["obligations_annexes"], r["tf2"][7]))
check("marchés publics : rappel des retenues de 1,5 % (art. 80.3)", any("80.3" in w and "État" in w for w in r["avertissements"]))

print(f"\n{sum(ok)}/{len(ok)} contrôles conformes")
sys.exit(0 if all(ok) else 1)
