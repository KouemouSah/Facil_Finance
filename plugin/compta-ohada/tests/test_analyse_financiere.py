# -*- coding: utf-8 -*-
"""Rapport d'analyse financière pluriannuel (analyse_financiere.py).

Prouve : ratios justes et JAMAIS calculés sur une base absente, nulle ou négative ; années tirées des DEF scannées
avec leurs seuls montants PROUVÉS (colonne N d'une DEF prioritaire sur la colonne N-1 d'une autre) ; variations
significatives = lignes de détail (pas les totaux), au-dessus des deux seuils, expliquées par les comptes qui portent
l'écart ; cause JAMAIS inventée (« à documenter ») et, une fois donnée, reprise avec son origine ; intégration à la
génération (livrable, questions facultatives).
Lancer : python tests/test_analyse_financiere.py      (code 0 = conforme)
"""
import glob
import os
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "tests"))
sys.path.insert(0, os.path.join(ROOT, "skills", "dsf-syscohada", "scripts"))
import analyse_financiere as af  # noqa: E402

ok = []


def check(nom, cond, detail=""):
    ok.append(bool(cond))
    print(f"  {'OK   ' if cond else 'ÉCHEC'} {nom}" + ("" if cond else f"  — {str(detail)[:400]}"))


def main():
    print("== 1. Ratios")
    v = {"XB": 1000, "XD": 200, "XI": 100, "XC": 400, "RK": -100, "CP": 500, "BZ": 2000, "DD": 300, "DT": 100,
         "BK": 600, "BT": 200, "DP": 400, "BI": 250, "DJ": 90, "RA": -100, "RC": -100, "RE": -160, "BB": 50,
         "XE": 150, "DF": 800, "AZ": 900}
    R = af.ratios(v)
    check("marge d'EBE, marge nette, autonomie, endettement, liquidité justes",
          R["Marge d'EBE (EBE / CA)"] == 0.2 and R["Marge nette (résultat / CA)"] == 0.1
          and R["Autonomie financière (capitaux propres / total bilan)"] == 0.25
          and R["Endettement (dettes financières + trésorerie passif) / capitaux propres"] == 0.8
          and R["Liquidité générale (actif circulant + trésorerie) / passif à court terme"] == 1.6, R)
    check("délais en jours (360) : clients 90, fournisseurs 90", R["Délai clients (jours de CA)"] == 90
          and R["Délai fournisseurs (jours d'achats)"] == 90, R)
    check("fonds de roulement, BFR, trésorerie nette", R["Fonds de roulement (ressources stables − actif immobilisé)"] == -100
          and R["Besoin en fonds de roulement (actif circulant − passif circulant)"] == 200
          and R["Trésorerie nette (trésorerie actif − trésorerie passif)"] == 100, R)
    R2 = af.ratios({**v, "XC": -50, "XB": 0, "XD": -10})
    check("base négative ou nulle : ratio NON calculé (valeur ajoutée < 0, CA = 0, EBE < 0)",
          R2["Charges de personnel / valeur ajoutée"] is None and R2["Marge d'EBE (EBE / CA)"] is None
          and R2["Capacité de remboursement (dettes financières / EBE), années"] is None, R2)
    check("donnée absente : ratio NON calculé", af.ratios({"XB": 100})["Marge nette (résultat / CA)"] is None)

    print("== 2. Années depuis les DEF scannées")
    D = {"pN": {"XB": 1000, "BZ": 2000}, "p1": {"BZ": 1800, "XB": 900}, "src_n1": "ouvertures"}

    def scan(ex, n, n1, prouve=True):
        return {"exercice": ex, "etats": {c: {"n": {"valeur": a, "prouve": prouve}, "n1": {"valeur": b, "prouve": prouve}}
                                          for c, (a, b) in {"XB": (n, n1), "RK": (-n // 10, -n1 // 10)}.items()}}
    A = af.annees_depuis(D, 2025, [("d2024.pdf", scan(2024, 900, 800)), ("d2023.pdf", scan(2023, 805, 700))], "scan 2024")
    check("années 2022 à 2025", sorted(A) == [2022, 2023, 2024, 2025], sorted(A))
    check("N et N-1 viennent du calcul (balance, ouvertures + scan retenu)", A[2025]["v"]["XB"] == 1000 and A[2024]["v"]["XB"] == 900)
    check("colonne N d'une DEF prioritaire sur la colonne N-1 d'une autre (2023 : 805 de la DEF 2023, pas 800)",
          A[2023]["v"]["XB"] == 805 and "DEF 2023" in A[2023]["source"] and "colonne N " in A[2023]["source"], A[2023])
    check("charges ramenées en négatif comme dans le calcul", A[2022]["v"]["RK"] < 0, A[2022]["v"])
    B = af.annees_depuis(D, 2025, [("d2023.pdf", scan(2023, 805, 700, prouve=False))], None)
    check("montants NON prouvés d'un scan : jamais repris", sorted(B) == [2024, 2025], sorted(B))

    print("== 3. Variations expliquées")
    postes = {"BI": {"libelle": "Clients"}, "BK": {"libelle": "TOTAL ACTIF CIRCULANT", "total": ["BI"]},
              "BZ": {"libelle": "TOTAL", "total": ["BK"]}, "RK": {"libelle": "Personnel"}, "XB": {"libelle": "CA", "total": ["TC"]},
              "TC": {"libelle": "Travaux"}, "BB": {"libelle": "Stocks"}}
    an = {2024: {"v": {"BI": 100, "BK": 100, "BZ": 1000, "RK": -100, "TC": 1000, "XB": 1000, "BB": 100}},
          2025: {"v": {"BI": 200, "BK": 200, "BZ": 1100, "RK": -105, "TC": 1000, "XB": 1000, "BB": 104}}}
    cN = {"BI": {"411100": 150, "411200": 50}}
    c1 = {"BI": {"411100": 60, "411200": 40}}
    V = af.variations(an, postes, cN, c1, {"411100": "CLIENTS LOCAUX"}, {})
    refs = [x["ref"] for x in V]
    check("totaux exclus (BK, BZ) ; ligne de détail retenue (BI)", refs == ["BI"], refs)
    check("sous les seuils (RK +5 %, BB +4 %) : non retenues", "RK" not in refs and "BB" not in refs)
    x = V[0]
    check("faits : le compte qui porte l'écart, avec sa part", "411100 CLIENTS LOCAUX +90" in x["faits"] and "90%" in x["faits"], x["faits"])
    check("points à vérifier propres au poste (clients : provisions, encaissements, coupure)",
          "créances anciennes" in x["verifier"], x["verifier"])
    check("cause JAMAIS inventée sans réponse", x["cause"] is None)

    print("== 4. Rapport (classeur)")
    import openpyxl
    t = tempfile.mkdtemp()
    p = os.path.join(t, "a.xlsx")
    Dx = {"postes": postes, "libelles": {"411100": "CLIENTS LOCAUX"}, "agg": {"detail": [("411100", "BI", "brut", 150)]},
          "agg1": {"detail": [("411100", "BI", "brut", 60)]}}
    r = af.construire(Dx, {k: {**w, "source": f"source {k}"} for k, w in an.items()}, p, None,
                      {"BI": {"texte": "Retard de paiement de l'État", "origine": "réponse de l'entreprise, le 2026-10-19"}}, "ALPHA")
    wb = openpyxl.load_workbook(p)
    check("six feuilles", wb.sheetnames == ["Synthèse", "Bilan comparé", "Compte de résultat comparé", "Soldes et ratios",
                                            "Variations expliquées", "Sources et méthode"], wb.sheetnames)
    texte = " | ".join(str(c) for ws in wb.worksheets for row in ws.iter_rows(values_only=True) for c in row if c is not None)
    check("cause donnée : reprise avec son origine et sa date", "Retard de paiement de l'État — réponse de l'entreprise, le 2026-10-19" in texte)
    check("donnée absente affichée « — », jamais estimée", "—" in texte)
    check("sources de chaque année citées", "source 2024" in texte and "source 2025" in texte)

    print("== 5. Génération")
    import test_reprise_validation as trv
    ed = trv.ed
    reel = ed.generer
    ed.generer = lambda d, **k: reel(d, **{**k, "pdf": False})
    d, _ = trv.balance_seule()
    g = ed.generer(d)
    af_res = g.get("analyse_financiere", {})
    livr = ed.lire_json(os.path.join(d, "manifeste.json"), {})["versions"][-1]["livrables"]
    check("rapport produit à chaque version, dans les livrables", af_res.get("produite") and any("analyse-financiere" in x for x in livr), af_res)
    qs = [q for q in ed.lire_json(os.path.join(d, "04_questions", "questions.json"), []) if q["id"].startswith("analyse.")]
    check("causes demandées en questions FACULTATIVES, sans total", qs and all(not q["obligatoire"] for q in qs)
          and not any(q["id"] in ("analyse.BZ", "analyse.DZ", "analyse.BT", "analyse.CP") for q in qs), [q["id"] for q in qs])
    ed.repondre(d, qs[0]["id"], "Encaissement d'un gros chantier en fin d'année", canal="formulaire")
    g = ed.generer(d)
    x = glob.glob(os.path.join(d, "02_xls", "*_v2_analyse-financiere*.xlsx"))
    texte = " | ".join(str(c) for ws in openpyxl.load_workbook(x[0]).worksheets for row in ws.iter_rows(values_only=True)
                       for c in row if c is not None) if x else ""
    check("réponse reprise au rapport de la version suivante, avec son origine",
          "Encaissement d'un gros chantier en fin d'année — réponse de l'entreprise, le " in texte, x)
    check("la validation n'attend PAS ces questions (facultatives)", not any(q["id"].startswith("analyse.")
                                                                          for q in ed.ouvertes(ed.lire_json(os.path.join(d, "04_questions", "questions.json"), []))))
    print(f"{sum(ok)}/{len(ok)} contrôles conformes")
    return 0 if all(ok) else 1


if __name__ == "__main__":
    sys.exit(main())
