# -*- coding: utf-8 -*-
"""Garde du générateur de DEF (skills/dsf-syscohada/scripts/def_ge.py).

Balance 6 colonnes (ouverture, mouvements, soldes) construite À LA MAIN, écritures connues :
  vente 1 000 000 encaissée ; achat 450 000 à crédit ; règlement fournisseur 400 000 ;
  acquisition de matériel 100 000 payée ; dotation 70 000 ; report à nouveau débiteur 200 000 à l'ouverture.
Attendus calculés à la main (commentaires). Aucun impôt en comptabilité : il doit être INTÉGRÉ.
Lancer : python tests/test_def_ge.py     (code 0 = conforme ; le PDF n'est pas testé ici : dépend d'Excel)
"""
import argparse
import csv
import json
import os
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "skills", "dsf-syscohada", "scripts"))
import def_ge  # noqa: E402

# compte : (ouverture D-C, mouvement débit, mouvement crédit)
B = {"10130000": (-1_000_000, 0, 0), "12100000": (200_000, 0, 0),
     "24100000": (600_000, 100_000, 0), "28410000": (-200_000, 0, 70_000),
     "52110000": (500_000, 1_000_000, 500_000), "40110000": (-100_000, 400_000, 450_000),
     "70510000": (0, 0, 1_000_000), "60400000": (0, 450_000, 0), "68130000": (0, 70_000, 0)}
LIB = {"10130000": "CAPITAL", "12100000": "REPORT A NOUVEAU", "24100000": "MATERIEL", "28410000": "AMORT MATERIEL",
       "52110000": "BANQUE", "40110000": "FOURNISSEURS", "70510000": "TRAVAUX", "60400000": "ACHATS MATIERES",
       "68130000": "DOTATIONS AMORT"}

ok = []


def check(nom, cond, detail=""):
    ok.append(bool(cond))
    print(f"  {'OK   ' if cond else 'ÉCHEC'} {nom}" + ("" if cond else f"  — {detail}"))


def main():
    d = tempfile.mkdtemp()
    p = os.path.join(d, "balance.csv")
    with open(p, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Compte", "Intitulé", "Ouverture D", "Ouverture C", "Mvt D", "Mvt C", "Solde D", "Solde C"])
        for a, (o, md, mk) in B.items():
            s = o + md - mk
            w.writerow([a, LIB[a], max(o, 0), max(-o, 0), md, mk, max(s, 0), max(-s, 0)])
    assert sum(o + md - mk for o, md, mk in B.values()) == 0
    prof = os.path.join(d, "profil.json")
    json.dump({"identite": {"raison_sociale": "ENTITE TEST SA", "nif": "TEST-01"}, "faits": {"nouvelle_societe": False}},
              open(prof, "w", encoding="utf-8"))
    a = argparse.Namespace(balance=p, exercice=2025, profil=prof, decisions=None, grand_livre=None, tiers=None,
                           def_precedente=None, colonne_precedente="N", gabarit=None)
    D = def_ge.calculer(a)
    pN, t, n3 = D["pN"], D["tft"], D["n3"]
    print("== IS intégré (résultat avant impôt 480 000 -> IS 25 % = 120 000 > CMF 15 000)")
    check("RS = -120 000, XI = CJ = 360 000, DK = 120 000",
          (pN["RS"], pN["XI"], pN["CJ"], pN["DK"]) == (-120_000, 360_000, 360_000, 120_000), (pN["RS"], pN["XI"], pN["CJ"], pN["DK"]))
    check("écriture simulée signalée (à passer en comptabilité)", D["info_is"]["mode"] == "ecriture_simulee")
    print("== Bilan")
    # actif : AM net 700 000 - 270 000 = 430 000 ; BS 1 000 000 -> BZ 1 430 000
    # passif : CA 1 000 000 ; CH -200 000 ; CJ 360 000 ; DJ 150 000 ; DK 120 000 -> DZ 1 430 000
    check("BZ = DZ = 1 430 000", pN["BZ"] == pN["DZ"] == 1_430_000, (pN["BZ"], pN["DZ"]))
    print("== TFT (N-1 = soldes d'ouverture)")
    # ZA 500 000 ; FA = 360 000 + 70 000 = 430 000 ; FE = (150 000 + 120 000) - 100 000 = 170 000 ; ZB 600 000
    # FG -100 000 ; FN 0 ; ZH = 500 000 + 600 000 - 100 000 = 1 000 000 = trésorerie du bilan
    attendu = {"ZA": 500_000, "FA": 430_000, "FE": 170_000, "ZB": 600_000, "FG": -100_000, "ZC": -100_000,
               "FN": 0, "ZG": 500_000, "ZH": 1_000_000}
    check("TFT : ZA, FA, FE, ZB, FG, ZC, FN, ZG, ZH", all(abs(t[k] - v) < 0.01 for k, v in attendu.items()),
          {k: t[k] for k in attendu})
    check("TFT réconcilié (écart 0)", abs(D["tft_ecart"]) < 0.01, D["tft_ecart"])
    print("== Note 3")
    am = n3["AM"]
    check("AM : brut 600 000 + 100 000 = 700 000 ; amort 200 000 + 70 000 = 270 000",
          (am["b_ouv"], am["b_aug"], am["b_clo"], am["a_ouv"], am["a_dot"], am["a_clo"]) == (600_000, 100_000, 700_000, 200_000, 70_000, 270_000), am)
    print("== Contrôles et couverture")
    ko = [x["controle"] for x in D["controles"] if not x["ok"]]
    check("tous les contrôles au vert", not ko, ko)
    check("identité incomplète déclarée « partiel »", any(el.startswith("Couverture") and st == "partiel" for el, st, _ in D["couverture"]))
    print("== Classeur")
    x = os.path.join(d, "DEF.xlsx")
    feuilles = def_ge.ecrire_xlsx(D, x)
    check("feuilles de la DEF", {"Couverture", "Bilan actif", "Bilan passif", "Compte de résultat", "TFT", "Note 3", "Notes", "TF1", "TF2"}
          <= set(feuilles), feuilles)
    from openpyxl import load_workbook
    wb = load_workbook(x)
    check("annexe interne « Controles » présente mais HORS liste DEF (jamais exportée avec la DEF)",
          "Controles" in wb.sheetnames and "Controles" not in feuilles)
    f = [c.value for n in ("Bilan actif", "Bilan passif", "Compte de résultat", "TFT") for r in wb[n].iter_rows() for c in r
         if isinstance(c.value, str) and c.value.startswith("=")]
    check("totaux en formules vivantes (≥ 30)", len(f) >= 30, len(f))
    print(f"\n{sum(ok)}/{len(ok)} contrôles conformes")
    sys.exit(0 if all(ok) else 1)


if __name__ == "__main__":
    main()
